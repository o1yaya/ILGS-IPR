from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph


SOURCE = Path(r"D:\wlxy\基于改进ILGS的开放词汇目标抓取系统.docx")
OUTPUT = Path(r"D:\wlxy\基于改进ILGS的开放词汇目标抓取系统_大论文第2章相关工作_20261019.docx")


INTRO = (
    "本文面向由自然语言目标描述到机械臂抓取执行的完整链路，相关研究可归纳为四个彼此衔接的方向："
    "开放词汇三维场景表示负责把二维视觉语言知识提升到三维空间；语言引导与目标驱动抓取负责将语义目标转化为抓取对象或对象部件；"
    "基于3D Gaussian Splatting（3DGS）的机器人操作进一步利用显式高斯表示的渲染效率、几何属性与可编辑性；"
    "抓取可行性和运动规划则保证候选姿态能够由具体机械臂安全执行。以下围绕这四条研究线梳理代表性工作，并说明它们与本文研究范围的关系。"
)


SECTIONS = [
    (
        "2.1 开放词汇三维场景表示",
        [
            "开放词汇三维场景表示的目标，是把视觉语言模型在二维图像中的语义知识组织为可跨视角查询的三维表示，使用户无需预先限定类别集合即可通过自然语言定位物体、属性或部件。该方向需要同时处理二维语义特征到三维空间的提升、不同视角间的一致性以及对象边界的保持。早期方法多以神经辐射场为载体，随后逐渐转向能够显式表示场景并进行快速光栅化的3DGS。",
            "LERF将CLIP的多尺度语言特征蒸馏到NeRF中，通过沿训练射线体渲染语言嵌入，构建可由文本交互查询的稠密三维语言场[1]。其优势是无需预定义类别或对象提议即可响应长尾概念、视觉属性和抽象描述，但NeRF的隐式表示在训练与更新效率方面受限，而且连续相关性场本身缺少明确的对象身份边界。当语言场被用于下游抓取时，局部高响应区域不一定等同于完整、可操作的目标实例。",
            "ILGS进一步将身份信息引入语言高斯表示，使同一身份对应的语言嵌入在潜在空间内相互接近，不同身份相互分离，并通过渐进式掩码扩展改善目标边界[2]。与仅依赖像素级语言特征的3D语言场相比，身份感知表示更适合输出对象级掩码，是本文选用ILGS作为上游感知框架的直接原因。然而，ILGS的身份分类分支依赖由二维分割与跟踪产生的多视角身份伪标签；当标签包含背景泄漏、边界偏移、相邻物体粘连或错误身份时，这些误差会被写入三维身份场。本文提出的IPR正是针对这一输入监督可靠性问题，而不是重新设计语言编码器或3DGS主干。",
            "GaussianGrasper把3DGS进一步用于开放词汇机器人抓取。该方法从少量RGB-D视角重建兼具语义与几何信息的高斯场，通过高效特征蒸馏将CLIP特征和SAM分割先验注入高斯，并利用渲染法向对预训练抓取模型生成的候选姿态进行筛选[3]。它表明显式高斯表示能够同时服务于语言定位和几何抓取，但其目标定位质量仍取决于蒸馏语义与对象边界是否可靠。",
            "GraspSplats采用深度监督和参考特征计算构建特征增强高斯表示，并直接利用高斯的语义与几何属性完成对象部件定位、实时抓取采样以及动态场景中的点跟踪[4]。这类工作说明，3DGS不仅是新视角合成工具，也可以作为感知与操作共享的中间表示。与GaussianGrasper和GraspSplats不同，本文不以提出新的高斯抓取表示为主要目标，而是先提高ILGS身份监督与目标Gaussian掩码的可靠性，再将较干净的目标区域交给现有抓取后端，以验证上游表示质量对下游抓取的影响。",
        ],
    ),
    (
        "2.2 语言引导与目标驱动机器人抓取",
        [
            "语言引导抓取要求机器人根据自然语言识别目标对象或对象部件，并在开放类别和杂乱环境中生成合适的抓取姿态。与传统类别级抓取相比，这类系统不仅要回答“抓哪个物体”，还要处理属性、空间关系、用途和部件等开放描述。目标驱动抓取则进一步强调候选姿态必须集中在指定目标上，而不是仅从整幅场景中寻找任意稳定抓取。",
            "LERF-TOGO在LERF的三维语言场上引入对象级空间分组：首先利用DINO特征提取三维对象掩码，再在该掩码内进行条件语言查询，用文本描述的对象部件对现成抓取规划器输出的候选进行重排序[5]。这一设计缓解了LERF缺少对象性的局限，并把开放词汇查询从目标定位推进到任务导向部件抓取。不过，其目标掩码仍由单独的对象分组过程提供，感知误差会直接影响后续条件查询和抓取分布。GaussianGrasper则在统一的3D高斯场中完成语言定位、候选生成与法向引导筛选[3]，体现了语义与几何联合建模的另一条路线。",
            "近年来，视觉语言模型也被用于更开放的意图理解与抓取决策。OVGrasp融合RGB-D视觉、开放词汇提示和语音命令，在可穿戴软体外骨骼上进行多模态意图识别与抓取辅助[6]；其研究重点是用户意图、抓取与释放状态以及辅助设备的人机交互，与桌面机械臂的六自由度目标抓取在执行机构和规划约束上有所不同。OWG则采用视觉语言模型、实例分割与抓取合成模型，依次完成开放描述指代、基于场景关系的抓取规划以及接触推理驱动的候选排序，并能在目标受阻时先移除遮挡物[7]。这说明开放词汇抓取正在从“类别识别加抓取”转向语义、关系和接触几何的联合推理。",
            "目标驱动抓取还必须处理杂乱环境中的遮挡与碰撞。Murali等提出的六自由度目标驱动抓取方法从部分点云中定位指定对象，并利用学习式碰撞检测模块评估候选抓取与场景的碰撞关系，从而在未知物体和杂乱桌面中规划抓取与放置[8]。该类方法通常假设目标点云或目标实例已经较为准确，因此其性能上限受到上游目标分割质量约束。",
            "综上，现有语言抓取方法已形成“目标语义定位—对象或部件分割—候选抓取生成—候选排序”的基本范式，但对上游目标掩码的可靠性处理通常不是研究中心。本文沿用这一任务链路，把改进ILGS得到的目标Gaussian掩码转换为目标局部点云，再使用统一抓取后端生成候选；通过与原始ILGS掩码和整场景点云对比，单独检验身份伪标签改进是否能够降低候选落在背景或相邻物体上的概率。",
        ],
    ),
    (
        "2.3 基于3DGS的机器人操作",
        [
            "3DGS具有显式基元、快速可微光栅化和可直接编辑等特点，使其逐渐从静态重建扩展到机器人感知、抓取和动态场景更新。相较于隐式NeRF，单个高斯携带位置、尺度、旋转、不透明度与外观属性，便于附加语义、可供性或物理状态，并可从三维基元直接构建对象集合和局部几何。",
            "Splat-MOVER是较完整的3DGS开放词汇操作系统，由ASK-Splat、SEE-Splat和Grasp-Splat三个模块组成[9]。ASK-Splat在高斯场中蒸馏语言语义与抓取可供性；SEE-Splat通过三维语义掩码和场景编辑更新数字孪生；Grasp-Splat调用抓取模型生成候选，并依据语义与可供性信息重排序。该系统能够支持多阶段操作，但也需要较完整的语义、可供性、场景编辑和机器人执行模块。本文的研究范围更收敛，重点是目标身份表示的可靠性以及它对单阶段抓取候选和执行成功率的影响。",
            "GraspSplats利用特征增强高斯直接支持对象部件抓取、实时候选采样和通过点跟踪实现的动态操作[4]，体现了显式高斯基元在效率与可操作性上的优势。Physically Embodied Gaussian Splatting则提出Gaussian-Particle双重表示：粒子承担几何与物理模拟，高斯负责视觉渲染，再利用观测图像产生校正作用，使表示能够在动态交互中预测未来状态并与真实场景同步[10]。这条路线关注物理一致性和在线可校正世界模型，超出了本文当前静态桌面场景和抓取前重建的设定。",
            "现有3DGS操作研究分别从语义、可供性、动态编辑和物理建模扩展了高斯表示，但系统越完整，需要控制的变量也越多。本文不直接复现完整的Splat-MOVER或物理高斯世界模型，而是在ILGS身份感知高斯场上增加可靠性建模，并把输出接入可替换的抓取和规划后端。这样既保留3DGS作为共享三维表示的优势，也能通过受控实验明确区分“目标表示改善”和“抓取规划器能力”各自带来的影响。",
        ],
    ),
    (
        "2.4 抓取可行性与运动规划",
        [
            "从目标点云生成的高分抓取姿态并不必然能够由具体机械臂执行。候选还需要满足逆运动学可解、夹爪与机械臂无碰撞、接近和撤离方向可行、轨迹长度与关节运动合理等条件。因此，抓取系统通常需要把候选生成与可行性验证分开：前者估计抓取稳定性，后者结合机器人模型和场景约束筛除不可执行姿态。",
            "MoveIt为机械臂操作提供了运动学、碰撞检测与运动规划框架[11]。其中MoveIt Grasps按照“候选生成—可达性过滤—接近、抬升和撤离轨迹规划”的顺序组织抓取流程，可基于逆运动学和笛卡尔路径验证候选[12]。其几何抓取生成主要面向长方体或圆柱等简化对象，并不直接建模开放词汇语义、摩擦锥或复杂抓取动力学。因此，在本文系统中，MoveIt更适合作为统一的执行约束验证工具，而不是语义抓取方法本身。",
            "面向动态目标，Akinola等提出可达性与运动感知抓取框架：以有符号距离场快速描述机器人可达空间，以运动条件抓取质量筛选候选，并通过目标运动预测和上一时刻轨迹热启动实现在线重规划[13]。该方法表明，静态抓取质量不能替代对机器人可达性和目标运动的显式考虑。本文当前聚焦静态桌面场景，但其“先快速过滤、再对少量候选规划”的思想可用于降低多候选验证成本。",
            "在多物体杂乱场景中，Zhao等以obstruction degree刻画其他物体对候选抓取和目标对象的阻碍程度，并联合规划抓取顺序与具体抓取构型，以减少因支撑关系和相互遮挡导致的不安全操作[14]。该方法强调场景级对象关系，但假设物体三维模型、身份和位姿已知。对于开放词汇系统，这些前提再次说明目标身份和三维区域必须足够可靠，否则阻挡关系和候选可达性都会建立在错误对象上。",
            "神经运动规划研究尝试用端到端策略、学习式采样、轨迹优化热启动和神经碰撞检测提高高维规划效率。相关综述指出，神经规划器具有推理快和表达多模态解的潜力，但在分布外场景泛化、安全保证和跨机器人迁移方面仍存在挑战[15]。基于这一认识，本文不重新训练通用运动规划器，而采用“硬约束剔除+软评分排序”的Grasp Feasibility Ranking：先删除IK无解、明显碰撞或超出工作空间的候选，再综合抓取质量、IPR目标置信度、语义一致性、碰撞裕量和轨迹代价排序。该设计把本文创新集中在感知置信度与执行约束的接口，同时保留成熟规划器的安全验证能力。",
        ],
    ),
]


REFERENCES = [
    "[1] KERR J, KIM C M, GOLDBERG K, et al. LERF: Language Embedded Radiance Fields[C]//Proceedings of the IEEE/CVF International Conference on Computer Vision. 2023: 19729-19739.",
    "[2] JANG S, KIM W. Identity-aware Language Gaussian Splatting for Open-vocabulary 3D Semantic Segmentation[C]//Proceedings of the IEEE/CVF International Conference on Computer Vision. 2025: 20467-20476.",
    "[3] ZHENG Y, CHEN X, ZHENG Y, et al. GaussianGrasper: 3D Language Gaussian Splatting for Open-Vocabulary Robotic Grasping[J]. IEEE Robotics and Automation Letters, 2024, 9(9): 7827-7834. DOI: 10.1109/LRA.2024.3432348.",
    "[4] JI M, QIU R Z, ZOU X, et al. GraspSplats: Efficient Manipulation with 3D Feature Splatting[C]//Proceedings of the 8th Conference on Robot Learning. PMLR, 2025, 270: 1443-1460.",
    "[5] RASHID A, SHARMA S, KIM C M, et al. Language Embedded Radiance Fields for Zero-Shot Task-Oriented Grasping[C]//Proceedings of the 7th Conference on Robot Learning. PMLR, 2023, 229: 178-200.",
    "[6] HU C, LUO S, GIONFRIDA L. OVGrasp: Open-Vocabulary Grasping Assistance via Multimodal Intent Detection[EB/OL]. arXiv:2509.04324, 2025.",
    "[7] TZIAFAS G, KASAEI H. Towards Open-World Grasping with Large Vision-Language Models[C]//Proceedings of the 8th Conference on Robot Learning. PMLR, 2025, 270: 3304-3332.",
    "[8] MURALI A, MOUSAVIAN A, EPPNER C, et al. 6-DOF Grasping for Target-driven Object Manipulation in Clutter[C]//2020 IEEE International Conference on Robotics and Automation. 2020: 6232-6238. DOI: 10.1109/ICRA40945.2020.9197318.",
    "[9] SHORINWA O, TUCKER J, SMITH A, et al. Splat-MOVER: Multi-Stage, Open-Vocabulary Robotic Manipulation via Editable Gaussian Splatting[C]//Proceedings of the 8th Conference on Robot Learning. PMLR, 2025, 270: 4748-4770.",
    "[10] ABOU-CHAKRA J, RANA K, DAYOUB F, et al. Physically Embodied Gaussian Splatting: A Visually Learnt and Physically Grounded 3D Representation for Robotics[C]//Proceedings of the 8th Conference on Robot Learning. PMLR, 2025, 270: 513-530.",
    "[11] COLEMAN D, SUCAN I A, CHITTA S, et al. Reducing the Barrier to Entry of Complex Robotic Software: A MoveIt! Case Study[J]. Journal of Software Engineering for Robotics, 2014, 5(1): 3-16.",
    "[12] MOVEIT. MoveIt Grasps[EB/OL]. https://moveit.picknik.ai/main/doc/examples/moveit_grasps/moveit_grasps_tutorial.html, 2026-10-19.",
    "[13] AKINOLA I, XU J, SONG S, et al. Dynamic Grasping with Reachability and Motion Awareness[C]//2021 IEEE/RSJ International Conference on Intelligent Robots and Systems. 2021: 9422-9429.",
    "[14] ZHAO W, WANG J, CHEN W, et al. Planning for Grasping Cluttered Objects Based on Obstruction Degree[J]. International Journal of Advanced Robotic Systems, 2021, 18(6): 1-13. DOI: 10.1177/17298814211040632.",
    "[15] SOLEYMANZADEH D, LOPEZ-SANCHEZ I, SU H, et al. Toward Generalist Neural Motion Planners for Robotic Manipulators: Challenges and Opportunities[J]. IEEE Transactions on Automation Science and Engineering, 2026, 23: 4488-4531. DOI: 10.1109/TASE.2026.3660830.",
]


def set_keep_with_next(paragraph, value=True):
    ppr = paragraph._p.get_or_add_pPr()
    keep = ppr.find(qn("w:keepNext"))
    if value and keep is None:
        keep = OxmlElement("w:keepNext")
        ppr.append(keep)
    elif not value and keep is not None:
        ppr.remove(keep)


def set_keep_lines(paragraph):
    ppr = paragraph._p.get_or_add_pPr()
    if ppr.find(qn("w:keepLines")) is None:
        ppr.append(OxmlElement("w:keepLines"))


def insert_before(anchor, text, style):
    paragraph = anchor.insert_paragraph_before(text, style=style)
    if style.startswith("Heading"):
        set_keep_with_next(paragraph)
    else:
        set_keep_lines(paragraph)
    return paragraph


def insert_after(anchor, text, style="Normal"):
    node = OxmlElement("w:p")
    anchor._p.addnext(node)
    paragraph = Paragraph(node, anchor._parent)
    paragraph.style = style
    paragraph.add_run(text)
    return paragraph


def add_page_break_before(paragraph):
    ppr = paragraph._p.get_or_add_pPr()
    if ppr.find(qn("w:pageBreakBefore")) is None:
        ppr.append(OxmlElement("w:pageBreakBefore"))


def main():
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)

    doc = Document(SOURCE)

    for p in doc.paragraphs:
        if p.text.strip() == "2. 研究问题链路":
            p.text = "2. 相关工作"

    section_heading = next(
        p for p in doc.paragraphs if p.style.name == "Heading 1" and p.text.strip() == "2. 相关工作"
    )
    next_heading = next(
        p for p in doc.paragraphs if p.style.name == "Heading 1" and p.text.strip() == "3. 三个核心研究问题"
    )
    set_keep_with_next(section_heading)
    add_page_break_before(next_heading)

    insert_before(next_heading, INTRO, "Normal")
    for heading, paragraphs in SECTIONS:
        insert_before(next_heading, heading, "Heading 2")
        for text in paragraphs:
            insert_before(next_heading, text, "Normal")

    # The original second table summarizes the three-problem chain. Move it to
    # the beginning of section 3 so the existing planning material is preserved.
    chain_table = doc.tables[1]._element
    parent = chain_table.getparent()
    parent.remove(chain_table)
    next_heading._p.addnext(chain_table)

    ref_heading = doc.add_paragraph("11. 参考文献", style="Heading 1")
    add_page_break_before(ref_heading)
    set_keep_with_next(ref_heading)
    for entry in REFERENCES:
        p = doc.add_paragraph(entry, style="Normal")
        set_keep_lines(p)

    # Refresh the simple, manually written contents list at the front.
    first_main_heading = next(
        i for i, p in enumerate(doc.paragraphs)
        if p.style.name == "Heading 1" and p.text.strip() == "1. 总体定位"
    )
    front = doc.paragraphs[:first_main_heading]
    for p in list(front):
        if p.text.strip() in {"9. 可行性与降级方案", "10. 最终建议", "11. 参考文献"}:
            p._element.getparent().remove(p._element)
    toc_prompt = next(
        p for p in doc.paragraphs[:first_main_heading]
        if p.text.strip() == "8. 可直接使用的 Prompt"
    )
    toc_anchor = insert_after(toc_prompt, "9. 可行性与降级方案")
    toc_anchor = insert_after(toc_anchor, "10. 最终建议")
    insert_after(toc_anchor, "11. 参考文献")

    doc.core_properties.title = "基于改进ILGS的开放词汇目标抓取系统：相关工作扩充版"
    doc.core_properties.subject = "开放词汇三维表示、语言引导抓取、3DGS机器人操作与抓取可行性规划"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
