"""Update the ILGS-IPR draft with the completed 3D-OVS three-seed study."""

from __future__ import annotations

import csv
import math
import statistics
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(r"D:\PycharmProjects\PythonProject\ILGS_release-main")
SOURCE = (
    ROOT
    / "outputs/paper_revision_20260930/ILGS-IPR_中文_3DOVS十场景补充_20260930.docx"
)
OUTPUT = (
    ROOT
    / "outputs/paper_revision_20261007/ILGS-IPR_中文_3DOVS三种子补充_20261007.docx"
)
RAW_RESULTS = Path(r"D:\downloads\controlled_ablation_raw_3dovs.csv")

SCENE_ORDER = [
    ("bed", "Bed"),
    ("bench", "Bench"),
    ("blue_sofa", "Blue Sofa"),
    ("covered_desk", "Covered Desk"),
    ("lawn", "Lawn"),
    ("office_desk", "Office Desk"),
    ("room", "Room"),
    ("snacks", "Snacks"),
    ("sofa", "Sofa"),
    ("table", "Table"),
]


def paragraph_with_prefix(doc: Document, prefix: str):
    matches = [p for p in doc.paragraphs if p.text.startswith(prefix)]
    if len(matches) != 1:
        raise ValueError(
            f"Expected one paragraph starting with {prefix!r}, found {len(matches)}"
        )
    return matches[0]


def replace_text(paragraph, text: str) -> None:
    if not paragraph.runs:
        paragraph.add_run(text)
        return
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def insert_after(anchor, text: str, template):
    element = OxmlElement("w:p")
    if template._p.pPr is not None:
        element.append(deepcopy(template._p.pPr))
    anchor._p.addnext(element)
    from docx.text.paragraph import Paragraph

    paragraph = Paragraph(element, anchor._parent)
    run = paragraph.add_run(text)
    if template.runs and template.runs[0]._r.rPr is not None:
        run._r.insert(0, deepcopy(template.runs[0]._r.rPr))
    return paragraph


def load_results():
    rows = list(csv.DictReader(RAW_RESULTS.open(encoding="utf-8-sig")))
    expected = {
        (scene, variant, seed)
        for scene, _ in SCENE_ORDER
        for variant in ("official", "ipr")
        for seed in (0, 1, 2)
    }
    observed = {
        (row["scene"], row["variant"], int(row["seed"])) for row in rows
    }
    if observed != expected:
        missing = sorted(expected - observed)
        extra = sorted(observed - expected)
        raise ValueError(f"Unexpected result grid; missing={missing}, extra={extra}")

    data = {
        (row["scene"], row["variant"], int(row["seed"])): (
            float(row["mIoU"]),
            float(row["mBIoU"]),
        )
        for row in rows
    }

    summary = {}
    pair_diffs = []
    scene_diffs = []
    for scene, _ in SCENE_ORDER:
        for variant in ("official", "ipr"):
            values = [data[(scene, variant, seed)] for seed in (0, 1, 2)]
            summary[(scene, variant)] = (
                statistics.fmean(value[0] for value in values),
                statistics.stdev(value[0] for value in values),
                statistics.fmean(value[1] for value in values),
                statistics.stdev(value[1] for value in values),
            )

        diffs = []
        for seed in (0, 1, 2):
            official = data[(scene, "official", seed)]
            ipr = data[(scene, "ipr", seed)]
            diff = (ipr[0] - official[0], ipr[1] - official[1])
            pair_diffs.append((scene, seed, *diff))
            diffs.append(diff)
        scene_diffs.append(
            (
                scene,
                statistics.fmean(item[0] for item in diffs),
                statistics.fmean(item[1] for item in diffs),
            )
        )

    overall = {}
    for variant in ("official", "ipr"):
        overall[variant] = (
            statistics.fmean(summary[(scene, variant)][0] for scene, _ in SCENE_ORDER),
            statistics.fmean(summary[(scene, variant)][2] for scene, _ in SCENE_ORDER),
        )

    stats = {
        "pair_wins_miou": sum(item[2] > 0 for item in pair_diffs),
        "pair_wins_mbiou": sum(item[3] > 0 for item in pair_diffs),
        "scene_wins_miou": sum(item[1] > 0 for item in scene_diffs),
        "scene_wins_mbiou": sum(item[2] > 0 for item in scene_diffs),
    }
    for name, index in (("miou", 1), ("mbiou", 2)):
        values = [item[index] for item in scene_diffs]
        mean = statistics.fmean(values)
        sd = statistics.stdev(values)
        half_width = 2.262157 * sd / math.sqrt(len(values))
        stats[f"delta_{name}"] = mean
        stats[f"ci_{name}"] = (mean - half_width, mean + half_width)

    return summary, overall, stats


def set_cell(cell, text: str, bold: bool = False, header: bool = False) -> None:
    cell.text = text
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for paragraph in cell.paragraphs:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.first_line_indent = 0
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        for run in paragraph.runs:
            run.bold = bold
            run.font.name = "Times New Roman"
            run.font.size = Pt(7.6 if not header else 8.3)
            run.font.color.rgb = RGBColor(0, 0, 0)
    if header or bold:
        tc_pr = cell._tc.get_or_add_tcPr()
        shade = tc_pr.find(qn("w:shd"))
        if shade is None:
            shade = OxmlElement("w:shd")
            tc_pr.append(shade)
        shade.set(qn("w:fill"), "E7E7E7" if header else "F2F2F2")


def metric_cell(values) -> str:
    return f"{values[0]:.2f} ± {values[1]:.2f} / {values[2]:.2f} ± {values[3]:.2f}"


def update_table(table, summary, overall) -> None:
    for row in list(table.rows)[1:]:
        table._tbl.remove(row._tr)
    table.autofit = False
    widths = (Cm(3.8), Cm(5.4), Cm(5.4))
    for column, width in zip(table.columns, widths):
        column.width = width
    for cell, text in zip(
        table.rows[0].cells,
        ("场景", "ILGS受控基线", "IPR"),
    ):
        set_cell(cell, text, bold=True, header=True)
    header_pr = table.rows[0]._tr.get_or_add_trPr()
    if header_pr.find(qn("w:tblHeader")) is None:
        header_pr.append(OxmlElement("w:tblHeader"))

    for key, label in SCENE_ORDER:
        row = table.add_row()
        for cell, text in zip(
            row.cells,
            (label, metric_cell(summary[(key, "official")]), metric_cell(summary[(key, "ipr")])),
        ):
            set_cell(cell, text)

    average = table.add_row()
    for cell, text in zip(
        average.cells,
        (
            "十场景平均",
            f"{overall['official'][0]:.2f} / {overall['official'][1]:.2f}",
            f"{overall['ipr'][0]:.2f} / {overall['ipr'][1]:.2f}",
        ),
    ):
        set_cell(cell, text, bold=True)

    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = width
        tr_pr = row._tr.get_or_add_trPr()
        if tr_pr.find(qn("w:cantSplit")) is None:
            tr_pr.append(OxmlElement("w:cantSplit"))


def main() -> None:
    if not SOURCE.is_file():
        raise FileNotFoundError(SOURCE)
    if not RAW_RESULTS.is_file():
        raise FileNotFoundError(RAW_RESULTS)

    summary, overall, stats = load_results()
    doc = Document(SOURCE)
    math_before = len(doc.element.body.xpath(".//m:oMath"))

    replace_text(
        paragraph_with_prefix(doc, "摘要："),
        "摘要：身份感知语言高斯依赖跨视角对象伪标签构建语义一致性，但自动生成的身份区域可能包含边界噪声、背景泄漏或多对象混合。仅依赖训练交叉熵难以区分可信身份与可被模型拟合的错误标签。本文提出身份空间保持式伪标签可靠性建模方法（Identity-aware Pseudo-label Refinement，IPR），在已有对象ID集合内进行局部清理，利用跨视角裁剪语义响应和连通区域面积分布估计监督可靠性，并通过像素级置信图加权二维身份分类损失。在统一损失归一化与查询解码协议下，三个LERF场景的三种子实验中，IPR将平均mIoU由75.89%提高至79.63%，mBIoU由71.25%提高至74.46%。在固定参数的3D-OVS十场景三种子实验中，平均mIoU由82.01%提高至87.07%，mBIoU由75.69%提高至81.22%，且两个指标的场景均值均在9/10个场景提高；Covered Desk仍出现退化。结果表明，IPR能够在不改变身份集合的条件下缓解部分错误身份监督，但其收益具有场景依赖性。结论限于场景词表已知、无标注多视图可参与预处理的协议。",
    )
    replace_text(
        paragraph_with_prefix(doc, "本文的贡献包括三个方面。"),
        "本文的贡献包括三个方面。第一，明确界定既有身份空间中的监督可靠性问题，并给出不引入新ID的局部处理约束。第二，构建基于语义响应和大范围连通区域风险的置信度加权身份监督，保持ILGS的高斯表示及其他训练分支。第三，在三个LERF场景和3D-OVS十个场景上开展统一协议的三种子受控比较，同时报告配对增益、跨场景波动、退化案例与适用边界。",
    )

    replace_text(
        paragraph_with_prefix(doc, "3D-OVS [19]包含"),
        "3D-OVS [19]包含Bed、Bench、Blue Sofa、Covered Desk、Lawn、Office Desk、Room、Snacks、Sofa和Table十个场景，每个场景使用五个带标注视图评测。本节对ILGS受控基线与IPR分别使用训练种子0、1和2，共形成30组同场景同种子配对比较。两种方法共享相同DEVA轨迹来源、图像划分、语言自编码器和seed_ref查询解码器；IPR仅在原始身份标签基础上进行局部清理与置信度建模。DEVA使用短边480像素的图像副本生成轨迹，标签通过最近邻插值恢复到原图尺寸；语言特征与高斯训练采用原图宽高各缩小4倍的分辨率。",
    )
    replace_text(
        paragraph_with_prefix(doc, "表3  3D-OVS十场景"),
        "表3  3D-OVS十场景三种子受控对照（均值 ± 样本标准差；每格为mIoU / mBIoU，%）",
    )
    update_table(doc.tables[2], summary, overall)

    result = paragraph_with_prefix(doc, "表3显示，IPR的十场景")
    replace_text(
        result,
        "表3显示，IPR的十场景等权平均mIoU为87.07%，高于受控基线的82.01%；mBIoU为81.22%，高于基线的75.69%，分别增加5.06和5.53个百分点。按场景均值比较，两个指标均在9/10个场景提高；Covered Desk的mIoU和mBIoU分别下降2.15和1.99个百分点。Bench、Blue Sofa、Office Desk和Snacks贡献了主要增量，而Bed、Lawn、Sofa和Table的改善较小。因此，整体提升主要来自部分错误对象对应的恢复，不能解释为对所有场景的均匀改善。",
    )
    body_template = paragraph_with_prefix(doc, "3D-OVS [19]包含")
    robustness = insert_after(
        result,
        "在30组同场景同种子配对结果中，IPR分别有24/30次mIoU提高和22/30次mBIoU提高。以十个场景的三种子均值差作为统计单元，mIoU平均增量的95%置信区间为[0.33, 9.80]个百分点，mBIoU为[0.72, 10.34]个百分点。该区间反映十个场景之间的差异，而非把同一场景的三个种子视为完全独立样本。Bed、Blue Sofa和Room等场景仍表现出较大的跨种子标准差，说明训练随机性尚未被完全消除。",
        body_template,
    )
    robustness.paragraph_format.keep_together = True

    replace_text(
        paragraph_with_prefix(doc, "为检查平均数背后的对象级变化"),
        "为检查平均数背后的对象级变化，图2固定展示seed 0中的三个查询及其五视角平均IoU，并选取IoU变化中位视角进行可视化。Bench的Portuguese egg tart由0.007增至0.965，Blue Sofa的sunglasses由0.046增至0.939，Office Desk的a pack of pocket tissues由0.025增至0.945。受控基线分别偏向葡萄、大片背景和书本，而IPR对应到被查询物体。这些案例用于解释可能的身份关联改善，不代替表3的三种子统计，也不能判定某个离线权重分支单独导致了变化。",
    )
    replace_text(
        paragraph_with_prefix(doc, "图2  3D-OVS中的查询目标恢复案例"),
        "图2  3D-OVS中的查询目标恢复案例（seed 0）。每行从左至右为RGB、真实掩码、ILGS受控基线和IPR；所示视角按五个评测视角的IoU变化中位数选取，白色轮廓为真实边界。图中数值为该视角IoU。",
    )
    replace_text(
        paragraph_with_prefix(doc, "图3同时展示失败边界"),
        "图3以seed 0展示失败边界。Covered Desk的gerbera在五个视角中有四个与基线接近，但图像26对应视角的IoU从0.819降至0.078，使该种子下该类五视角平均IoU由0.912降至0.764；三种子场景均值同样下降2.15 mIoU和1.99 mBIoU。Sofa中的Gundam对两种方法都困难，seed 0的五视角平均IoU分别为0.031和0.021。图2的成功案例因此不能概括为普遍的边界改善；对小而复杂的对象以及个别视角的身份判别仍存在明显不足。",
    )
    replace_text(
        paragraph_with_prefix(doc, "图3  3D-OVS的剩余失败案例"),
        "图3  3D-OVS的剩余失败案例（seed 0）。前两行分别为Covered Desk中gerbera的典型视角和退化最大视角，第三行为Sofa中Gundam的IoU变化中位视角；列定义与图2相同。",
    )

    replace_text(
        paragraph_with_prefix(doc, "再次，最终配置是在LERF消融中选择的"),
        "再次，最终配置是在LERF消融中选择，并以固定参数迁移到3D-OVS。三种子结果支持总体平均收益，但场景级95%置信区间仍较宽，且Covered Desk持续退化；Bed、Blue Sofa和Room的跨种子波动也表明，当前证据不能被表述为对任意场景均稳定改善。较大的外部收益主要集中于若干查询的对象对应恢复，现有实验尚不能把因果贡献独立分配给局部清理、语义响应或大区域降权中的某一项。已有SAM2 [14]和DINOv3 [16]相关探索使用过不同标签及评测路径，未纳入本次受控主表。后续需在更多外部数据集上进行多种子验证，并补充受控ID合并、背景泄漏、缺失区域、权重分布以及预处理时间和存储开销分析。",
    )

    replace_text(
        paragraph_with_prefix(doc, "在统一损失归一化和查询解码的三个LERF场景"),
        "在统一损失归一化和查询解码的三个LERF场景三种子实验中，IPR取得79.63% mIoU与74.46% mBIoU，相对受控基线提高3.74和3.21个百分点；在固定参数的3D-OVS十场景三种子实验中，分别取得87.07%与81.22%，提高5.06和5.53个百分点，并在9/10个场景的均值上同时改善两项指标。消融结果表明，局部清理本身不足以保证收益，基础置信度与区域风险降权对总体结果有贡献，语义增量则具有场景依赖性。30组外部配对结果和场景级置信区间支持总体改进，但Covered Desk的退化与若干场景的跨种子波动仍限定了结论范围。当前方法适用于场景词表已知且无标注多视图可参与预处理的设置，进一步验证应面向独立词表、更多数据集和更严格的抗噪机制。",
    )

    math_after = len(doc.element.body.xpath(".//m:oMath"))
    if math_before != math_after:
        raise RuntimeError(f"Word equations changed: {math_before} -> {math_after}")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    check = Document(OUTPUT)
    if len(check.element.body.xpath(".//m:oMath")) != math_before:
        raise RuntimeError("Saved Word equations did not survive the round trip")
    if len(check.inline_shapes) != 2 or len(check.tables[2].rows) != 12:
        raise RuntimeError("Expected two figures and a 12-row external results table")

    print(OUTPUT)
    print(
        f"equations={math_before}, figures={len(check.inline_shapes)}, "
        f"table_rows={len(check.tables[2].rows)}"
    )
    print(
        "3D-OVS delta: "
        f"mIoU={stats['delta_miou']:.2f}, mBIoU={stats['delta_mbiou']:.2f}; "
        f"pair wins={stats['pair_wins_miou']}/30, {stats['pair_wins_mbiou']}/30"
    )


if __name__ == "__main__":
    main()
