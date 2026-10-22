from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.text.paragraph import Paragraph


def paragraph_text(paragraph: Paragraph) -> str:
    return "".join(run.text for run in paragraph.runs)


def find_paragraph(document: Document, needle: str) -> Paragraph:
    for paragraph in document.paragraphs:
        if needle in paragraph_text(paragraph):
            return paragraph
    raise ValueError(f"Paragraph not found: {needle}")


def clear_paragraph(paragraph: Paragraph) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def insert_paragraph_after(paragraph: Paragraph) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    return Paragraph(new_p, paragraph._parent)


def set_run_font(run, size_pt: float = 10.5) -> None:
    run.font.name = "Times New Roman"
    run.font.size = Pt(size_pt)
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:ascii"), "Times New Roman")
    fonts.set(qn("w:hAnsi"), "Times New Roman")
    fonts.set(qn("w:eastAsia"), "宋体")


def replace_text(paragraph: Paragraph, old: str, new: str) -> None:
    text = paragraph_text(paragraph)
    if old not in text:
        raise ValueError(f"Text to replace not found: {old}")
    clear_paragraph(paragraph)
    run = paragraph.add_run(text.replace(old, new))
    set_run_font(run)


def append_text(paragraph: Paragraph, text: str) -> None:
    run = paragraph.add_run(text)
    set_run_font(run)


def set_caption(paragraph: Paragraph, text: str) -> None:
    clear_paragraph(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_together = True
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(text)
    set_run_font(run, 9.0)


def add_picture(paragraph: Paragraph, image_path: Path, width_emu: int, descr: str) -> None:
    clear_paragraph(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_together = True
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run()
    inline = run.add_picture(str(image_path), width=width_emu)
    inline._inline.docPr.set("name", descr)
    inline._inline.docPr.set("descr", descr)


def insert_figure_after(
    anchor: Paragraph,
    image_path: Path,
    width_emu: int,
    descr: str,
    caption_text: str,
    caption_style,
) -> None:
    figure = insert_paragraph_after(anchor)
    add_picture(figure, image_path, width_emu, descr)
    caption = insert_paragraph_after(figure)
    caption.style = caption_style
    set_caption(caption, caption_text)


def revise(source: Path, fig3: Path, fig4: Path, output: Path) -> None:
    document = Document(str(source))
    section = document.sections[0]
    text_width = section.page_width - section.left_margin - section.right_margin
    caption_style = find_paragraph(document, "图1  IPR在ILGS训练流程中的位置").style

    ablation_anchor = find_paragraph(document, "额外加入V后")
    append_text(
        ablation_anchor,
        " 图3将组件级增益与真实中间结果并置：最终采用的无可见性项配置在两项指标上均为最优，"
        "而大面积墙面与地面身份区域被显式标记并在可靠性图中降权。",
    )
    insert_figure_after(
        ablation_anchor,
        fig3,
        int(text_width),
        "LERF ablation and IPR reliability diagnostic",
        "图3  组件消融与可靠性诊断。a，LERF三场景三种子等权平均相对官方ILGS的增益，"
        "浅绿色底纹标出最终采用的无可见性项配置；圆点和方点分别表示mIoU和mBIoU。"
        "b-e，Bench第12帧的RGB、原始身份ID、大区域标记和IPR可靠性图；"
        "墙面ID 6与地面ID 4分别占42.3%和43.3%。a中的数值为描述统计，不表示显著性检验。",
        caption_style,
    )

    paired_anchor = find_paragraph(document, "在30组同场景同种子配对结果中")
    append_text(
        paired_anchor,
        " 图4进一步给出每个场景的三种子均值与样本标准差，显示主要增益来自Blue Sofa、Bench、"
        "Snacks、Office Desk和Room，同时保留Covered Desk退化与部分场景方差较大的事实。",
    )
    insert_figure_after(
        paired_anchor,
        fig4,
        int(text_width),
        "Three-seed robustness across ten 3D-OVS scenes",
        "图4  3D-OVS十场景三种子稳定性比较。a、b分别为mIoU和mBIoU。"
        "点表示三次独立训练的均值，误差线表示样本标准差，连线连接同一场景的官方ILGS与IPR；"
        "右侧数字为IPR减官方ILGS。场景按mIoU增益排序，不作单场景显著性检验。",
        caption_style,
    )

    for paragraph in document.paragraphs:
        text = paragraph_text(paragraph)
        if text.startswith("为检查平均数背后的对象级变化，图3"):
            replace_text(paragraph, "图3", "图5")
        elif text.startswith("图3  3D-OVS中的查询目标恢复案例"):
            replace_text(paragraph, "图3", "图5")
        elif text.startswith("图4以seed 0展示失败边界"):
            replace_text(paragraph, "图4", "图6")
            if "图2的成功案例" in paragraph_text(paragraph):
                replace_text(paragraph, "图2的成功案例", "图5的成功案例")
        elif text.startswith("图4  3D-OVS的剩余失败案例"):
            replace_text(paragraph, "图4", "图6")
            if "图2相同" in paragraph_text(paragraph):
                replace_text(paragraph, "图2相同", "图5相同")

    limitations = find_paragraph(document, "首先，语义可靠性依赖给定场景词表")
    if "图3的成功案例" in paragraph_text(limitations):
        replace_text(limitations, "图3的成功案例", "图5的成功案例")

    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output))
    print(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--fig3", type=Path, required=True)
    parser.add_argument("--fig4", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    revise(args.source, args.fig3, args.fig4, args.output)


if __name__ == "__main__":
    main()
