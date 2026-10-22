"""Add the completed seed-0 3D-OVS comparison to the current IPR draft."""

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(r"D:\PycharmProjects\PythonProject\ILGS_release-main")
SOURCE = ROOT / "outputs/paper_revision_20260929/ILGS-IPR_中文_受控实验与方法修订_20260929.docx"
OUTPUT = ROOT / "outputs/paper_revision_20260930/ILGS-IPR_中文_3DOVS十场景补充_20260930.docx"
FIGURES = Path(r"D:\downloads\seed0_cases_v2")

SCENES = [
    ("Bed", "97.58 / 92.39", "97.59 / 93.05"),
    ("Bench", "81.48 / 76.36", "95.85 / 91.07"),
    ("Blue Sofa", "78.28 / 74.01", "96.51 / 92.42"),
    ("Covered Desk", "79.09 / 69.74", "76.62 / 68.67"),
    ("Lawn", "97.09 / 92.78", "97.12 / 92.88"),
    ("Office Desk", "78.21 / 70.19", "96.03 / 90.43"),
    ("Room", "88.14 / 83.50", "88.21 / 83.10"),
    ("Snacks", "85.31 / 79.89", "85.60 / 80.95"),
    ("Sofa", "73.00 / 66.64", "72.73 / 65.62"),
    ("Table", "81.91 / 72.45", "82.27 / 73.64"),
    ("十场景平均", "84.01 / 77.79", "88.85 / 83.18"),
]


def paragraph_with_prefix(doc, prefix):
    matches = [p for p in doc.paragraphs if p.text.startswith(prefix)]
    if len(matches) != 1:
        raise ValueError(f"Expected one paragraph starting with {prefix!r}, found {len(matches)}")
    return matches[0]


def replace_text(paragraph, text):
    if not paragraph.runs:
        paragraph.add_run(text)
        return
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def insert_after(anchor, text, template):
    element = OxmlElement("w:p")
    if template._p.pPr is not None:
        element.append(deepcopy(template._p.pPr))
    anchor._p.addnext(element)
    from docx.text.paragraph import Paragraph

    paragraph = Paragraph(element, anchor._parent)
    run = paragraph.add_run(text)
    if template.runs and template.runs[0]._r.rPr is not None:
        run._r.insert(0, deepcopy(template.runs[0]._r.rPr))
    return paragraph


def add_figure_after(anchor, image_path, caption, caption_template, body_template):
    figure = insert_after(anchor, "", body_template)
    figure.alignment = WD_ALIGN_PARAGRAPH.CENTER
    figure.paragraph_format.first_line_indent = 0
    figure.paragraph_format.keep_with_next = True
    figure.paragraph_format.space_before = Pt(8)
    figure.paragraph_format.space_after = Pt(2)
    figure.add_run().add_picture(str(image_path), width=Cm(14.55))

    label = insert_after(figure, caption, caption_template)
    label.alignment = WD_ALIGN_PARAGRAPH.CENTER
    label.paragraph_format.first_line_indent = 0
    label.paragraph_format.keep_with_next = False
    label.paragraph_format.space_after = Pt(7)
    return label


def set_cell(cell, text, bold=False, header=False):
    cell.text = text
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for paragraph in cell.paragraphs:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.first_line_indent = 0
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.1
        for run in paragraph.runs:
            run.bold = bold
            run.font.name = "Times New Roman"
            run.font.size = Pt(8.8 if header else 9.0)
            run.font.color.rgb = RGBColor(0, 0, 0)
    if header or bold:
        tc_pr = cell._tc.get_or_add_tcPr()
        shade = tc_pr.find(qn("w:shd"))
        if shade is None:
            shade = OxmlElement("w:shd")
            tc_pr.append(shade)
        shade.set(qn("w:fill"), "E7E7E7" if header else "F2F2F2")


def update_table(table):
    for row in list(table.rows)[1:]:
        table._tbl.remove(row._tr)
    table.autofit = False
    for column, width in zip(table.columns, (Cm(4.0), Cm(5.3), Cm(5.3))):
        column.width = width
    for cell, text in zip(table.rows[0].cells, ("场景", "ILGS受控基线", "IPR")):
        set_cell(cell, text, bold=True, header=True)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for scene, official, ipr in SCENES:
        row = table.add_row()
        row.height = None
        for cell, text in zip(row.cells, (scene, official, ipr)):
            set_cell(cell, text, bold=(scene == "十场景平均"))
    for row in table.rows:
        for cell, width in zip(row.cells, (Cm(4.0), Cm(5.3), Cm(5.3))):
            cell.width = width
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)


def main():
    if not SOURCE.is_file():
        raise FileNotFoundError(SOURCE)
    for name in ("fig_3dovs_recovery.png", "fig_3dovs_limitations.png"):
        if not (FIGURES / name).is_file():
            raise FileNotFoundError(FIGURES / name)

    doc = Document(SOURCE)
    math_before = len(doc.element.body.xpath(".//m:oMath"))

    replace_text(paragraph_with_prefix(doc, "摘要："),
        "摘要：身份感知语言高斯依赖跨视角对象伪标签构建语义一致性，但自动生成的身份区域可能包含边界噪声、背景泄漏或多对象混合。仅依赖训练交叉熵难以区分可信身份与可被模型拟合的错误标签。本文提出身份空间保持式伪标签可靠性建模方法（Identity-aware Pseudo-label Refinement，IPR），在已有对象ID集合内进行局部清理，利用跨视角裁剪语义响应和连通区域面积分布估计监督可靠性，并通过像素级置信图加权二维身份分类损失。在统一损失归一化与查询解码协议下，三个LERF场景的三种子实验中，IPR将平均mIoU由75.89%提高至79.63%，mBIoU由71.25%提高至74.46%。固定参数的3D-OVS十场景单种子实验中，平均mIoU由84.01%提高至88.85%，mBIoU由77.79%提高至83.18%；增益集中于若干原本发生对象身份误匹配的查询，但个别视角出现明显退化。结论限于场景词表已知、无标注多视图可参与预处理的协议；外部数据集的跨种子稳定性仍需验证。")
    replace_text(paragraph_with_prefix(doc, "本文的贡献包括三个方面。"),
        "本文的贡献包括三个方面。第一，明确界定既有身份空间中的监督可靠性问题，并给出不引入新ID的局部处理约束。第二，构建基于语义响应和大范围连通区域风险的置信度加权身份监督，保持ILGS的高斯表示及其他训练分支。第三，在三个LERF场景、三个训练种子和统一查询协议下开展受控比较，并以固定配置在3D-OVS十个场景上进行单种子外部验证，同时报告改善集中处、退化案例与适用边界。")

    replace_text(paragraph_with_prefix(doc, "4.5 3D-OVS"), "4.5 3D-OVS外部验证")
    replace_text(paragraph_with_prefix(doc, "3D-OVS [19]包含"),
        "3D-OVS [19]包含Bed、Bench、Blue Sofa、Covered Desk、Lawn、Office Desk、Room、Snacks、Sofa和Table十个场景，每个场景使用五个带标注视图评测。本节报告训练种子0的完整十场景对照，两种方法共享相同DEVA轨迹来源、图像划分、语言自编码器和seed_ref查询解码器；IPR仅在原始身份标签基础上进行清理与置信度建模。DEVA使用短边480像素的图像副本生成轨迹，标签通过最近邻插值恢复到原图尺寸；语言特征与高斯训练采用原图宽高各缩小4倍的分辨率。")
    replace_text(paragraph_with_prefix(doc, "当前适配对SIMPLE_RADIAL"),
        "当前适配对SIMPLE_RADIAL相机使用焦距对应的针孔近似，未显式校正径向畸变；基线与IPR采用相同近似。外部阶段沿用在LERF消融中确定的配置，不按3D-OVS场景分别调整阈值。该结果属于当前适配协议，不能直接与采用其他相机处理、分辨率或查询解码方式的文献数值排名。")
    replace_text(paragraph_with_prefix(doc, "表3  Bed先导实验"),
        "表3  3D-OVS十场景受控对照（seed 0，mIoU / mBIoU，%）")
    update_table(doc.tables[2])

    result = paragraph_with_prefix(doc, "截至本稿修订时，已得到Bed")
    replace_text(result,
        "表3显示，IPR的十场景等权平均mIoU为88.85%，高于受控基线的84.01%；mBIoU为83.18%，高于基线的77.79%，分别增加4.84和5.39个百分点。逐场景比较中，mIoU在8/10个场景提高，mBIoU在7/10个场景提高。Bench、Blue Sofa和Office Desk贡献了最明显的mIoU增量，分别为14.37、18.23和17.82个百分点；Covered Desk和Sofa的mIoU分别下降2.47和0.27个百分点。因此，平均改进不能解释为对所有场景均有收益。")
    body_template = paragraph_with_prefix(doc, "3D-OVS [19]包含")
    caption_template = paragraph_with_prefix(doc, "表3  3D-OVS十场景")
    case_p = insert_after(result,
        "为检查平均数背后的对象级变化，图2给出三个查询的五视角平均IoU及其中位改善视角。Bench的Portuguese egg tart由0.007增至0.965，Blue Sofa的sunglasses由0.046增至0.939，Office Desk的a pack of pocket tissues由0.025增至0.945。可视化中，受控基线分别偏向葡萄、大片背景和书本，而IPR对应到被查询物体；这与身份关联改善相符，但仅凭终端掩码不能判定某个离线权重分支单独导致了该变化。", body_template)
    fig2_caption = (
        "图2  3D-OVS中的查询目标恢复案例。每行从左至右为RGB、真实掩码、ILGS受控基线和IPR；"
        "所示视角按五个评测视角的IoU变化中位数选取，白色轮廓为真实边界。图中数值为该视角IoU。"
    )
    fig2 = add_figure_after(case_p, FIGURES / "fig_3dovs_recovery.png",
                            fig2_caption, caption_template, body_template)

    limit_p = insert_after(fig2,
        "图3同时展示失败边界。Covered Desk的gerbera在五个视角中有四个与基线接近，但图像26对应视角的IoU从0.819降至0.078，使该类五视角平均IoU由0.912降至0.764。Sofa中的Gundam对两种方法都困难，五视角平均IoU为0.031和0.021。图2的成功案例因此不能概括为普遍的边界改善；对小而复杂的对象以及个别视角的身份判别仍存在明显不足。", body_template)
    fig3_caption = (
        "图3  3D-OVS的剩余失败案例。前两行分别为Covered Desk中gerbera的典型视角和退化最大视角，"
        "第三行为Sofa中Gundam的IoU变化中位视角；列定义与图2相同。"
    )
    add_figure_after(limit_p, FIGURES / "fig_3dovs_limitations.png",
                     fig3_caption, caption_template, body_template)

    replace_text(paragraph_with_prefix(doc, "首先，语义可靠性依赖给定场景词表"),
        "首先，语义可靠性依赖给定场景词表，未涵盖的对象或多个相似候选可能产生较低置信度。候选词表与评测查询重合，以及无标注留出图像可参与预处理，限定了当前结果的适用协议；独立词表与严格训练视图预处理仍需单独验证。其次，面积规则可能误伤真实的大对象，高置信度的错误身份也可能被保留，完全遗漏的对象无法通过降权恢复。3D-OVS中Covered Desk的单视角大幅退化和Sofa的持续失败表明，外部场景中的错误并未被本方法普遍消除。")
    replace_text(paragraph_with_prefix(doc, "再次，最终配置是在LERF消融中选择的"),
        "再次，最终配置是在LERF消融中选择的；3D-OVS虽然固定参数并覆盖十个场景，但目前仅报告一个训练种子，不能据此声称跨种子稳定。较大的外部平均收益主要集中于若干查询的对象对应恢复，现有证据尚不能把因果贡献分配给局部清理、语义响应或大区域降权中的某一项。已有SAM2 [14]和DINOv3 [16]相关探索使用过不同标签及评测路径，未纳入本次受控主表。后续需增加3D-OVS配对多种子实验，并补充受控ID合并、背景泄漏、缺失区域及权重分布分析；预处理时间和存储开销也尚未形成完整统计。")

    replace_text(paragraph_with_prefix(doc, "在统一损失归一化和查询解码的三个LERF场景"),
        "在统一损失归一化和查询解码的三个LERF场景三种子实验中，IPR取得79.63% mIoU与74.46% mBIoU，相对受控基线提高3.74和3.21个百分点；在固定参数的3D-OVS十场景单种子对照中，分别取得88.85%与83.18%，提高4.84和5.39个百分点。消融结果表明，局部清理本身不足以保证收益，基础置信度与区域风险降权对总体结果有贡献；语义增量依赖场景。外部收益集中于若干原本对应错误的查询，且存在个别视角显著退化。当前结论仅适用于所述场景词表和多视图预处理协议，跨种子外部稳定性与直接抗噪机制仍待验证。")
    paragraph_with_prefix(doc, "本文研究ILGS既有身份空间中的伪标签可靠性问题。").paragraph_format.keep_together = True

    math_after = len(doc.element.body.xpath(".//m:oMath"))
    if math_before != math_after:
        raise RuntimeError(f"Word equations changed: {math_before} -> {math_after}")
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    check = Document(OUTPUT)
    if len(check.element.body.xpath(".//m:oMath")) != math_before:
        raise RuntimeError("Saved Word equations did not survive the round trip")
    if len(check.inline_shapes) != 2 or len(check.tables[2].rows) != 12:
        raise RuntimeError("Expected two figures and a 12-row external results table")
    print(OUTPUT)
    print(f"equations={math_before}, figures={len(check.inline_shapes)}, table_rows={len(check.tables[2].rows)}")


if __name__ == "__main__":
    main()
