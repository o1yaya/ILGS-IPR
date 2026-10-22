from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.text.paragraph import Paragraph


def paragraph_text(paragraph: Paragraph) -> str:
    return "".join(run.text for run in paragraph.runs)


def find_paragraph(document: Document, needle: str) -> Paragraph:
    for paragraph in document.paragraphs:
        if needle in paragraph_text(paragraph):
            return paragraph
    raise ValueError(f"Paragraph not found: {needle}")


def clear_paragraph(paragraph: Paragraph) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def remove_paragraph(paragraph: Paragraph) -> None:
    parent = paragraph._element.getparent()
    parent.remove(paragraph._element)
    paragraph._p = paragraph._element = None


def insert_paragraph_after(paragraph: Paragraph) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    return Paragraph(new_p, paragraph._parent)


def set_run_font(run, size_pt: float = 10.5) -> None:
    run.font.name = "Times New Roman"
    run.font.size = Pt(size_pt)
    r_fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    r_fonts.set(qn("w:ascii"), "Times New Roman")
    r_fonts.set(qn("w:hAnsi"), "Times New Roman")
    r_fonts.set(qn("w:eastAsia"), "宋体")


def replace_text(paragraph: Paragraph, old: str, new: str) -> None:
    text = paragraph_text(paragraph)
    if old not in text:
        raise ValueError(f"Text to replace not found: {old}")
    clear_paragraph(paragraph)
    run = paragraph.add_run(text.replace(old, new))
    set_run_font(run)


def set_caption(paragraph: Paragraph, text: str) -> None:
    clear_paragraph(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_together = True
    run = paragraph.add_run(text)
    set_run_font(run, 9.0)


def add_picture(paragraph: Paragraph, image_path: Path, width_emu: int, descr: str) -> None:
    clear_paragraph(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_together = True
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run()
    inline = run.add_picture(str(image_path), width=width_emu)
    inline._inline.docPr.set("name", descr)
    inline._inline.docPr.set("descr", descr)


def revise(source: Path, pipeline: Path, method: Path, output: Path) -> None:
    document = Document(str(source))
    section = document.sections[0]
    text_width = section.page_width - section.left_margin - section.right_margin

    placeholders = [
        find_paragraph(document, "原始对象ID掩码"),
        find_paragraph(document, "对象连通区域"),
        find_paragraph(document, "基础置信图 × 对象可靠性"),
        find_paragraph(document, "ILGS身份渲染"),
        find_paragraph(document, "RGB重建、语言特征监督"),
    ]
    pipeline_caption = find_paragraph(document, "IPR在ILGS训练流程中的位置")
    add_picture(
        placeholders[0],
        pipeline,
        int(text_width),
        "IPR position in the ILGS training pipeline",
    )
    for paragraph in placeholders[1:]:
        remove_paragraph(paragraph)
    set_caption(
        pipeline_caption,
        "图1  IPR在ILGS训练流程中的位置。a，输入RGB图像、含噪身份标签和CLIP特征；"
        "b，IPR在既有身份空间内生成清理标签Y'与可靠性权重W；c，身份感知语言高斯的表示保持不变；"
        "d，仅二维身份分类项由清理标签和像素可靠性加权，语言与RGB分支以及原有动态离群过滤保持不变。",
    )

    method_anchor = find_paragraph(document, "式（10）表示以式（9）替换身份分类项")
    appended = method_anchor.add_run(
        " 图2以同一查询示例概括了可靠性估计、置信加权与下游分割之间的作用路径；"
        "图中定量变化仅用于说明现象，主要结论以第4节的多种子统计为准。"
    )
    set_run_font(appended)
    method_figure = insert_paragraph_after(method_anchor)
    add_picture(
        method_figure,
        method,
        int(text_width),
        "IPR reliability modeling and its effect on ILGS segmentation",
    )
    method_caption = insert_paragraph_after(method_figure)
    method_caption.style = pipeline_caption.style
    set_caption(
        method_caption,
        "图2  IPR的可靠性建模与作用路径。a，原始身份标签中的背景泄漏；b，结构清理、跨视角语义证据和面积先验共同形成像素置信图，"
        "并用于加权身份损失；c，置信加权训练后的身份感知语言高斯在新视角产生语言、身份与分段相似度；"
        "d，同一文本查询下，错误身份区域受到抑制，预测集中到目标对象。定量结论见表3。",
    )

    for paragraph in document.paragraphs:
        text = paragraph_text(paragraph)
        if text.startswith("为检查平均数背后的对象级变化，图2"):
            replace_text(paragraph, "图2", "图3")
        elif text.startswith("图2  3D-OVS中的查询目标恢复案例"):
            replace_text(paragraph, "图2", "图3")
        elif text.startswith("图3以seed 0展示失败边界"):
            replace_text(paragraph, "图3", "图4")
        elif text.startswith("图3  3D-OVS的剩余失败案例"):
            replace_text(paragraph, "图3", "图4")

    limitations = find_paragraph(document, "首先，语义可靠性依赖给定场景词表")
    if "图2的成功案例" in paragraph_text(limitations):
        replace_text(limitations, "图2的成功案例", "图3的成功案例")

    protocol = find_paragraph(document, "主要实验使用LERF中的Figurines")
    clear_paragraph(protocol)
    protocol_run = protocol.add_run(
        "3DGS训练包含随机视角采样、参数初始化和自适应密度控制，单次运行可能将方法差异与优化波动混在一起。"
        "主要实验使用LERF中的Figurines、Ramen和Teatime三个场景。每个设置分别以训练种子0、1、2独立训练，"
        "并保持图像划分、输入分辨率、语言特征及自编码器配置一致，训练均进行30000次迭代。实验使用两张NVIDIA RTX 3090，"
        "每个训练任务占用一张GPU。IPR在LERF消融结果上确定最终配置，再以固定参数开展3D-OVS外部验证；"
        "因此，LERF结果同时承担方法选择与评估功能，不能视为独立于配置选择的测试。"
    )
    set_run_font(protocol_run)

    metrics = find_paragraph(document, "mIoU和mBIoU首先按类别汇总")
    clear_paragraph(metrics)
    metrics_run = metrics.add_run(
        "mIoU和mBIoU首先按类别汇总各评测视图，再对场景内类别等权平均。边界带宽采用图像对角线的2%，最少1像素；"
        "预测尺寸与GT不一致时使用最近邻插值。每个场景报告三种子的均值和样本标准差。由于n=3不足以可靠刻画单场景结果分布，"
        "这些量仅作为运行间波动的描述统计。跨场景平均为场景均值的等权平均，不表示像素加权汇总；"
        "配对胜率按同场景同种子计算，不能替代显著性检验。"
    )
    set_run_font(metrics_run)

    evidence_boundary = find_paragraph(document, "再次，最终配置是在LERF消融中选择")
    replace_text(
        evidence_boundary,
        "三种子结果支持总体平均收益，但场景级95%置信区间仍较宽，且Covered Desk持续退化；",
        "三次重复能够暴露训练波动，但不足以稳定估计单场景方差或建立分布假设，因此本文不据此作单场景显著性声明。"
        "总体平均改善与场景级95%置信区间相符，但区间仍较宽，且Covered Desk持续退化；",
    )

    conclusion = find_paragraph(document, "在统一损失归一化和查询解码的三个LERF场景三种子实验中")
    if "30组外部配对结果和场景级置信区间支持总体改进" in paragraph_text(conclusion):
        replace_text(
            conclusion,
            "30组外部配对结果和场景级置信区间支持总体改进",
            "30组外部配对结果和场景级置信区间与总体改进相符",
        )

    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output))
    print(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--pipeline", type=Path, required=True)
    parser.add_argument("--method", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    revise(args.source, args.pipeline, args.method, args.output)


if __name__ == "__main__":
    main()
