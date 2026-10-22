from html import escape
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from docx.shared import Pt


DOC_DIR = Path(r"D:\wlxy")
INPUT = DOC_DIR / "ILGS_IPR_formula_fixed.docx"
FALLBACK_INPUT = DOC_DIR / "ILGS_IPR_论文初稿.docx"
OUTPUT = DOC_DIR / "ILGS_IPR_formula_omml.docx"
MATH_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"


def r(text):
    return f"<m:r><m:rPr><m:sty m:val=\"p\"/></m:rPr><m:t>{escape(text)}</m:t></m:r>"


def sub(base, lower):
    return f"<m:sSub><m:e>{base}</m:e><m:sub>{lower}</m:sub></m:sSub>"


def sup(base, upper):
    return f"<m:sSup><m:e>{base}</m:e><m:sup>{upper}</m:sup></m:sSup>"


def subsup(base, lower, upper):
    return f"<m:sSubSup><m:e>{base}</m:e><m:sub>{lower}</m:sub><m:sup>{upper}</m:sup></m:sSubSup>"


def frac(num, den):
    return (
        "<m:f><m:fPr><m:type m:val=\"bar\"/></m:fPr>"
        f"<m:num>{num}</m:num><m:den>{den}</m:den></m:f>"
    )


def sqrt(expr):
    return (
        "<m:rad><m:radPr><m:degHide m:val=\"1\"/></m:radPr>"
        f"<m:deg></m:deg><m:e>{expr}</m:e></m:rad>"
    )


def group(expr):
    return f"<m:d><m:dPr><m:begChr m:val=\"(\"/><m:endChr m:val=\")\"/></m:dPr><m:e>{expr}</m:e></m:d>"


def omath(expr):
    return parse_xml(f"<m:oMath xmlns:m=\"{MATH_NS}\">{expr}</m:oMath>")


def S_i():
    return sub(r("S"), r("i"))


def Y_t():
    return sub(r("Y"), r("t"))


def w_t():
    return sub(r("w"), r("t"))


def p_t():
    return sub(r("p"), r("t"))


def A_it():
    return subsup(r("A"), r("i"), r("t"))


def v_it():
    return subsup(r("v"), r("i"), r("t"))


def formula_exprs():
    c = r("c")
    x = r("x")
    k = r("k")
    sum_k = sub(r("∑"), k)
    sum_c = sub(r("∑"), c)
    max_c = sub(r("max"), c)
    c1 = sub(r("c"), r("1"))
    c2 = sub(r("c"), r("2"))
    sigma2 = sup(r("σ"), r("2"))
    A_imed = sub(r("A"), r("i,med"))

    expr1 = (
        Y_t() + group(x) + r(" = i,    ")
        + w_t() + group(x) + r(" ∈ [0, 1]")
    )

    expr2 = (
        S_i() + group(c) + r(" = ")
        + sum_k
        + subsup(r("v"), r("i"), r("k"))
        + subsup(r("p"), r("i"), r("k")) + group(c)
    )

    expr3 = (
        sub(r("P"), r("i")) + r(" = ")
        + frac(max_c + r(" ") + S_i() + group(c), sum_c + r(" ") + S_i() + group(c))
    )

    expr4 = (
        sub(r("M"), r("i")) + r(" = ")
        + frac(
            S_i() + group(c1) + r(" − ") + S_i() + group(c2),
            sum_c + r(" ") + S_i() + group(c),
        )
    )

    expr5_num_inner = sqrt(A_it()) + r(" − ") + sqrt(A_imed)
    expr5_exp = r("−") + frac(sup(group(expr5_num_inner), r("2")), r("2") + sigma2)
    expr5 = v_it() + r(" = ") + A_it() + r(" · exp") + group(expr5_exp)

    expr6_num = sub(r("∑"), x) + r(" ") + w_t() + group(x) + r(" CE") + group(p_t() + group(x) + r(", ") + Y_t() + group(x))
    expr6_den = sub(r("∑"), x) + r(" ") + w_t() + group(x) + r(" + ε")
    expr6 = sub(sup(r("L"), r("′")), r("id")) + r(" = ") + frac(expr6_num, expr6_den)

    return {
        "(1)": expr1,
        "(2)": expr2,
        "(3)": expr3,
        "(4)": expr4,
        "(5)": expr5,
        "(6)": expr6,
    }


def clear_paragraph_runs(paragraph):
    for run in list(paragraph.runs):
        run._element.getparent().remove(run._element)


def set_number_cell(cell, number):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run(number)
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(10.5)


def convert():
    src = INPUT if INPUT.exists() else FALLBACK_INPUT
    doc = Document(str(src))
    exprs = formula_exprs()
    converted = 0
    for table in doc.tables:
        if len(table.rows) != 1 or len(table.columns) != 2:
            continue
        number = table.cell(0, 1).text.strip()
        if number not in exprs:
            continue
        formula_cell = table.cell(0, 0)
        formula_cell.text = ""
        paragraph = formula_cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        clear_paragraph_runs(paragraph)
        paragraph._p.append(omath(exprs[number]))
        set_number_cell(table.cell(0, 1), number)
        converted += 1

    doc.save(str(OUTPUT))
    print(f"input={src}")
    print(f"output={OUTPUT}")
    print(f"converted={converted}")


if __name__ == "__main__":
    convert()
