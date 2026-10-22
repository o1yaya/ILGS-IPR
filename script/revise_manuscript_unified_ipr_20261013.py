"""Revise the retained Chinese manuscript using supplied strict-protocol results."""

from copy import deepcopy
from html import escape
from pathlib import Path
import hashlib
import json
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


SOURCE = Path(r"D:\wlxy\ILGS-IPR_中文1010_标准四场景LERFOVS补充版.docx")
OUTPUT = Path(r"D:\wlxy\ILGS-IPR_中文_统一方法与协议修订_20261018.docx")
QA = Path(r"D:\PycharmProjects\PythonProject\ILGS_release-main\output\manuscript\unified_ipr_20261018")
NS_M = "http://schemas.openxmlformats.org/officeDocument/2006/math"
NS_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def mr(text):
    return '<m:r><w:rPr><w:rFonts w:ascii="Cambria Math" w:hAnsi="Cambria Math"/><w:sz w:val="22"/></w:rPr><m:t xml:space="preserve">' + escape(text) + '</m:t></m:r>'


def sub(base, lower):
    return f'<m:sSub><m:e>{base}</m:e><m:sub>{lower}</m:sub></m:sSub>'


def sup(base, upper):
    return f'<m:sSup><m:e>{base}</m:e><m:sup>{upper}</m:sup></m:sSup>'


def frac(num, den):
    return f'<m:f><m:num>{num}</m:num><m:den>{den}</m:den></m:f>'


def token(base, lower):
    return sub(mr(base), mr(lower))


def cases(rows):
    return '<m:d><m:dPr><m:begChr m:val="{"/><m:endChr m:val=""/></m:dPr><m:e><m:m>' + ''.join('<m:mr>' + ''.join('<m:e>' + x + '</m:e>' for x in row) + '</m:mr>' for row in rows) + '</m:m></m:e></m:d>'


def font_run(run, size=11, bold=False):
    run.font.name = 'Times New Roman'
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '宋体')


def set_text(p, text, size=11):
    p.clear()
    inline = {name: token(base, lower) for name, base, lower in [
        ('T_train','T','train'), ('I_t','I','t'), ('Ω_t','Ω','t'),
        ('Y_t','Y','t'), ('w_t','w','t'), ('N_i','N','i'), ('q_i','q','i'),
        ('u_i','u','i'), ('s_i','s','i'), ('h_t','h','t'), ('ℓ_t','ℓ','t'),
        ('o_t','o','t'), ('E_img','E','img'), ('C_ik','C','ik'),
        ('J_i','J','i'), ('t_ik','t','ik'), ('R_ik','R','ik'),
    ]}
    inline['c_t⁰'] = sup(token('c', 't'), mr('0'))
    inline['Ỹ_t'] = '<m:acc><m:accPr><m:chr m:val="̃"/></m:accPr><m:e>' + token('Y', 't') + '</m:e></m:acc>'
    pattern = r'(?<![A-Za-z0-9_])(' + '|'.join(re.escape(x) for x in sorted(inline,key=len,reverse=True)) + r')(?![A-Za-z0-9_])'
    for part in re.split(pattern, text):
        if part in inline:
            p._p.append(parse_xml(f'<m:oMath xmlns:m="{NS_M}" xmlns:w="{NS_W}">{inline[part]}</m:oMath>'))
        elif part:
            font_run(p.add_run(part), size)
    p.paragraph_format.keep_together = False
    p.paragraph_format.widow_control = True
    return p


def equation(p, expr, number):
    p.clear()
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.keep_together = True
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.tab_stops.clear_all()
    p.paragraph_format.tab_stops.add_tab_stop(Inches(2.884), WD_TAB_ALIGNMENT.CENTER)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(5.768), WD_TAB_ALIGNMENT.RIGHT)
    p.add_run('\t')
    p._p.append(parse_xml(f'<m:oMath xmlns:m="{NS_M}" xmlns:w="{NS_W}">{expr}</m:oMath>'))
    font_run(p.add_run('\t(' + str(number) + ')'))


def build():
    source_hash = hashlib.sha256(SOURCE.read_bytes()).hexdigest()
    d = Document(SOURCE)
    originals = [deepcopy(p._p) for p in d.paragraphs]
    old_tables = [deepcopy(t._tbl) for t in d.tables]
    ps = d.paragraphs
    replacements = {
        1: '摘要：身份感知语言高斯依赖跨视角对象伪标签组织语义监督，但自动生成的身份区域可能含有边界噪声、背景泄漏或多对象混合。本文提出身份空间保持式伪标签可靠性建模方法（Identity-aware Pseudo-label Refinement，IPR），在既有ID集合内进行局部清理，仅基于训练视角中的对象裁剪一致性及连通区域面积风险构造像素置信图，并重加权ILGS的二维身份分类损失。可靠性估计不读取评测查询文本或真实像素标注，保留原有高斯表示与推理流程。在LERF-Mask三个场景的匹配三种子受控实验中，平均mIoU由70.77%提高至74.46%，mBIoU由66.35%提高至69.91%。在固定一致性阈值0.75的3D-OVS十场景三种子受控实验中，平均mIoU由82.01%提高至86.78%，mBIoU由75.69%提高至80.93%，两项指标均在9个场景改善。组件对照进一步显示，一致性与大范围风险门控在3D-OVS上带来明显额外收益，但在LERF-Mask上与基础置信的组合收益并非单调叠加。结果支持该方法在当前受控设置下提高平均分割质量，同时表明收益与运行波动仍具有场景依赖性。',
        7: '本文提出IPR，将局部标签清理与查询无关的可靠性估计连接到ILGS身份分类分支。局部清理仅复用已有ID；随后，对同一ID的训练视角裁剪计算CLIP图像嵌入的一致性，并以大范围连通区域比例作为补充风险代理。对象权重与基础置信图相乘后用于身份监督。本文统一采用这一配置，不使用候选文本词表、文本类别投票或额外的逐视角连续可见性乘子。',
        8: '本文的贡献包括三个方面。第一，在既有身份分类空间内定义伪标签可靠性干预，区分局部标签修正与监督强度调整；第二，结合查询无关的同身份裁剪一致性、局部结构置信度及大范围区域风险，对ILGS二维身份监督进行加权；第三，通过LERF-Mask与3D-OVS三种子比较及补充的标准LERF-OVS评测，报告方法收益、运行波动与失败边界。多种子复验用于评价可靠性，不作为独立算法贡献。',
        15: '与上述工作相比，本文在已有对象ID空间内估计离线监督权重，不引入新的全局ID匹配或合并。语义证据来自同一ID裁剪的CLIP图像嵌入一致性，不要求预先给出目标类别名称。训练误差过滤与离线可靠性建模针对不同信号：前者识别当前难以拟合的像素，后者利用结构、外观一致性和面积风险调节监督。二者均不能保证识别所有错误标签。',
        17: 'IPR以训练视角RGB图像及其既有身份掩码为输入，输出清理后的标签和像素置信图，整体位置如图1。相机与原有语言特征继续用于ILGS训练。可靠性预处理仅读取训练视角子集，不读取评测查询或真实像素标注；文本查询只在推理时用于目标检索。',
        19: '图1  IPR在ILGS中的干预位置。a，输入RGB、身份标签和CLIP特征；b，局部清理、同身份裁剪一致性和面积风险形成清理标签与可靠性权重；c，身份感知语言高斯表示；d，置信权重仅直接作用于二维身份分类项，RGB、语言监督、三维一致性及动态离群过滤沿用共同代码。图中的语义可靠性在本文中统一指图像裁剪一致性。',
        20: '本文将“训练视角限定”具体定义为：清理统计、裁剪选择、CLIP一致性和面积风险估计均只遍历训练子集中的掩码。该定义约束IPR新增预处理；既有跟踪ID及共享语言特征的上游生成过程仍需单独记录，不能仅凭对子集的筛选推断整个数据准备链已隔离留出视角。',
        22: '设训练视角索引集合为T_train，第t幅图像为I_t，像素域为Ω_t，原始身份标签为Y_t，清理后为Ỹ_t，浮点置信图为w_t。局部清理只使用原始图像中的已有ID，后续可靠性估计不再改变清理标签，因此满足：',
        27: '基础置信图c_t⁰由三类局部因素构成：小连通分量乘以0.25，发生重分配时再乘以0.5；身份边界及其扩展邻域乘以0.65；若清理后某ID的总像素面积低于该ID在原始训练视角中出现面积中位数的0.15倍，或高于其6倍，则该ID区域再乘以0.5。边界由3×3局部最大值与最小值不一致的位置检测，并向外膨胀一次。各因素相乘后截断至[0.05, 1]。',
        29: '3.3 查询无关的同身份裁剪一致性',
        30: '对清理后的每个ID i，在训练视角内收集面积不小于500像素的连通分量，第k个分量的像素集合记为R_ik。按像素面积降序选取至多40个裁剪，记数量为N_i。每个裁剪向外扩展12像素，并以对象像素的均值颜色填充掩码外背景。实现按连通分量枚举裁剪，同一视角可以贡献多个分量，因此下述分数同时包含跨视角与视角内分量间的一致性。小于面积门槛的分量只保留基础置信度。',
        31: '使用冻结的OpenCLIP ViT-B/16图像编码器E_img提取裁剪C_ik的特征，并进行L2归一化：',
        33: '当N_i≥2时，以所有不同裁剪对的余弦相似度均值计算原始一致性u_i；每个裁剪对等权，不使用面积投票：',
        36: '将u_i截断到[0, 1]得到q_i。只有一个合格裁剪时，实现令q_i=1，表示不因缺少成对证据施加一致性降权，并不表示该身份被验证为正确。没有合格裁剪的ID不进入该分支，其像素保留基础置信度。',
        37: '3.4 训练视角约束与评分边界',
        38: '设J_i为最终选中的分量索引集，t_ik表示分量k的来源视角。裁剪集合必须满足：',
        40: '对象出现的所有合格训练分量用于面积风险统计，式（5）选中的子集才用于CLIP编码。q_i是同身份外观一致性的启发式量，而非校准的标签正确概率：跨视角外观变化可能降低正确对象的一致性；持续混合的背景或错误身份也可能具有较高一致性。因此该分数与局部置信度及面积风险共同使用。本文不使用文本softmax、类别间隔或面积加权语义投票；这些属于附录B中的历史词表版本。',
        42: '大面积背景或多个对象可能被归入同一ID。本文使用合格连通分量相对整幅图像的面积作为风险代理，而不把它视为真实混合区域的判定器。统计覆盖该ID在训练子集中的全部合格分量，不受40个裁剪的上限限制：',
        44: '其中P90为分量面积比例的90%分位数，1[·]为指示函数。实现还包含中位面积比例≥0.55的条件，该条件蕴含最大比例≥0.45，故式（6）给出等价规则。触发大范围风险的ID跳过CLIP裁剪编码，直接赋予较低权重；其他ID根据q_i选择权重：',
        46: '大范围区域优先取0.35；其余ID在q_i<0.75时取0.65，否则取1。两个附加风险权重按互斥优先级选择，随后才与基础置信度相乘。真实大对象也可能被面积规则降权，这构成方法边界。',
        48: '令h_t(x)在像素属于已收集的合格分量时取对应对象权重s_i，其余位置取1。最终浮点权重为：',
        50: '令ℓ_t(x)为渲染身份分布与清理标签之间的逐像素交叉熵。第n次迭代的动态阈值为当前图像平均交叉熵的(2−n/30000)倍；o_t(x)在交叉熵不高于阈值时为1，否则为0。令C为身份分类器的输出通道数，加权身份损失为：',
        52: '分母按有效置信质量归一化，下限为1；随后除以log C。无置信图对照采用相同代码，令w_t(x)=1，并使用原始身份标签。由于这一归一化与过滤后直接对全图像素取均值的实现不同，本文统一称该对照为“ILGS受控基线”；日志中的official或strict_official是分支名，不代表与原论文逐项一致。',
        54: '式（10）仅替换二维身份分类项。RGB重建、语言监督和三维身份感知语义一致性保持共同实现及其迭代调度；特别地，原代码在指定三维一致性迭代与其他迭代使用不同损失组合，本文不把它们改写为每一步均同时启用的固定求和式。置信图对身份与语言表示的影响通过训练间接产生，推理阶段不额外读取IPR权重。',
    }
    for i, text in replacements.items():
        set_text(ps[i], text)
    # Preserve existing editable equations where their meaning remains unchanged.
    fi = token('f', 'ik')
    ci = token('C', 'ik')
    ni = token('N', 'i')
    ui = token('u', 'i')
    qi = token('q', 'i')
    enc = token('E', 'img') + mr('(') + ci + mr(')')
    equation(ps[32], fi + mr(' = ') + frac(enc, mr('‖') + enc + token('‖', '2')), 2)
    equation(ps[34], ui + mr(' = ') + frac(mr('2'), ni + mr('(') + ni + mr('−1)')) + sub(mr('∑'), mr('1≤k<l≤') + ni) + mr('⟨') + fi + mr(', ') + token('f', 'il') + mr('⟩'), 3)
    equation(ps[35], qi + mr(' = ') + cases([(mr('clip(') + ui + mr(', 0, 1)'), ni + mr('≥2')), (mr('1'), ni + mr('=1'))]), 4)
    equation(ps[39], token('t', 'ik') + mr('∈') + token('T', 'train') + mr('  (k∈') + token('J', 'i') + mr('),    ') + ni + mr('=|') + token('J', 'i') + mr('|≤40'), 5)
    equation(ps[45], token('s', 'i') + mr(' = ') + cases([(mr('0.35'), token('g', 'i') + mr('=1')), (mr('0.65'), token('g', 'i') + mr('=0, ') + qi + mr('<0.75')), (mr('1'), mr('otherwise'))]), 7)

    modified = [deepcopy(p._p) for p in ps]
    body = d._element.body
    for child in list(body):
        if child.tag != qn('w:sectPr'):
            body.remove(child)

    def append_node(node):
        body.insert(len(body) - 1, deepcopy(node))

    def para(text='', style=None):
        p = d.add_paragraph(style=style)
        set_text(p, text)
        fmt = p.paragraph_format
        fmt.space_after = Pt(6)
        fmt.line_spacing = 1.15
        if style and style.startswith('Heading'):
            fmt.keep_with_next = True
            fmt.space_before = Pt(12)
            fmt.first_line_indent = Pt(0)
            for run in p.runs:
                font_run(run, 13 if style == 'Heading 1' else 11.5, True)
        else:
            fmt.first_line_indent = Pt(22)
        return p

    def caption(text):
        p = para(text)
        p.paragraph_format.first_line_indent = Pt(0)
        p.paragraph_format.keep_together = True
        for r in p.runs:
            font_run(r, 9.5)
        return p

    def table(title, rows, widths=None, size=9):
        cap = caption(title)
        cap.paragraph_format.keep_with_next = True
        t = d.add_table(rows=0, cols=len(rows[0]))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        if widths is None:
            widths = {
                3: [1.00, 2.38, 2.38],
                4: [1.65, 2.25, 0.94, 0.94],
                5: [1.42, 1.09, 1.09, 1.09, 1.09],
            }[len(rows[0])]
        for c,w in zip(t.columns,widths):
            c.width = Inches(w)
        for j, row in enumerate(rows):
            cells = t.add_row().cells
            for c,w,text in zip(cells,widths,row):
                c.width = Inches(w)
                c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                p = set_text(c.paragraphs[0], text, size)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.first_line_indent = Pt(0)
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(4)
                p.paragraph_format.line_spacing = 1.05
                for r in p.runs:
                    r.bold = j == 0
                pr = c._tc.get_or_add_tcPr()
                sh = OxmlElement('w:shd')
                sh.set(qn('w:fill'), 'EDEDED' if j == 0 else 'FFFFFF')
                pr.append(sh)
            trpr = t.rows[-1]._tr.get_or_add_trPr()
            trpr.append(OxmlElement('w:cantSplit'))
            if j == 0:
                trpr.append(OxmlElement('w:tblHeader'))
        borders = OxmlElement('w:tblBorders')
        for edge in ['top','left','bottom','right','insideH','insideV']:
            e = OxmlElement('w:' + edge)
            e.set(qn('w:val'), 'single')
            e.set(qn('w:sz'), '4')
            e.set(qn('w:color'), 'D9D9D9')
            borders.append(e)
        t._tbl.tblPr.append(borders)
        return t

    # Keep the author's title, introduction, related work, pipeline and method.
    for node in modified[:55]:
        append_node(node)

    para('4 实验与分析', 'Heading 1')
    para('4.1 数据集与比较范围', 'Heading 2')
    para('主文以LERF-Mask和3D-OVS为两套评测。LERF-Mask包含Ramen、Figurines和Teatime；3D-OVS包含Bed、Bench、Blue Sofa、Covered Desk、Lawn、Office Desk、Room、Snacks、Sofa和Table。标准四场景LERF-OVS作为补充协议单列于附录A，不与LERF-Mask混合排名或平均。其场景与LERF-Mask部分重叠，不能视为四个完全独立的新场景。')
    para('本文所称IPR均指第3节的训练视角裁剪一致性配置，实验标识为strict_trainview_consistency，一致性阈值固定为0.75。ILGS受控基线使用原始身份掩码，不加载IPR置信图；两组共享训练代码、图像划分、分辨率、语言特征、自编码器及每个基准内部的解码设置。候选词表参与权重构造的旧版本仅在附录B归档，不替代本方法的完整实验或组件消融。')
    para('4.2 实现与评测协议', 'Heading 2')
    para('训练共30000次迭代，使用两张NVIDIA RTX 3090，每个任务使用一张GPU。LERF严格版采用原图分辨率，3D-OVS采用宽高各缩小4倍的训练分辨率。IPR处理输入掩码尺寸下的连通分量；最小面积500像素、40个裁剪上限、外扩12像素、低一致性权重0.65、大范围权重0.35及置信下限0.05均在不同场景间保持一致。语义编码使用OpenCLIP ViT-B/16；阈值0.75必须显式记录，不能用脚本的其他默认值替代已运行配置。')
    para('LERF-Mask使用既有images_train划分，三个场景分别保留4、4、3幅测试图像；该划分不是每8帧留出一次的LLFF划分。3D-OVS沿用场景的训练与测试划分。当前3D-OVS适配将SIMPLE_RADIAL相机按针孔近似处理；此近似、输入分辨率及损失归一化共同限定本文受控比较的范围，不能仅凭相同数据集名称与其他文献数值直接排名。')
    para('LERF-Mask与3D-OVS报告mIoU和mBIoU：先按查询类别汇总评测视图，再对场景内类别等权平均，最后对场景均值等权平均。边界带宽取图像对角线的2%，至少1像素，尺寸不一致时以最近邻插值对齐预测。LERF严格评测脚本默认采用official解码，3D-OVS当前脚本固定采用seed_ref解码；各基准内部必须对两种方法使用同一模式，解码器不是IPR贡献。具体运行模式须以原始命令和日志为准，不能仅由汇总表名推断。')
    para('LERF-Mask与3D-OVS均使用种子0、1、2，表1和表2报告均值与样本标准差。多种子用于描述运行波动，不作单场景显著性推断；LERF-Mask的九次训练和3D-OVS的三十次训练也不能分别视为九个或三十个独立场景。配置选择曾参考LERF实验，故LERF结果不能解释为独立于方法开发的盲测。')
    para('4.3 LERF-Mask三种子结果', 'Heading 2')
    table('表1  LERF-Mask严格版三种子受控比较，均值±样本标准差（%；每格为mIoU / mBIoU）', [
        ['场景','ILGS受控基线','ILGS + IPR'],
        ['Ramen','74.03 ± 0.69 / 66.50 ± 0.25','79.60 ± 2.73 / 71.68 ± 2.23'],
        ['Figurines','67.84 ± 11.64 / 65.48 ± 11.42','72.14 ± 2.02 / 69.84 ± 2.06'],
        ['Teatime','70.42 ± 4.12 / 67.09 ± 3.78','71.65 ± 5.15 / 68.21 ± 5.31'],
        ['场景平均','70.77 / 66.35','74.46 / 69.91'],
    ])
    para('IPR的场景平均mIoU提高3.69个百分点，mBIoU提高3.56个百分点。Ramen、Figurines和Teatime的mIoU分别提高5.57、4.30和1.23个百分点，mBIoU分别提高5.18、4.36和1.12个百分点。三个场景的均值均改善，但改善幅度并不一致。')
    para('Figurines的运行间波动较基线减小，而Ramen与Teatime的两项标准差均增大。因此，当前数据支持平均分割质量改善，不支持“普遍提高训练稳定性”的结论。仅凭均值和标准差也无法恢复逐种子配对胜率；旧版本的8/9等配对统计不移用于本表。')
    para('4.4 3D-OVS十场景三种子结果', 'Heading 2')
    para('为检验统一IPR在完整基准上的行为，本文对3D-OVS十个场景进行种子0、1、2的受控比较。两种方法采用相同的训练划分、分辨率、seed_ref解码和评测代码；目标差异是IPR产生的清理标签与置信图。表2先在每个场景内对三个种子取均值并计算样本标准差，再对十个场景的均值等权平均。')
    table('表2  3D-OVS十场景三种子受控比较，均值±样本标准差（%；每格为mIoU / mBIoU）', [
        ['场景','ILGS受控基线','ILGS + IPR'],
        ['Bed','92.11 ± 9.36 / 86.34 ± 10.65','97.63 ± 0.01 / 93.27 ± 0.05'],
        ['Bench','81.36 ± 0.30 / 76.27 ± 0.33','95.82 ± 0.00 / 90.98 ± 0.07'],
        ['Blue Sofa','72.11 ± 8.15 / 66.97 ± 9.02','84.11 ± 11.22 / 79.68 ± 12.76'],
        ['Covered Desk','78.90 ± 1.59 / 69.76 ± 1.18','79.20 ± 2.72 / 70.70 ± 2.55'],
        ['Lawn','97.09 ± 0.04 / 92.78 ± 0.29','97.19 ± 0.04 / 93.10 ± 0.38'],
        ['Office Desk','90.15 ± 10.35 / 83.82 ± 11.80','96.04 ± 0.02 / 90.32 ± 0.05'],
        ['Room','77.26 ± 9.42 / 72.03 ± 9.95','88.39 ± 0.41 / 83.73 ± 0.44'],
        ['Snacks','77.47 ± 7.75 / 71.59 ± 8.12','74.80 ± 9.43 / 68.84 ± 10.46'],
        ['Sofa','71.69 ± 1.74 / 64.94 ± 1.98','72.56 ± 0.38 / 65.59 ± 0.56'],
        ['Table','81.98 ± 0.09 / 72.44 ± 0.09','82.06 ± 0.09 / 73.13 ± 0.04'],
        ['场景平均','82.01 / 75.69','86.78 / 80.93'],
    ])
    para('十场景平均mIoU提高4.77个百分点，mBIoU提高5.24个百分点。Bench、Blue Sofa和Room的mIoU分别提高14.46、12.00和11.13个百分点，mBIoU分别提高14.71、12.71和11.70个百分点；Office Desk与Bed也取得5.89 / 6.50和5.52 / 6.93个百分点的提升。Covered Desk、Lawn、Sofa和Table的两项指标小幅改善，Snacks的mIoU和mBIoU分别下降2.67和2.75个百分点。')
    para('两项指标均在9/10个场景改善，但方差变化并不一致。IPR在Bed、Bench、Office Desk、Room和Sofa上同时降低两项指标的运行间标准差，在Blue Sofa、Covered Desk和Snacks上则同时增大；Lawn与Table表现为持平或单项变化。因此，结果支持统一IPR提高三种子平均分割质量，但不支持其普遍降低训练波动。该比较也不能证明上游数据链不存在留出视角影响。')
    para('附录B的旧词表协议在相同十场景三种子聚合下达到87.07 / 81.22，统一查询无关配置为86.78 / 80.93，两项均相差0.29个百分点。该差异说明移除评测词表与留出视角参与后仍保留了大部分平均收益，但新旧版本同时改变了视角范围和可靠性定义，不能将0.29个百分点解释为单一组件的因果代价。')
    para('4.5 组件对照与适用边界', 'Heading 2')
    para('局部清理、基础置信度、CLIP裁剪一致性与大范围区域降权共同组成当前IPR。为避免将旧词表协议的消融移植到新方法，本节仅使用训练视角、查询无关的统一实现。strict_trainview_cleanup仅替换清理后标签；cleanup_conf在此基础上加载局部结构置信图；consistency_only和large_only分别在共享的清理与基础置信上只启用裁剪一致性或大范围风险分支；完整IPR同时启用两个附加分支。')
    table('表3  3D-OVS十场景三种子机制对照（%）', [
        ['设置','新增可靠性机制','mIoU','mBIoU'],
        ['ILGS受控基线','无IPR','82.01','75.69'],
        ['局部清理+基础置信','清理标签+c_t⁰','83.00','76.91'],
        ['完整IPR','再加一致性与面积风险','86.78','80.93'],
    ])
    para('在3D-OVS上，局部清理与基础置信相对基线提高0.99 mIoU和1.22 mBIoU个百分点。进一步启用裁剪一致性与大范围风险后，完整IPR相对前者再提高3.78和4.02个百分点，且两项指标都在9/10个场景更高。唯一例外是Snacks，完整IPR相对清理加基础置信分别下降4.89和5.02个百分点。因此，两个高层风险分支在该基准上提供了主要额外增益，但仍存在场景依赖的负效应。')
    table('表4  LERF-Mask统一协议组件对照，三种子均值±样本标准差（%；每格为mIoU / mBIoU）', [
        ['设置','Ramen','Figurines','Teatime','三场景均值'],
        ['ILGS受控基线','74.03 ± 0.69 / 66.50 ± 0.25','67.84 ± 11.64 / 65.48 ± 11.42','70.42 ± 4.12 / 67.09 ± 3.78','70.77 / 66.35'],
        ['局部清理','74.79 ± 3.05 / 67.46 ± 3.02','66.93 ± 0.78 / 64.89 ± 1.18','71.16 ± 3.21 / 68.34 ± 3.81','70.96 / 66.90'],
        ['清理+基础置信','75.28 ± 2.83 / 68.04 ± 2.83','74.32 ± 5.42 / 72.11 ± 4.91','74.12 ± 2.46 / 71.22 ± 2.30','74.57 / 70.46'],
        ['仅一致性分支','75.24 ± 3.31 / 67.41 ± 3.58','73.40 ± 0.71 / 71.31 ± 0.61','71.95 ± 6.19 / 68.75 ± 6.35','73.53 / 69.15'],
        ['仅大范围分支','77.78 ± 3.17 / 70.01 ± 2.83','73.20 ± 3.43 / 70.25 ± 3.10','72.52 ± 4.23 / 69.03 ± 5.09','74.50 / 69.76'],
        ['完整IPR','79.60 ± 2.73 / 71.68 ± 2.23','72.14 ± 2.02 / 69.84 ± 2.06','71.65 ± 5.15 / 68.21 ± 5.31','74.46 / 69.91'],
    ], size=8)
    para('在LERF-Mask上，局部清理单独仅带来0.19 mIoU和0.55 mBIoU个百分点的平均增益，加入基础置信后增益扩大到3.80和4.11个百分点。完整IPR仍比基线高3.69和3.56个百分点，但比“清理+基础置信”低0.11和0.55个百分点；仅一致性或仅大范围分支的场景均值也未同时超过基础置信。这表明不同风险信号的作用不是简单叠加：基础置信是本组场景的主要改善来源，而额外门控的收益随场景和数据特性变化。')
    para('外观一致不等于身份正确。持续混合的区域可能得到高一致性分数，真实对象因遮挡、反光或视角变化也可能被降权。面积规则可能误伤大对象；仅有一个裁剪的身份缺少成对一致性证据；原始跟踪遗漏的对象无法通过调整权重恢复。像素面积门槛还依赖预处理分辨率。')
    para('训练视角约束目前明确作用于IPR新增预处理。既有DEVA轨迹是否由全序列生成、语言特征及自编码器是否使用留出图像，需要与场景划分一起审计。标准LERF-OVS中已标注帧是否参与训练也应逐帧说明。本文不把标注帧可被渲染等同于该帧必然未见过，也不将补充协议的平均提升解释为已证明的新视角泛化能力。')
    para('补充的标准LERF-OVS四场景平均mIoU提高1.51个百分点、mAcc@0.25提高1.82个百分点，但主要增量来自Waldo Kitchen；Ramen与Figurines的mIoU略有下降。该结果与主文共同说明，IPR收益具有场景和评测协议依赖性。现有数据尚未量化新增预处理时间、显存及存储成本，也不支持对这些开销作零成本声明。')

    para('5 结论', 'Heading 1')
    para('本文将IPR统一为训练视角内、查询无关的身份伪标签可靠性模块。方法在既有ID集合内进行局部清理，以同身份CLIP裁剪一致性和大范围区域风险形成像素置信图，并直接作用于ILGS二维身份分类损失。其作用是调节监督强度，不能保证修复所有错误身份。')
    para('在LERF-Mask三场景三种子受控比较中，平均mIoU和mBIoU分别提高3.69和3.56个百分点；在3D-OVS十场景三种子比较中，两项指标分别提高4.77和5.24个百分点，并均在9个场景改善。3D-OVS机制对照显示，一致性与大范围风险门控在清理加基础置信之上进一步提高3.78 mIoU和4.02 mBIoU个百分点；但LERF-Mask组件对照中，完整IPR未超过清理加基础置信。标准LERF-OVS补充评测也只给出有限且非均匀的收益。现有证据支持统一IPR提高当前受控设置下的平均分割质量，但Snacks仍出现退化，若干场景的运行间波动增大，各可靠性信号的增益也不是普遍单调叠加，因此不能声称逐场景稳定改善或普遍适用。')

    para('参考文献', 'Heading 1').paragraph_format.page_break_before = True
    for node in originals[108:]:
        append_node(node)

    h = para('附录A 标准四场景LERF-OVS评测', 'Heading 1')
    h.paragraph_format.page_break_before = True
    para('本附录使用Ramen、Figurines、Teatime和Waldo Kitchen的LERF-OVS标注，以帧–查询对为聚合单元。mIoU为每个场景内帧–查询对IoU的均值；mAcc@0.25为IoU严格大于0.25的配对比例，不是像素分类准确率。各场景结果先对种子0、1、2取均值，再计算四场景等权平均。标注与聚合不同于主文LERF-Mask，不能直接比较两表的数值高低。')
    para('IPR采用第3节的同一查询无关权重定义，阈值0.75。结果来自standard_officialdecode评测；每个场景内两组使用相同解码和训练设置。表中基线按当前受控实现命名，不作为ILGS原论文数值。完整的训练帧与标注帧关系、基线代码版本及逐运行配置仍应随可复现实验记录一同提供。')
    table('表A1  标准LERF-OVS三种子结果，均值±样本标准差（%；每格为mIoU / mAcc@0.25）', [
        ['场景','ILGS受控基线','ILGS + IPR'],
        ['Ramen','51.32 ± 0.52 / 57.28 ± 0.81','51.19 ± 0.64 / 56.81 ± 0.81'],
        ['Figurines','53.97 ± 1.84 / 59.52 ± 2.73','53.77 ± 0.36 / 60.12 ± 1.03'],
        ['Teatime','66.70 ± 2.01 / 72.88 ± 2.94','67.24 ± 1.84 / 74.01 ± 1.96'],
        ['Waldo Kitchen','30.04 ± 0.91 / 43.94 ± 2.62','35.88 ± 1.20 / 50.00 ± 0.00'],
        ['场景平均','50.51 / 58.41','52.02 / 60.23'],
    ])
    para('四场景mIoU和mAcc@0.25分别增加1.51和1.82个百分点。Waldo Kitchen分别增加5.84和6.06个百分点；Ramen与Figurines的mIoU分别下降0.13和0.20个百分点，Teatime提高0.54个百分点。因此，总体收益不能解释为每个场景均改善，也不能仅凭较大提升反推某个场景的噪声程度更高。Waldo Kitchen的准确率标准差0.00仅表示三个种子得到相同的阈值成功比例，不代表预测掩码完全相同。')

    h = para('附录B 旧词表协议的结果与定性归档', 'Heading 1')
    h.paragraph_format.page_break_before = True
    para('本附录保留已有结果以说明方法开发过程与失败边界。这里的“旧词表版IPR”使用评测查询词表进行语义评分，并允许无标注留出视角参与预处理；它与第3节的查询无关配置不同。以下表格和图片不作为统一方法的最终结果，不与表1、表2或表A1混合平均。原文件中的no_visibility也只对应本附录的旧配置。')
    para('B.1 旧可靠性定义与LERF结果', 'Heading 2')
    para('旧版本对裁剪与候选文本计算缩放系数35的余弦softmax，以面积权重汇总类别响应，使用最高类别归一化得分0.45和前两类间隔0.08作为降权阈值。面积投票尺度为100、最低投票权重0.25；低语义和大区域权重分别为0.65和0.35。这套文本评分及其消融不能视为当前图像一致性分支的同义实现。')
    caption('表B1  旧词表协议下的LERF三种子结果（%；每格为mIoU / mBIoU）').paragraph_format.keep_with_next = True
    append_node(old_tables[0])
    para('在旧协议下，场景均值由75.89 / 71.25提高至79.63 / 74.46。该差异与主文表1的统计对象及配置不同，不能用数值高低判断新旧方法优劣；尤其不能以旧版较高绝对值替换严格版结果。')
    caption('表B2  旧词表协议的组件消融（三场景三种子平均，%）').paragraph_format.keep_with_next = True
    append_node(old_tables[2])
    para('表B2中S为文本语义降权，G为大范围降权，V为附加连续可见性乘子；各设置均在已有基础置信图上比较。“基础置信度＋S＋G”是旧版本选中的配置。大区域若已被跳过CLIP编码，关闭G也不会补算文本语义，因此该消融还受门控实现影响。它支持旧协议下的组合比较，不能独立证明新q_i分数有效。')
    append_node(originals[84])
    caption('图B1  旧词表协议的组件消融与可靠性示例。a为旧LERF三场景三种子统计；b-e为Bench第12帧的RGB、原始ID、大区域标记及旧版权重。墙面ID 6与地面ID 4分别占42.3%和43.3%。图内“官方”或“最终配置”沿用旧图标注，应按本附录的受控基线与旧词表版本解释。')

    para('B.2 旧3D-OVS十场景三种子结果', 'Heading 2')
    caption('表B3  旧词表协议下的3D-OVS结果，均值±样本标准差（%；每格为mIoU / mBIoU）').paragraph_format.keep_with_next = True
    append_node(old_tables[3])
    para('旧协议下，十场景平均提高5.06 mIoU和5.53 mBIoU个百分点，两个指标的场景均值均在9/10个场景改善；Covered Desk分别下降2.15和1.99个百分点。这些统计仅归属于旧版本。主文表2已给出统一配置的十场景三种子结果；两表使用相同种子聚合，但预处理视角范围和可靠性定义同时变化，因此新旧差异只能作为整体版本比较，不能作为单因素消融。')
    append_node(originals[93])
    caption('图B2  旧词表协议的3D-OVS三种子分布。点为种子均值，误差线为样本标准差。图中IPR均指旧词表版本，不能作为表2的误差线或配对统计。')
    para('B.3 历史方法示例与失败案例', 'Heading 2')
    append_node(originals[55])
    caption('图B3  历史方法示意与蛋挞查询示例。图中裁剪、掩码及定量标注来自旧版本，用于保留原始展示记录；其中语义证据面板不能代替当前的同ID成对图像一致性公式，图内总体增益不属于本文严格版主结果。')
    append_node(originals[96])
    caption('图B4  旧词表协议的目标恢复案例（seed 0）。列为RGB、真实掩码、受控基线和旧版IPR。Bench蛋挞、Blue Sofa太阳镜及Office Desk纸巾的五视角平均IoU分别由0.007、0.046、0.025提高至0.965、0.939、0.945；图内数值为所示视角IoU。尚不能把这些旧预测标为当前一致性模型的输出。')
    append_node(originals[99])
    caption('图B5  旧词表协议的失败案例（seed 0）。Covered Desk的gerbera在图像26对应视角由0.819降至0.078，Sofa的Gundam五视角平均IoU为0.031与0.021。当前严格版Covered Desk场景均值有所变化，但没有新逐视角证据时，不能声称图中的失败已被修复。')

    # Label legacy table headers and repeat headers without changing old numbers.
    for ti,t in enumerate(d.tables):
        widths = {
            3: [1.0,2.38,2.38],
            4: [1.65,2.25,0.94,0.94],
            5: [1.42,1.09,1.09,1.09,1.09],
        }[len(t.columns)]
        t.autofit = False
        for col,w in zip(t.columns,widths):
            col.width = Inches(w)
        for cell in t.rows[0].cells:
            if cell.text.strip() == 'IPR':
                set_text(cell.paragraphs[0], '旧词表版IPR', 9)
        for ri,row in enumerate(t.rows):
            for cell,w in zip(row.cells,widths):
                cell.width = Inches(w)
                for p in cell.paragraphs:
                    p.paragraph_format.keep_with_next = ri < len(t.rows)-1
                    p.paragraph_format.space_before = Pt(2)
                    p.paragraph_format.space_after = Pt(2)
                    p.paragraph_format.line_spacing = 1.05
                    p.paragraph_format.first_line_indent = Pt(0)
                    for r in p.runs:
                        font_run(r,9,ri == 0)
            pr = row._tr.get_or_add_trPr()
            for h in pr.findall(qn('w:trHeight')):
                pr.remove(h)
            if pr.find(qn('w:cantSplit')) is None:
                pr.append(OxmlElement('w:cantSplit'))
        if t.rows[0]._tr.get_or_add_trPr().find(qn('w:tblHeader')) is None:
            t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))

    # Use the retained template; keep image paragraphs with their captions.
    for p in d.paragraphs:
        if p._p.xpath('.//w:drawing'):
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.keep_together = True
            p.paragraph_format.first_line_indent = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            snap = OxmlElement('w:snapToGrid')
            snap.set(qn('w:val'), '0')
            p._p.get_or_add_pPr().append(snap)
        if p.style.name.startswith('Heading'):
            p.paragraph_format.keep_with_next = True
            for r in p.runs:
                font_run(r,13 if p.style.name == 'Heading 1' else 11.5,True)
        if re.match(r'^\[\d+\]', p.text):
            p.paragraph_format.keep_together = True
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.0
            snap = OxmlElement('w:snapToGrid')
            snap.set(qn('w:val'), '0')
            p._p.get_or_add_pPr().append(snap)
            for r in p.runs:
                font_run(r,9.5)
    historical_method = d.inline_shapes[3]
    ratio = historical_method.height / historical_method.width
    historical_method.width = Inches(4.8)
    historical_method.height = int(historical_method.width * ratio)
    d.core_properties.title = '面向ILGS的身份空间保持式伪标签可靠性建模'
    d.core_properties.subject = '统一查询无关IPR定义及受控实验协议'
    d.save(OUTPUT)
    assert hashlib.sha256(SOURCE.read_bytes()).hexdigest() == source_hash

    QA.mkdir(parents=True, exist_ok=True)
    verification = Document(OUTPUT)
    assert len(verification.inline_shapes) == 6
    assert len(verification.tables) == 8
    numbers = [p.text.strip() for p in verification.paragraphs if p._p.xpath('./m:oMath') and p.text.strip().startswith('(')]
    assert numbers == [f'({i})' for i in range(1,11)], numbers
    main_text = '\n'.join(p.text for p in verification.paragraphs).split('附录A')[0]
    assert '79.63%' not in main_text and '87.07%' not in main_text
    (QA / 'structural_checks.json').write_text(json.dumps({
        'source': str(SOURCE), 'source_sha256': source_hash,
        'output': str(OUTPUT), 'images_preserved': 6, 'tables': 8,
        'editable_equation_numbers': numbers,
        'pending': ['3D-OVS其余单分支的十场景三种子完整对照', '逐运行解码配置审计', '上游轨迹及语言特征视角来源审计', '预处理开销', '新版定性图', '逐种子配对统计'],
        'evidence': '用户提供的三种子LERF-Mask组件对照、四场景LERF-OVS、十场景三种子3D-OVS主结果及清理加基础置信对照；本地方法实现；保留源稿中的历史结果。',
    }, ensure_ascii=False, indent=2), encoding='utf-8')
    (QA / 'revision_notes.md').write_text('''# 2026-10-18 修订核对

- 源稿保持不变；修订稿另存于 D:/wlxy。
- 主方法统一为训练视角内的同身份CLIP图像裁剪一致性，阈值0.75。
- 按实际代码写入等权裁剪对余弦相似度；单裁剪取1仅表示不额外降权。
- 保留可编辑Word公式，替换旧文本softmax与面积投票公式。
- 主文表1替换为严格版LERF-Mask三种子结果74.46/69.91。
- 主文表2更新为统一配置的3D-OVS十场景三种子结果86.78/80.93。
- 主文表3加入3D-OVS清理加基础置信对照：83.00/76.91；完整IPR在其上再提高3.78/4.02个百分点。
- 主文表4加入LERF-Mask六设置三种子组件对照，明确基础置信的主要作用与附加门控的非单调收益。
- 标准四场景LERF-OVS保留全部均值和标准差，移至附录A。
- 旧词表结果及五幅旧定性/定量图归档于附录B，明确不属于当前方法。
- 基线统一称ILGS受控基线。日志分支official不是原论文复现的充分证据。
- 新版定性图、逐种子配对统计、预处理开销与3D-OVS其余单分支的完整十场景对照仍待补齐。
- 本次未从远程服务器获取原始日志，运行配置仍需与源日志核对。
- IPR训练子集限制不能证明上游DEVA跟踪、共享语言特征或自编码器未使用留出视角。
''', encoding='utf-8')
    print(OUTPUT)
    print('Structural checks passed; source preserved; 6 images; 8 tables; 10 editable equations.')


if __name__ == '__main__':
    build()
