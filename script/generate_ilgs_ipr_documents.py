from pathlib import Path
import zipfile

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor as PRGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches as PInches, Pt as PPt


OUT = Path(r"D:\wlxy")
FIG = Path(r"C:\Users\hp\Downloads\paper_figures_final")
DOCX = OUT / "ILGS_IPR_draft.docx"
PPTX = OUT / "ILGS_IPR_slides.pptx"
QA = OUT / "ILGS_IPR_generation_QA.md"

FIGS = {
    "pipeline": FIG / "fig_method_pipeline.png",
    "ablation": FIG / "fig_ablation_average.png",
    "teatime": FIG / "fig_teatime_mask_comparison.png",
    "ramen": FIG / "fig_ramen_mask_comparison.png",
    "figurines": FIG / "fig_figurines_mask_comparison.png",
}

MAIN_TABLE = [
    ["方法", "teatime mIoU", "teatime mBIoU", "figurines mIoU", "figurines mBIoU", "ramen mIoU", "ramen mBIoU", "Avg. mIoU", "Avg. mBIoU"],
    ["ILGS official-fixed", "72.87", "69.55", "65.99", "64.22", "79.21", "72.00", "72.69", "68.59"],
    ["Semantic-area IPR", "77.35", "73.20", "76.30", "73.02", "87.86", "78.02", "80.50", "74.75"],
]

ABLATION_TABLE = [
    ["方法", "figurines mIoU", "figurines mBIoU", "ramen mIoU", "ramen mBIoU", "teatime mIoU", "teatime mBIoU", "Avg. mIoU", "Avg. mBIoU"],
    ["Official-fixed", "65.99", "64.22", "79.21", "72.00", "72.87", "69.55", "72.69", "68.59"],
    ["Only area", "70.53", "67.61", "75.69", "66.33", "72.86", "69.62", "73.03", "67.85"],
    ["Only semantic", "70.21", "68.11", "78.30", "70.61", "73.40", "70.90", "73.97", "69.87"],
    ["Semantic-area no conf", "71.08", "68.10", "80.06", "71.62", "72.92", "69.77", "74.69", "69.83"],
    ["Full Semantic-area IPR", "76.30", "73.02", "87.86", "78.02", "77.35", "73.20", "80.50", "74.75"],
]

REFERENCES = [
    "Kerbl B, Kopanas G, Leimkühler T, et al. 3D Gaussian Splatting for Real-Time Radiance Field Rendering. ACM Transactions on Graphics, 2023.",
    "Kerr J, Kim C M, Goldberg K, et al. LERF: Language Embedded Radiance Fields. ICCV, 2023.",
    "Qin M, Li W, Zhou J, et al. LangSplat: 3D Language Gaussian Splatting. CVPR, 2024.",
    "Ye M, Danelljan M, Yu F, et al. Gaussian Grouping: Segment and Edit Anything in 3D Scenes. ECCV, 2024.",
    "Jang 等. Identity-aware Language Gaussian Splatting for Open-vocabulary 3D Semantic Segmentation. ICCV, 2025.",
    "Kirillov A, Mintun E, Ravi N, et al. Segment Anything. ICCV, 2023.",
    "Ravi N, Gabeur V, Hu Y T, et al. SAM 2: Segment Anything in Images and Videos. 2024.",
    "Cheng H K, Oh S W, Price B, et al. Tracking Anything with Decoupled Video Segmentation. ICCV, 2023.",
    "Oquab M, Darcet T, Moutakanni T, et al. DINOv2: Learning Robust Visual Features without Supervision. TMLR, 2024.",
]


def set_font(run, size=10.5, name="宋体"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)


def add_paragraph(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Pt(21)
    p.paragraph_format.line_spacing = 1.25
    run = p.add_run(text)
    set_font(run)
    return p


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell(cell, text, bold=False, size=8):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(str(text))
    run.bold = bold
    set_font(run, size=size)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, rows):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            cell = table.cell(i, j)
            set_cell(cell, value, bold=(i == 0 or i == len(rows) - 1), size=7)
            if i == 0:
                shade_cell(cell, "D9EAF7")
            elif i == len(rows) - 1:
                shade_cell(cell, "EAF6EA")
    return table


def add_figure(doc, path, caption, width=6.2):
    if not path.exists():
        add_paragraph(doc, f"[缺失图片：{path}]")
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Inches(width))
    cap = doc.add_paragraph(caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in cap.runs:
        set_font(run, size=9)


def build_docx():
    OUT.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.85)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    styles = doc.styles
    styles["Normal"].font.name = "宋体"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    styles["Normal"].font.size = Pt(10.5)
    for style_name in ["Heading 1", "Heading 2", "Heading 3"]:
        styles[style_name].font.name = "黑体"
        styles[style_name]._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("面向开放词汇三维语义分割的身份感知伪标签优化方法")
    run.bold = True
    set_font(run, size=16, name="黑体")

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = sub.add_run("Identity-aware Pseudo-label Refinement for Open-vocabulary 3D Semantic Segmentation")
    run.italic = True
    set_font(run, size=11, name="Times New Roman")

    doc.add_heading("摘要", level=1)
    add_paragraph(
        doc,
        "开放词汇三维语义分割旨在根据自然语言查询定位三维场景中的目标区域。近年来，基于 3D Gaussian Splatting 的语言场方法在渲染效率和开放词汇查询方面取得了较大进展。其中，ILGS 通过引入 identity-aware semantic consistency，有效缓解了多视角语言特征不一致问题。然而，ILGS 的训练依赖自动生成的 object mask 和 identity pseudo-labels，当这些伪标签存在身份混淆、背景混入或跨视角不一致时，错误监督会限制最终分割性能。针对这一问题，本文提出一种身份感知伪标签优化方法 Identity-aware Pseudo-label Refinement, IPR。该方法在不增加人工标注和不改变 ILGS 主干结构的条件下，从语义置信度、面积一致性和大区域混合检测三个方面评估 object identity 的可靠性，并生成 refined object masks 与 confidence maps，用于加权 ILGS 的身份一致性监督。实验结果表明，在 LERF 的 teatime、figurines 和 ramen 场景上，本文方法相较修复后的 ILGS official 流程平均 mIoU 从 72.69% 提升至 80.50%，平均 mBIoU 从 68.59% 提升至 74.75%。消融实验进一步表明，语义约束、面积约束和置信度加权训练共同构成了性能提升的主要来源。"
    )
    add_paragraph(doc, "关键词：开放词汇三维语义分割；3D Gaussian Splatting；伪标签优化；身份一致性；语义分割")

    doc.add_heading("1 引言", level=1)
    for text in [
        "开放词汇三维场景理解是三维视觉和视觉语言模型交叉领域的重要任务，其目标是使三维场景能够响应自然语言查询，并定位相应的物体或区域。早期方法通常基于 NeRF 表示构建三维语言场，例如 LERF 将 CLIP 等视觉语言模型的特征蒸馏到辐射场中，使模型能够根据文本查询生成三维相关性热力图。该类方法提高了三维场景的开放词汇理解能力，但 NeRF 表示训练和渲染成本较高，且查询结果常表现为连续响应区域，难以获得边界清晰的实例级分割结果。",
        "3D Gaussian Splatting 以显式高斯点表示三维场景，在保持高质量新视角合成的同时显著提升了渲染效率。基于这一表示，LangSplat 将语言特征嵌入 3D Gaussian 表示中，提升了开放词汇查询速度；Gaussian Grouping 进一步为高斯点引入 identity encoding，使三维场景具备 object-level grouping 能力。ILGS 在此基础上强调 identity-aware language field，通过身份一致性约束解决不同视角中语言特征不稳定的问题，并利用 progressive mask expanding 改善边界分割效果。",
        "尽管 ILGS 提升了开放词汇三维语义分割性能，但其训练过程仍依赖预先生成的 object mask 和 identity labels。这些标签通常来自自动化零样本跟踪或分组结果，并非人工标注，因此容易存在三类问题：同一 object id 可能覆盖多个语义对象或背景区域；真实物体在不同视角中的面积变化可能导致身份标签不稳定；小目标或边界区域容易被错误合并或遗漏。由于 ILGS 的 identity-aware supervision 直接使用这些伪标签，伪标签质量成为影响模型上限的重要因素。",
        "为此，本文提出一种面向 ILGS 的身份感知伪标签优化方法。与直接引入更强二维分割模型重新生成全部 mask 不同，本文保留原始 object id 的跨视角身份结构，仅对其可靠性进行评估和加权。该策略可以减少破坏原有 identity correspondence 的风险，并以较低代价提升训练监督质量。",
    ]:
        add_paragraph(doc, text)

    doc.add_heading("2 相关工作", level=1)
    for text in [
        "开放词汇三维语义分割方法主要可分为三类。第一类基于 NeRF 表示构建语言场，如 LERF 通过体渲染方式学习多尺度 CLIP 特征，使三维空间具备文本查询能力。该类方法具有较好的开放词汇表达能力，但训练和渲染效率受 NeRF 表示限制，且输出区域边界不够稳定。",
        "第二类方法基于 3D Gaussian Splatting 构建高效语言场。LangSplat 使用 3D Gaussian 表示语言特征，并通过语言 autoencoder 降低高维 CLIP 特征带来的存储和优化成本。Gaussian Grouping 则通过引入 identity encoding 将二维 SAM mask 监督提升到三维高斯点上，实现开放世界场景的分组和编辑。这类方法提高了效率和实例表达能力，但其性能高度依赖二维 mask 或伪标签质量。",
        "第三类方法关注跨视角一致性和身份约束。ILGS 通过 identity-aware semantic consistency loss 约束同一物体在不同视角中的语言表示，并使用 mask expanding 提升边界质量。该方法证明了 identity 信息对开放词汇三维分割的重要性。然而，ILGS 默认使用已有 object mask 作为 identity supervision，对伪标签本身的噪声处理仍然不足。本文工作正是针对这一空缺展开：不改变 ILGS 的三维语言建模框架，而是优化其输入伪标签和训练权重。",
    ]:
        add_paragraph(doc, text)

    doc.add_heading("3 方法", level=1)
    add_paragraph(doc, "本文方法的整体流程如图 1 所示。给定训练图像及其原始 object mask，首先统计每个 object id 在多视角中的语义响应和面积分布；然后根据语义置信度、面积一致性和大区域混合检测结果估计其可靠性；最后生成 refined object masks 和 confidence maps，并将其输入 ILGS 训练过程。")
    add_figure(doc, FIGS["pipeline"], "图 1 身份感知伪标签优化流程。")

    doc.add_heading("3.1 语义置信度评估", level=2)
    add_paragraph(doc, "对于每个 object id，本文从其多视角可见区域中裁剪图像 patch，并利用视觉语言模型计算其与场景查询类别之间的相似度。设第 i 个 object id 对类别 c 的累积响应为 S_i(c)，则其语义置信度定义为：P_i = max_c S_i(c) / sum_c S_i(c)。")
    add_paragraph(doc, "同时，为衡量 top-1 类别与其他类别之间的区分度，定义语义间隔：M_i = (S_i(c_1)-S_i(c_2)) / sum_c S_i(c)。当 P_i 或 M_i 较低时，说明该 object id 的语义归属不稳定，可能存在身份混淆。")

    doc.add_heading("3.2 面积一致性分析", level=2)
    add_paragraph(doc, "仅使用语义置信度无法识别大面积背景混合区域。为此，本文进一步统计每个 object id 在不同图像中的面积比例 a_i^t = |M_i^t|/(HW)。本文使用面积比例的 median、p90 和 max 描述其跨视角分布。当某个 object id 的最大面积或高分位面积异常偏大时，将其判定为 large mixed region，并在训练中降低其监督权重。")

    doc.add_heading("3.3 置信度加权身份监督", level=2)
    add_paragraph(doc, "根据语义和面积分析结果，本文为每个像素生成 confidence map。可靠 object 区域保持较高权重，低语义置信度区域、大面积混合区域和面积异常区域则降低权重。ILGS 原始身份一致性损失可写为 L_id，本文将其改写为 L'_id = w_i^t L_id，其中 w_i^t 表示第 t 帧中 object id i 的置信度权重。该设计并不直接删除可疑标签，而是降低其对训练的影响，从而避免错误伪标签被等权放大。")

    doc.add_heading("4 实验与分析", level=1)
    add_paragraph(doc, "本文在 LERF 数据集的 teatime、figurines 和 ramen 三个场景上进行实验。评价指标采用 mean Intersection over Union, mIoU 和 mean Boundary IoU, mBIoU。所有实验均基于 ILGS 官方流程复现，并在相同训练设置下比较不同伪标签处理策略。实验硬件为 NVIDIA RTX 3090 24GB，训练迭代数为 30000。")

    doc.add_heading("4.1 主结果比较", level=2)
    add_paragraph(doc, "表 1 给出了本文方法与 ILGS official-fixed 流程的比较。可以看到，Semantic-area IPR 在三个场景上均取得提升，平均 mIoU 从 72.69% 提升至 80.50%，平均 mBIoU 从 68.59% 提升至 74.75%。")
    add_paragraph(doc, "表 1 主结果比较（%）")
    add_table(doc, MAIN_TABLE)
    add_figure(doc, FIGS["teatime"], "图 2 teatime 场景分割可视化结果。")
    add_figure(doc, FIGS["ramen"], "图 3 ramen 场景分割可视化结果。")

    doc.add_heading("4.2 消融实验", level=2)
    add_paragraph(doc, "为验证各模块的作用，本文比较了 only area、only semantic、semantic-area no confidence 和完整 IPR。结果表明，单独使用面积约束或语义约束均只能带来有限提升。Only area 在 ramen 场景中甚至出现下降，说明面积统计无法充分区分真实小目标与异常区域。Only semantic 能够改善部分 identity 混淆，但对边界质量提升有限。Semantic-area no confidence 相比 official-fixed 有一定提升，但明显低于完整方法，说明仅修正 mask 不足以抑制训练阶段的噪声监督。完整 IPR 同时利用语义、面积和 confidence weighting，因此获得最优结果。")
    add_paragraph(doc, "表 2 消融实验结果（%）")
    add_table(doc, ABLATION_TABLE)
    add_figure(doc, FIGS["ablation"], "图 4 消融实验平均结果。")

    doc.add_heading("4.3 讨论与失败案例", level=2)
    add_paragraph(doc, "本文还探索了 SAM2、DEVA 和 DINOv3+SAM2 V2-anchor 等更强二维分割或跨视角匹配策略。实验发现，这些方法能够改善个别类别，例如 red apple 或 old camera，但整体结果并不稳定。图 5 展示了 figurines 场景中的部分现象：某些目标被改善，但 old camera 等类别仍容易失败。这说明对 ILGS 而言，二维 mask 边界质量并不是唯一关键，跨视角 identity consistency 同样重要。直接使用强二维分割模型重新生成 mask，可能破坏原始 object id 的跨视角对应关系。")
    add_figure(doc, FIGS["figurines"], "图 5 figurines 场景讨论案例。")

    doc.add_heading("5 结论", level=1)
    add_paragraph(doc, "本文针对 ILGS 中 object mask 和 identity pseudo-label 质量不足的问题，提出一种身份感知伪标签优化方法。该方法通过语义置信度、面积一致性和大区域混合检测评估 object identity 的可靠性，并利用 confidence map 对训练监督进行加权。实验结果表明，该方法在不增加人工标注和不改变 ILGS 主干结构的条件下，能够稳定提升 LERF 多场景开放词汇三维语义分割性能。进一步分析表明，强二维分割模型并不必然带来更好的三维分割结果，保持跨视角 identity consistency 是提升 identity-aware 3D language field 的关键因素之一。")

    doc.add_heading("参考文献", level=1)
    for i, ref in enumerate(REFERENCES, 1):
        p = doc.add_paragraph(f"[{i}] {ref}")
        for run in p.runs:
            set_font(run, size=9)

    doc.save(DOCX)


def add_textbox(slide, x, y, w, h, text, size=24, bold=False, color=None, align=None):
    color = color or PRGBColor(35, 35, 35)
    box = slide.shapes.add_textbox(PInches(x), PInches(y), PInches(w), PInches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    p = tf.paragraphs[0]
    if align:
        p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = PPt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Microsoft YaHei"
    return box


def add_title(slide, title, subtitle=None):
    blue = PRGBColor(36, 83, 164)
    add_textbox(slide, 0.55, 0.35, 12.3, 0.55, title, 25, True, blue)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, PInches(0.55), PInches(0.98), PInches(12.2), PInches(0.03))
    line.fill.solid()
    line.fill.fore_color.rgb = blue
    line.line.fill.background()
    if subtitle:
        add_textbox(slide, 0.6, 1.05, 12, 0.35, subtitle, 12, False, PRGBColor(95, 95, 95))


def add_bullets(slide, x, y, w, h, bullets, size=18):
    box = slide.shapes.add_textbox(PInches(x), PInches(y), PInches(w), PInches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    for idx, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = bullet
        p.font.size = PPt(size)
        p.font.name = "Microsoft YaHei"
        p.font.color.rgb = PRGBColor(35, 35, 35)
        p.space_after = PPt(8)
    return box


def add_image_fit(slide, path, x, y, w, h):
    path = Path(path)
    if not path.exists():
        add_textbox(slide, x, y, w, h, f"缺失图片\n{path.name}", 16, False, PRGBColor(95, 95, 95), PP_ALIGN.CENTER)
        return None
    with Image.open(path) as image:
        iw, ih = image.size
    scale = min(w / iw, h / ih)
    nw, nh = iw * scale, ih * scale
    return slide.shapes.add_picture(str(path), PInches(x + (w - nw) / 2), PInches(y + (h - nh) / 2), width=PInches(nw), height=PInches(nh))


def add_table_ppt(slide, rows, x, y, w, h, font_size=11):
    shape = slide.shapes.add_table(len(rows), len(rows[0]), PInches(x), PInches(y), PInches(w), PInches(h))
    table = shape.table
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            cell = table.cell(i, j)
            cell.text = str(value)
            for p in cell.text_frame.paragraphs:
                p.alignment = PP_ALIGN.CENTER
                for run in p.runs:
                    run.font.name = "Microsoft YaHei"
                    run.font.size = PPt(font_size)
                    run.font.bold = i == 0 or i == len(rows) - 1
            cell.fill.solid()
            if i == 0:
                cell.fill.fore_color.rgb = PRGBColor(221, 235, 247)
            elif i == len(rows) - 1:
                cell.fill.fore_color.rgb = PRGBColor(226, 239, 218)
            else:
                cell.fill.fore_color.rgb = PRGBColor(255, 255, 255)


def build_pptx():
    prs = Presentation()
    prs.slide_width = PInches(13.333)
    prs.slide_height = PInches(7.5)
    blank = prs.slide_layouts[6]
    blue = PRGBColor(36, 83, 164)
    green = PRGBColor(65, 156, 119)
    gray = PRGBColor(95, 95, 95)

    slide = prs.slides.add_slide(blank)
    add_textbox(slide, 0.75, 1.45, 11.8, 0.9, "面向开放词汇三维语义分割的\n身份感知伪标签优化方法", 34, True, blue, PP_ALIGN.CENTER)
    add_textbox(slide, 1.4, 3.25, 10.5, 0.45, "Identity-aware Pseudo-label Refinement for ILGS", 20, False, gray, PP_ALIGN.CENTER)
    add_textbox(slide, 1.4, 5.7, 10.5, 0.35, "开放词汇 3D 语义分割 · 3D Gaussian Splatting · 伪标签优化", 14, False, gray, PP_ALIGN.CENTER)

    slide = prs.slides.add_slide(blank)
    add_title(slide, "研究背景：3D 语言场正在从 NeRF 走向 3DGS")
    add_bullets(slide, 0.75, 1.45, 5.7, 4.8, ["开放词汇 3D 分割：用自然语言查询三维场景目标", "LERF 将 CLIP 特征蒸馏进 NeRF，但训练和渲染成本较高", "3D Gaussian Splatting 提供更高效的显式三维表示", "LangSplat、Gaussian Grouping、ILGS 将语言和身份信息引入 3DGS"], 18)
    add_textbox(slide, 7.0, 1.8, 5.2, 0.5, "核心趋势", 22, True, green)
    add_bullets(slide, 7.0, 2.45, 5.2, 2.7, ["从像素级语言热力图", "走向物体级 identity-aware language field", "但伪 object mask 质量成为新瓶颈"], 20)

    slide = prs.slides.add_slide(blank)
    add_title(slide, "关键问题：ILGS 的上限受 object mask / identity labels 限制")
    add_bullets(slide, 0.8, 1.45, 5.6, 4.9, ["ILGS 依赖自动生成的 object mask 作为 identity supervision", "这些伪标签不是人工标注，可能存在身份混淆和背景混入", "错误 identity labels 会被训练过程放大，影响最终 mIoU/mBIoU"], 19)
    add_bullets(slide, 7.0, 1.45, 5.5, 4.9, ["本文目标：不改变 ILGS 主干结构", "不增加人工标注", "只优化伪标签可靠性与训练权重", "提升开放词汇 3D 分割质量"], 19)

    slide = prs.slides.add_slide(blank)
    add_title(slide, "方法总览：Identity-aware Pseudo-label Refinement")
    add_image_fit(slide, FIGS["pipeline"], 0.55, 1.25, 12.2, 3.2)
    add_bullets(slide, 1.0, 4.9, 11.3, 1.2, ["保留原始 object IDs，避免破坏跨视角 identity correspondence", "通过语义置信度、面积一致性和大区域混合检测生成 confidence map", "训练阶段对不可靠区域降权，而不是简单删除或重新生成全部 mask"], 17)

    slide = prs.slides.add_slide(blank)
    add_title(slide, "关键设计：从“是否有 mask”转向“mask 是否可靠”")
    add_bullets(slide, 0.8, 1.35, 3.7, 4.9, ["语义置信度", "统计 object id 对文本查询类别的响应", "低置信或低 margin 表示身份语义不稳定"], 17)
    add_bullets(slide, 4.9, 1.35, 3.7, 4.9, ["面积一致性", "统计 median / p90 / max area ratio", "识别异常大区域和背景混合 object id"], 17)
    add_bullets(slide, 9.0, 1.35, 3.7, 4.9, ["置信度加权", "生成 confidence map", "训练时降低噪声 identity supervision 权重"], 17)

    slide = prs.slides.add_slide(blank)
    add_title(slide, "实验设置：LERF 三个场景 + mIoU/mBIoU")
    small_main = [["方法", "teatime", "figurines", "ramen", "Avg."], ["Official-fixed mIoU", "72.87", "65.99", "79.21", "72.69"], ["Semantic-area IPR mIoU", "77.35", "76.30", "87.86", "80.50"], ["Official-fixed mBIoU", "69.55", "64.22", "72.00", "68.59"], ["Semantic-area IPR mBIoU", "73.20", "73.02", "78.02", "74.75"]]
    add_table_ppt(slide, small_main, 0.8, 1.35, 11.7, 2.55, 13)
    add_bullets(slide, 0.9, 4.4, 11.4, 1.4, ["数据集：LERF teatime、figurines、ramen", "指标：mIoU 与 mBIoU", "训练：沿用 ILGS 主干和 30000 iterations，硬件为 RTX 3090 24GB"], 18)

    slide = prs.slides.add_slide(blank)
    add_title(slide, "主结果：平均 mIoU 提升 7.81，mBIoU 提升 6.16")
    add_table_ppt(slide, MAIN_TABLE, 0.35, 1.2, 12.65, 2.0, 8)
    add_bullets(slide, 0.9, 3.65, 11.4, 1.8, ["三场景均优于 official-fixed，说明方法不是单场景调参", "figurines 提升最大：mIoU +10.31，mBIoU +8.80", "ramen 达到最高 mIoU：87.86%"], 18)

    slide = prs.slides.add_slide(blank)
    add_title(slide, "定性结果：teatime 中小目标和错误大区域得到改善")
    add_image_fit(slide, FIGS["teatime"], 0.4, 1.2, 12.5, 5.6)

    slide = prs.slides.add_slide(blank)
    add_title(slide, "定性结果：ramen 中 glass of water / pork belly 更稳定")
    add_image_fit(slide, FIGS["ramen"], 0.4, 1.2, 12.5, 5.6)

    slide = prs.slides.add_slide(blank)
    add_title(slide, "消融实验：三个模块共同决定最终提升")
    add_image_fit(slide, FIGS["ablation"], 0.65, 1.15, 7.2, 5.7)
    add_bullets(slide, 8.1, 1.5, 4.7, 4.6, ["Only area：不稳定，ramen 下降", "Only semantic：有提升，但幅度有限", "No confidence：说明只修 mask 不够", "Full IPR：平均 mIoU/mBIoU 最优"], 17)

    slide = prs.slides.add_slide(blank)
    add_title(slide, "讨论：强二维分割模型不一定更适合 ILGS")
    add_image_fit(slide, FIGS["figurines"], 0.55, 1.2, 7.5, 5.5)
    add_bullets(slide, 8.35, 1.4, 4.4, 4.9, ["SAM2 / DEVA / V2-anchor 能改善个别类别", "但可能破坏跨视角 identity correspondence", "对 ILGS 而言，边界质量和 identity consistency 必须同时考虑"], 17)

    slide = prs.slides.add_slide(blank)
    add_title(slide, "结论与下一步")
    add_bullets(slide, 0.85, 1.35, 11.8, 4.7, ["提出 Semantic-area IPR：面向 ILGS 的轻量伪标签优化方法", "在不增加人工标注、不改变主干结构的条件下提升 mIoU/mBIoU", "消融证明：语义、面积和置信度加权训练共同发挥作用", "下一步：补充更多数据集、完善公式与参考文献、整理投稿格式"], 22)

    prs.save(PPTX)


def main():
    build_docx()
    build_pptx()
    with zipfile.ZipFile(PPTX, "r") as archive:
        slide_files = [n for n in archive.namelist() if n.startswith("ppt/slides/slide") and n.endswith(".xml")]
        media_files = [n for n in archive.namelist() if n.startswith("ppt/media/")]
    QA.write_text(
        f"""# ILGS IPR document generation QA

- Word draft: `{DOCX}`
- PPT draft: `{PPTX}`
- PPT slide count: {len(slide_files)}
- Embedded media files: {len(media_files)}
- Figures checked for existence: {sum(1 for p in FIGS.values() if p.exists())}/{len(FIGS)}
- Known limitations: formulas in Word are currently plain editable text, not native Office Math; references are an initial list and should be formatted to the target journal style later.
""",
        encoding="utf-8",
    )
    print(DOCX)
    print(PPTX)
    print(QA)


if __name__ == "__main__":
    main()
