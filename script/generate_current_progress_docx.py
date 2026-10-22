from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


OUT_DIR = Path(r"D:\wlxy")
DOCX_PATH = OUT_DIR / "ILGS_IPR_current_progress.docx"
FIG_DIR = Path(r"C:\Users\hp\Downloads\paper_figures_final")


MAIN_RESULTS = [
    ["方法", "teatime mIoU", "teatime mBIoU", "figurines mIoU", "figurines mBIoU", "ramen mIoU", "ramen mBIoU", "Avg. mIoU", "Avg. mBIoU"],
    ["Official-fixed baseline", "72.87", "69.55", "65.99", "64.22", "79.21", "72.00", "72.69", "68.59"],
    ["Semantic-area IPR", "77.35", "73.20", "76.30", "73.02", "87.86", "78.02", "80.50", "74.75"],
    ["提升", "+4.48", "+3.65", "+10.31", "+8.80", "+8.65", "+6.02", "+7.81", "+6.16"],
]

ABLATION_RESULTS = [
    ["消融版本", "figurines", "ramen", "teatime", "Avg. mIoU", "Avg. mBIoU"],
    ["Official-fixed", "65.99 / 64.22", "79.21 / 72.00", "72.87 / 69.55", "72.69", "68.59"],
    ["Only area", "70.53 / 67.61", "75.69 / 66.33", "72.86 / 69.62", "73.03", "67.85"],
    ["Only semantic", "70.21 / 68.11", "78.30 / 70.61", "73.40 / 70.90", "73.97", "69.87"],
    ["Semantic-area no conf", "71.08 / 68.10", "80.06 / 71.62", "72.92 / 69.77", "74.69", "69.83"],
    ["Full Semantic-area IPR", "76.30 / 73.02", "87.86 / 78.02", "77.35 / 73.20", "80.50", "74.75"],
]

EXPLORATORY_RESULTS = [
    ["探索版本", "代表结果", "观察结论"],
    ["SAM2 semantic masks", "figurines 74.06 / 70.99", "old camera、red apple 有改善，但 green toy chair、red toy chair 降低，说明强 2D mask 会破坏部分跨视角 identity。"],
    ["DEVA + SAM2 alignment", "figurines 70.58 / 68.22", "跨视角跟踪带来更一致的轨迹，但与 ILGS 原 anchor id 对齐不充分，整体低于 Semantic-area IPR。"],
    ["DINOv3/SAM2 V2-anchor strict", "figurines 71.19 / 68.13", "能做局部修复，但 cycle matching 仍会引入过多低置信区域，当前不作为主结果。"],
]

DEFAULT_PARAMS = [
    ["模块", "关键参数", "默认/当前设置", "作用"],
    ["Semantic crop scoring", "max_crops_per_id", "40", "每个 object id 最多采样 40 个高可见 crop，避免少数异常视角主导语义投票。"],
    ["Semantic crop scoring", "min_component_area", "500", "过滤过小连通域，减少碎片参与语义估计。"],
    ["Semantic crop scoring", "crop_padding", "12", "裁剪目标周围上下文，降低纯 mask crop 语义不足的问题。"],
    ["Visibility weighting", "visibility_sigma", "100.0", "TrackRef3D-style 面积可见性函数的尺度参数。"],
    ["Visibility weighting", "min_visibility_weight", "0.25", "面积异常视角最低保留权重，避免完全丢失稀有视角。"],
    ["Semantic confidence", "semantic_conf_threshold", "0.45", "低于该语义置信度则降权。"],
    ["Semantic confidence", "semantic_margin_threshold", "0.08", "top-1 与 top-2 语义差距过小时降权。"],
    ["Semantic confidence", "low_semantic_weight", "0.65", "语义不稳定 object 的训练监督权重。"],
    ["Large mixed region", "max_semantic_area_ratio / p90 / max", "0.55 / 0.35 / 0.45", "识别大面积背景/混合 object id。"],
    ["Large mixed region", "large_region_weight", "0.35", "大面积混合区域的监督降权。"],
    ["Training", "min_confidence", "0.05", "confidence map 的最小权重，防止分母过小。"],
    ["Training", "object CE weighting", "sum(c_i CE_i) / sum(c_i)", "将原 identity loss 改为 confidence-weighted loss。"],
]

FILES = [
    ["文件", "用途", "当前状态"],
    ["refine_object_masks.py", "早期 IPR：小连通域、边界、面积异常降权。", "保留，用于说明方法演进和对比。"],
    ["refine_object_masks_semantic_area.py", "主方法：语义置信度 + 面积一致性 + confidence map。", "主结果使用。"],
    ["script/train_lerf_semantic_area_refined_masks.sh", "使用 semantic-area refined masks 训练 ILGS。", "主训练脚本。"],
    ["train.py", "加入 object_confidence_path，并对 object identity CE loss 做 confidence weighting。", "核心代码改动。"],
    ["segmentation_mask_official_fixed.py", "只修复官方 mask 生成流程的兼容性，作为 fair baseline。", "baseline 使用。"],
    ["segmentation_mask_full_refine.py", "后处理：official、seed-ref、part-refine 候选 mask 选择。", "可作为辅助结果，不建议当前主打。"],
    ["generate_object_masks_sam2_semantic.py", "SAM2 + Florence/semantic prior 重新生成/修复 object mask。", "探索版本，暂不主打。"],
    ["align_deva_object_masks.py", "将 DEVA 跨视角轨迹 id 对齐到 ILGS anchor id。", "探索版本，暂不主打。"],
    ["refine_object_masks_v2anchor.py", "DINO feature cycle matching + SAM2 prompt 的 V2-anchor 尝试。", "探索版本，暂不主打。"],
]


def set_run_font(run, size=10.5, east_asia="宋体", latin="Times New Roman", bold=None):
    run.font.name = latin
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def set_doc_style(doc):
    styles = doc.styles
    styles["Normal"].font.name = "Times New Roman"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    styles["Normal"].font.size = Pt(10.5)
    for style_name, size in [("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 11)]:
        style = styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        style.font.size = Pt(size)
        style.font.bold = True


def add_p(doc, text, first_line=True):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.2
    if first_line:
        p.paragraph_format.first_line_indent = Pt(21)
    run = p.add_run(text)
    set_run_font(run)
    return p


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell(cell, text, bold=False, size=8):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(str(text))
    set_run_font(run, size=size, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, rows, font_size=8):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            cell = table.cell(i, j)
            set_cell(cell, value, bold=(i == 0 or i == len(rows) - 1), size=font_size)
            if i == 0:
                shade_cell(cell, "D9EAF7")
            elif i == len(rows) - 1:
                shade_cell(cell, "EAF6EA")
    return table


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_run_font(run, size=9)


def add_figure(doc, path, caption, width=6.2):
    if not path.exists():
        add_p(doc, f"[图缺失：{path}]", first_line=False)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Inches(width))
    add_caption(doc, caption)


def add_code_block(doc, code):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Pt(18)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(code)
    run.font.name = "Consolas"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(8.5)


def build_doc():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    set_doc_style(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("ILGS 优化当前进度与方法整理")
    set_run_font(run, size=16, east_asia="黑体", bold=True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Identity-aware Pseudo-label Refinement for Open-vocabulary 3D Semantic Segmentation")
    set_run_font(run, size=11, latin="Times New Roman")

    doc.add_heading("一、当前一句话结论", level=1)
    add_p(
        doc,
        "当前最适合写成小论文主线的结果，是 Semantic-area IPR：在不改变 ILGS 主干、不增加人工标注的前提下，利用语义置信度、跨视角面积一致性和 confidence-weighted identity loss 对 object identity pseudo-labels 进行保守优化。该方法在 LERF 的 teatime、figurines 和 ramen 三个场景上，将 official-fixed baseline 的平均 mIoU 从 72.69% 提升到 80.50%，平均 mBIoU 从 68.59% 提升到 74.75%。",
    )
    add_p(
        doc,
        "因此，当前论文的核心问题可以写为：ILGS 的性能上限受到自动生成 object_mask / coherent identity labels 质量限制；本文提出一种轻量的身份感知伪标签优化流程，在保持跨视角 object id 一致性的同时降低噪声伪标签对训练的污染。",
    )

    doc.add_heading("二、术语与版本锁定", level=1)
    ledger = [
        ["术语", "含义", "本文统一写法"],
        ["ILGS", "Identity-aware Language Gaussian Splatting", "ILGS"],
        ["IPR", "Identity-aware Pseudo-label Refinement", "IPR"],
        ["Official-fixed", "仅修复官方 mask 生成/评估兼容性，不加入启发式增强", "Official-fixed baseline"],
        ["Semantic-area IPR", "语义置信度 + 面积一致性 + confidence map 加权训练", "Semantic-area IPR"],
        ["Full-refine", "推理后 query-level mask 候选选择", "Full-refine / Query-aware refinement"],
        ["mIoU / mBIoU", "平均交并比 / 平均边界交并比", "mIoU / mBIoU"],
    ]
    add_table(doc, ledger, font_size=8)

    doc.add_heading("三、研究动机与问题定位", level=1)
    add_p(
        doc,
        "ILGS 的基本思想是将 language feature 与 identity-aware supervision 结合，使 3D Gaussian 场景不仅能响应自然语言查询，还能在不同视角中保持同一物体的语义一致性。这个设计依赖一个前提：训练用 object_mask 中的 identity labels 足够可靠。如果同一个 object id 混入背景或多个语义物体，或者同一个真实物体在不同视角中被分配为不同 id，那么 identity supervision 会把错误关系写入 3D language field。",
    )
    add_p(
        doc,
        "原始 ILGS 代码公开流程中并没有把 DEVA / zero-shot tracker 作为可运行模块重新生成 object_mask，而是默认使用数据集中已经准备好的 object_mask。因此，若论文要做有意义的优化，最直接也最稳的切入点不是重新设计 3DGS 或 language field，而是处理 object_mask 伪标签质量问题。这与 ILGS 论文中提到的 outlier filtering 思路是一致的：系统需要识别 rendered identity 与 pseudo label 不一致或不可靠的区域，并降低其训练影响。",
    )
    add_p(
        doc,
        "本项目已经验证，直接引入更强的 2D 分割/跟踪模型并不一定带来更好的 ILGS 结果。SAM2、DEVA 和 DINOv3/V2-anchor 在局部物体上有改善，但常常破坏原本跨视角的 identity anchor。说明 ILGS 的关键不是单帧 mask 越精细越好，而是 2D mask 边界质量与跨视角 identity consistency 之间的平衡。",
    )

    doc.add_heading("四、主方法：Semantic-area IPR", level=1)
    add_figure(doc, FIG_DIR / "fig_method_pipeline.png", "图 1 Semantic-area IPR 方法流程示意。")

    doc.add_heading("4.1 输入与输出", level=2)
    add_p(doc, "输入为每个 LERF 场景的原始或初步 refined object masks、可选 confidence maps、训练图像和文本查询类别；输出包括新的 object_mask_semantic_area_refined、object_confidence_semantic_area_refined，以及 object-level 语义报告。")
    add_code_block(
        doc,
        "输入:\n"
        "dataset/lerf/<scene>/object_mask_refined\n"
        "dataset/lerf/<scene>/object_confidence_refined\n"
        "dataset/lerf/<scene>/images\n\n"
        "输出:\n"
        "dataset/lerf/<scene>/object_mask_semantic_area_refined\n"
        "dataset/lerf/<scene>/object_confidence_semantic_area_refined\n"
        "dataset/lerf/<scene>/object_mask_semantic_area_report/object_semantics.json\n"
        "dataset/lerf/<scene>/object_mask_semantic_area_report/object_semantic_area_summary.csv",
    )

    doc.add_heading("4.2 语义置信度估计", level=2)
    add_p(
        doc,
        "对每个 object id，脚本先在所有视角中收集对应连通域。为避免极小碎片和异常视角影响判断，只保留面积大于 min_component_area 的连通域，并最多选取 max_crops_per_id 个面积较大的 crop。每个 crop 会在目标区域外保留 crop_padding 像素上下文，并使用 CLIP/open_clip 与场景查询词计算语义相似度。",
    )
    add_p(
        doc,
        "设 object id i 在第 k 个 crop 上对类别 c 的相似度 softmax 后为 p_i^k(c)，则该 object 的类别得分为可见性加权和：S_i(c)=Σ_k v_i^k p_i^k(c)。语义置信度为 top-1 类别得分占比，semantic margin 为 top-1 与 top-2 的差值占比。若置信度或 margin 过低，说明该 id 的语义不稳定，训练中需要降权。",
    )

    doc.add_heading("4.3 TrackRef3D-style 面积可见性", level=2)
    add_p(
        doc,
        "面积一致性参考 TrackRef3D 中 visibility-aware keyframe selection 的思想：一个 object 在不同视角中的可见面积不应长期极端偏大或偏小。本文不直接按照面积删除样本，而是把面积变成权重。设 A_t 是该 object 在视角 t 中的 mask 面积，A_med 是同一 object 多视角面积的中位数，则可见性权重可写为：",
    )
    add_p(doc, "v_t = A_t exp[-(sqrt(A_t)-sqrt(A_med))^2 / (2 sigma^2)]，并归一化到 [min_visibility_weight, 1]。")
    add_p(
        doc,
        "面积过小通常对应遮挡、碎片或低可见性；面积过大通常对应背景混入或大区域错误。使用该权重后，高质量视角对语义投票和训练监督贡献更大，异常视角仍被保留但影响降低。",
    )

    doc.add_heading("4.4 大面积混合区域抑制", level=2)
    add_p(
        doc,
        "实验中发现 object id=0、32 等区域常会覆盖大背景或多个物体，单纯语义投票可能给它一个看似合理的类别，但该 id 实际并不是 coherent object。因此脚本统计每个 object id 的面积比例 median、p90 和 max。当 median ≥ 0.55、p90 ≥ 0.35 或 max ≥ 0.45 时，将其标记为 large mixed region，并把该区域监督权重降为 large_region_weight=0.35。",
    )

    doc.add_heading("4.5 Confidence-weighted ILGS training", level=2)
    add_p(
        doc,
        "核心训练改动发生在 train.py。原始 ILGS 对 object identity CE loss 做逐像素平均。本文加入 object_confidence_path 后，从 object_confidence_semantic_area_refined 中读取每张图的 confidence map，并把 CE loss 改为加权平均：",
    )
    add_p(doc, "L_obj = Σ_i c_i CE(p_i, y_i) / max(Σ_i c_i, 1)。")
    add_p(
        doc,
        "其中 c_i 为第 i 个像素的伪标签置信度，y_i 为 object id pseudo-label，p_i 为渲染得到的 identity logits。这样做的好处是：不需要激进删除伪标签，也不会完全破坏跨视角 id 对应关系；不可靠区域仍参与训练，但对参数更新的影响显著降低。",
    )

    doc.add_heading("4.6 当前主方法参数", level=2)
    add_table(doc, DEFAULT_PARAMS, font_size=7)

    doc.add_heading("五、后处理方法：Full-refine 的当前定位", level=1)
    add_p(
        doc,
        "Full-refine 是 query-level 后处理，不改变训练结果，只在推理阶段基于 decoded feature、identity logits 和 CLIP similarity 生成多个候选 mask 并选择较可信者。候选包括 official mask、seed-ref mask 和 part-refine mask。seed-ref 用于修正 top-1 identity 选错导致的区域漂移；part-refine 用于 spoon handle 等小部件，只在 identity id 内部选择局部高相似区域。",
    )
    add_p(
        doc,
        "当前建议是：论文主线不要把 Full-refine 当成唯一贡献，而是把它作为 inference-side refinement 或附加模块讨论。原因是它对 teatime、ramen 的局部类别提升明显，但在 figurines 上可能拉低 IPR-only 的整体结果。因此主论文更稳的主线是 preprocessing-side Semantic-area IPR，Full-refine 可作为扩展实验或定性分析。",
    )

    doc.add_heading("六、探索但暂不主打的强版本", level=1)
    add_table(doc, EXPLORATORY_RESULTS, font_size=8)
    add_p(
        doc,
        "这些结果说明：更强的 2D 模型不能直接等价为更强的 ILGS 监督。SAM2 能提供更锐利的单帧边界，DEVA 能提供跨视角轨迹，DINOv3/V2-anchor 能做特征级匹配，但它们都需要重新对齐到 ILGS 的 identity anchor 空间。一旦这个对齐不稳定，训练阶段的 identity consistency 会被削弱，最终三维分割结果反而下降。",
    )

    doc.add_heading("七、实验结果汇总", level=1)
    add_p(doc, "表 1 是当前最重要的主结果，可直接作为论文核心实验表。")
    add_caption(doc, "表 1 Official-fixed baseline 与 Semantic-area IPR 对比（mIoU / mBIoU，%）")
    add_table(doc, MAIN_RESULTS, font_size=7)
    add_p(doc, "表 2 是当前已经补齐的必需消融实验，证明语义、面积和 confidence-weighted training 三部分均有作用。")
    add_caption(doc, "表 2 消融实验结果（每个场景格式为 mIoU / mBIoU，%）")
    add_table(doc, ABLATION_RESULTS, font_size=7)
    add_figure(doc, FIG_DIR / "fig_ablation_average.png", "图 2 消融实验平均结果。")
    add_figure(doc, FIG_DIR / "fig_teatime_mask_comparison.png", "图 3 teatime 定性对比。")
    add_figure(doc, FIG_DIR / "fig_ramen_mask_comparison.png", "图 4 ramen 定性对比。")
    add_figure(doc, FIG_DIR / "fig_figurines_mask_comparison.png", "图 5 figurines 定性对比与失败案例。")

    doc.add_heading("八、代码改动与文件用途", level=1)
    add_table(doc, FILES, font_size=7)

    doc.add_heading("九、可发表性判断", level=1)
    add_p(
        doc,
        "如果论文定位为“对 ILGS 的输入伪标签质量进行轻量优化”，目前结果已经具备小论文雏形：问题明确，方法与 ILGS 的已知弱点直接相关，三场景主结果和消融实验已经能支撑主要结论。最适合的贡献表述不是声称重新提出一个完整的 3D open-vocabulary segmentation 框架，而是提出一个可插拔的 identity-aware pseudo-label refinement module。",
    )
    add_p(
        doc,
        "仍需补强的部分包括：1）更严格的 related work 和引用，尤其是 ILGS、Gaussian Grouping、LangSplat、LERF、SAM、SAM2、DEVA、TrackRef3D/DINO 系列；2）明确数据集划分、训练迭代、GPU、超参数和评估脚本；3）补充失败案例解释，说明为什么 SAM2/DEVA/V2-anchor 不是最终主方法；4）若时间允许，加入 3D-OVS 或另一个非 LERF 数据集，验证泛化性。",
    )

    doc.add_heading("十、建议的论文结构", level=1)
    outline = [
        "题目：Identity-aware Pseudo-label Refinement for Open-vocabulary 3D Semantic Segmentation",
        "摘要：一句话背景、一句话 ILGS 瓶颈、一句话方法、一句话三场景提升、一句话边界。",
        "引言：从 3D language field 到 ILGS，再引出 object mask 伪标签质量限制。",
        "相关工作：3DGS language field、identity/grouping methods、pseudo-label refinement/track consistency。",
        "方法：Semantic scoring、area consistency、large mixed-region suppression、confidence-weighted ILGS training。",
        "实验：official-fixed baseline、主结果、消融、定性图、失败探索。",
        "讨论：为什么保守 refinement 比直接 SAM2/DEVA 重生成更适合 ILGS。",
    ]
    for item in outline:
        add_p(doc, item, first_line=False)

    doc.add_heading("十一、当前推荐命令流程", level=1)
    add_code_block(
        doc,
        "python refine_object_masks.py --dataset_path dataset/lerf/<scene>\n\n"
        "python refine_object_masks_semantic_area.py \\\n"
        "  --dataset_path dataset/lerf/<scene> \\\n"
        "  --src object_mask_refined \\\n"
        "  --conf_src object_confidence_refined\n\n"
        "CUDA_VISIBLE_DEVICES=0 bash script/train_lerf_semantic_area_refined_masks.sh lerf/<scene> 1\n\n"
        "CUDA_VISIBLE_DEVICES=0 python render_all.py \\\n"
        "  -m output/lerf/<scene>_semantic_area_refined_masks \\\n"
        "  --skip_train --train_split\n\n"
        "CUDA_VISIBLE_DEVICES=0 python segmentation_mask_official_fixed.py \\\n"
        "  -m output/lerf/<scene>_semantic_area_refined_masks\n\n"
        "python script/eval_lerf_mask.py -m output/lerf/<scene>_semantic_area_refined_masks",
    )

    doc.add_heading("十二、参考文献占位", level=1)
    refs = [
        "Kerbl B, Kopanas G, Leimkühler T, et al. 3D Gaussian Splatting for Real-Time Radiance Field Rendering. ACM TOG, 2023.",
        "Kerr J, Kim C M, Goldberg K, et al. LERF: Language Embedded Radiance Fields. ICCV, 2023.",
        "Qin M, Li W, Zhou J, et al. LangSplat: 3D Language Gaussian Splatting. CVPR, 2024.",
        "Ye M, Danelljan M, Yu F, et al. Gaussian Grouping: Segment and Edit Anything in 3D Scenes. ECCV, 2024.",
        "Jang et al. Identity-aware Language Gaussian Splatting for Open-vocabulary 3D Semantic Segmentation. ICCV, 2025.",
        "Kirillov A, Mintun E, Ravi N, et al. Segment Anything. ICCV, 2023.",
        "Ravi N, Gabeur V, Hu Y-T, et al. SAM 2: Segment Anything in Images and Videos. 2024.",
        "Cheng H K, Oh S W, Price B, et al. Tracking Anything with Decoupled Video Segmentation. ICCV, 2023.",
        "Oquab M, Darcet T, Moutakanni T, et al. DINOv2: Learning Robust Visual Features without Supervision. TMLR, 2024.",
        "TrackRef3D / DINOv3 相关文献：需后续根据最终引用版本核对 BibTeX 和页码。",
    ]
    for i, ref in enumerate(refs, 1):
        add_p(doc, f"[{i}] {ref}", first_line=False)

    doc.save(DOCX_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    build_doc()
