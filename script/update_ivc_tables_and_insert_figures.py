from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph


SOURCE = Path(r"D:\wlxy\ILGS-IPR_中文_受控基线与强基线补充_20261020.docx")
OUTPUT = Path(r"D:\wlxy\ILGS-IPR_中文_IVC表格与实验图整合_20261020.docx")
FIGURE_DIR = Path(r"D:\wlxy\ilgs\文章图\experiments_final_20261020")

FIGURES = {
    "lerf": FIGURE_DIR / "fig_exp_lerf_paired.png",
    "scene": FIGURE_DIR / "fig_exp_3dovs_scene_gains.png",
    "qual": FIGURE_DIR / "fig_exp_qualitative_success.png",
    "ablation": FIGURE_DIR / "fig_exp_component_ablation.png",
    "threshold": FIGURE_DIR / "fig_exp_threshold_sensitivity.png",
}

CAPTIONS = {
    "lerf": (
        "图4 LERF-Mask三场景三种子受控比较。浅色点表示各随机种子，实心标记与误差线表示均值±样本标准差；"
        "灰线连接同一场景和种子下的受控ILGS与ILGS + IPR。"
    ),
    "scene": (
        "图5 3D-OVS十场景受控比较。圆点和菱形分别表示三种子场景均值，数字为加入IPR后的百分点变化，"
        "虚线表示十场景平均值。"
    ),
    "qual": (
        "图6 3D-OVS真实测试视角的代表性改善案例（seed 0）。各行使用相同RGB视图和文本查询，依次展示真实掩码、"
        "受控ILGS预测及ILGS + IPR预测；预测掩码未经人工像素编辑。"
    ),
    "ablation": (
        "图7 LERF-Mask组件对照。点和误差线表示三个种子的三场景等权平均及样本标准差；完整IPR与各简化配置"
        "均使用相同训练视角和评测协议。"
    ),
    "threshold": (
        "图8 一致性阈值敏感性（seed 0）。红色空心圆和虚线标记本文采用的τ = 0.75；LERF-Mask覆盖三个场景，"
        "3D-OVS覆盖三个代表场景。"
    ),
}


def set_run_font(run, size: float, bold: bool | None = None, italic: bool | None = None) -> None:
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = None
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")


def set_paragraph_font(paragraph, size: float, bold: bool | None = None) -> None:
    if not paragraph.runs:
        paragraph.add_run("")
    for run in paragraph.runs:
        set_run_font(run, size, bold=bold)


def set_cell_border(cell, **edges) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)

    for edge_name, edge_data in edges.items():
        edge = tc_borders.find(qn(f"w:{edge_name}"))
        if edge is None:
            edge = OxmlElement(f"w:{edge_name}")
            tc_borders.append(edge)
        for key, value in edge_data.items():
            edge.set(qn(f"w:{key}"), str(value))


def remove_cell_shading(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    for shd in tc_pr.findall(qn("w:shd")):
        tc_pr.remove(shd)


def set_cell_margins(cell, top=70, start=90, bottom=70, end=90) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)
    tbl_header.set(qn("w:val"), "true")


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def clear_table_level_decoration(table) -> None:
    tbl_pr = table._tbl.tblPr
    for tag in ("w:tblBorders", "w:shd"):
        for node in tbl_pr.findall(qn(tag)):
            tbl_pr.remove(node)

    borders = OxmlElement("w:tblBorders")
    for edge_name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = OxmlElement(f"w:{edge_name}")
        edge.set(qn("w:val"), "nil")
        borders.append(edge)
    tbl_pr.append(borders)


def apply_three_line_table(table, font_size: float) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    clear_table_level_decoration(table)

    nil = {"val": "nil"}
    rule = {"val": "single", "sz": "8", "space": "0", "color": "000000"}
    thin = {"val": "single", "sz": "4", "space": "0", "color": "000000"}

    for row_index, row in enumerate(table.rows):
        prevent_row_split(row)
        if row_index == 0:
            set_repeat_table_header(row)
        for col_index, cell in enumerate(row.cells):
            remove_cell_shading(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cell, top=nil, left=nil, bottom=nil, right=nil, insideH=nil, insideV=nil)

            if row_index == 0:
                set_cell_border(cell, top=rule, bottom=thin)
            if row_index == len(table.rows) - 1:
                set_cell_border(cell, bottom=rule)

            for paragraph in cell.paragraphs:
                paragraph.alignment = (
                    WD_ALIGN_PARAGRAPH.LEFT if col_index == 0 and row_index > 0 else WD_ALIGN_PARAGRAPH.CENTER
                )
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
                set_paragraph_font(paragraph, font_size, bold=True if row_index == 0 else None)


def add_separator_above(table, row_index: int) -> None:
    thin = {"val": "single", "sz": "4", "space": "0", "color": "7F7F7F"}
    for cell in table.rows[row_index].cells:
        set_cell_border(cell, top=thin)


def bold_cell(cell, bold: bool = True) -> None:
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.bold = bold


def replace_cell_text(cell, text: str, size: float, bold: bool = False) -> None:
    cell.text = text
    for paragraph in cell.paragraphs:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        set_paragraph_font(paragraph, size, bold=bold)


def emphasize_pair(cell, left_bold: bool, right_bold: bool, size: float) -> None:
    text = cell.text.strip()
    parts = re.split(r"\s*/\s*", text, maxsplit=1)
    if len(parts) != 2:
        return
    paragraph = cell.paragraphs[0]
    paragraph.clear()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(parts[0])
    set_run_font(run, size, bold=left_bold)
    run = paragraph.add_run(" / ")
    set_run_font(run, size, bold=False)
    run = paragraph.add_run(parts[1])
    set_run_font(run, size, bold=right_bold)


def style_table_caption(paragraph) -> None:
    caption_text = paragraph.text.strip()
    match = re.match(r"^(表\s*[A-Za-z]?\d+)\s*(.*)$", caption_text)
    if match:
        paragraph.clear()
        label = paragraph.add_run(match.group(1))
        set_run_font(label, 9, bold=True)
        if match.group(2):
            remainder = paragraph.add_run("  " + match.group(2))
            set_run_font(remainder, 9, bold=False)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.keep_with_next = True
    set_paragraph_font(paragraph, 9)


def paragraph_after(paragraph: Paragraph) -> Paragraph:
    node = OxmlElement("w:p")
    paragraph._p.addnext(node)
    return Paragraph(node, paragraph._parent)


def paragraph_after_table(table) -> Paragraph:
    node = OxmlElement("w:p")
    table._tbl.addnext(node)
    return Paragraph(node, table._parent)


def insert_figure_after(anchor: Paragraph, image_path: Path, caption: str, width_inches: float = 6.45) -> Paragraph:
    picture_paragraph = paragraph_after(anchor)
    picture_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture_paragraph.paragraph_format.space_before = Pt(6)
    picture_paragraph.paragraph_format.space_after = Pt(2)
    picture_paragraph.paragraph_format.keep_with_next = True
    picture_paragraph.add_run().add_picture(str(image_path), width=Inches(width_inches))

    caption_paragraph = paragraph_after(picture_paragraph)
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_paragraph.paragraph_format.space_before = Pt(0)
    caption_paragraph.paragraph_format.space_after = Pt(6)
    caption_paragraph.paragraph_format.line_spacing = 1.0
    caption_paragraph.paragraph_format.keep_together = True
    run = caption_paragraph.add_run(caption)
    set_run_font(run, 9)
    run.bold = False
    return caption_paragraph


def find_paragraph(doc, fragment: str, last: bool = False) -> Paragraph:
    matches = [paragraph for paragraph in doc.paragraphs if fragment in paragraph.text]
    if matches:
        return matches[-1] if last else matches[0]
    raise RuntimeError(f"Paragraph not found: {fragment}")


def append_reference(paragraph: Paragraph, sentence: str) -> None:
    if sentence in paragraph.text:
        return
    run = paragraph.add_run(sentence)
    if paragraph.runs:
        reference = paragraph.runs[0]
        if reference.font.size:
            run.font.size = reference.font.size
        if reference.font.name:
            run.font.name = reference.font.name
        run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")


def apply_main_table_emphasis(doc) -> None:
    # Table 1: reported references and the controlled comparison are visually separated.
    table = doc.tables[0]
    replace_cell_text(table.cell(0, 2), "mIoU (%)↑", 8.2, bold=True)
    replace_cell_text(table.cell(0, 3), "mBIoU (%)↑", 8.2, bold=True)
    add_separator_above(table, 6)
    bold_cell(table.cell(5, 2))
    bold_cell(table.cell(5, 3))
    bold_cell(table.cell(7, 2))
    bold_cell(table.cell(7, 3))

    # Table 2: IPR improves every reported LERF-Mask mean.
    table = doc.tables[1]
    replace_cell_text(table.cell(0, 1), "ILGS受控基线\n(mIoU / mBIoU, %)", 7.5, bold=True)
    replace_cell_text(table.cell(0, 2), "ILGS + IPR\n(mIoU / mBIoU, %)", 7.5, bold=True)
    for row in range(1, len(table.rows)):
        bold_cell(table.cell(row, 2))
    add_separator_above(table, len(table.rows) - 1)

    # Table 3: the baseline remains better only for Snacks.
    table = doc.tables[2]
    replace_cell_text(table.cell(0, 1), "ILGS受控基线\n(mIoU / mBIoU, %)", 7.1, bold=True)
    replace_cell_text(table.cell(0, 2), "ILGS + IPR\n(mIoU / mBIoU, %)", 7.1, bold=True)
    for row in range(1, len(table.rows)):
        label = table.cell(row, 0).text.strip().lower()
        if "snacks" in label:
            bold_cell(table.cell(row, 1))
        else:
            bold_cell(table.cell(row, 2))
    add_separator_above(table, len(table.rows) - 1)

    # Table 4: fix the truncated mechanism label and emphasize the complete method.
    table = doc.tables[3]
    replace_cell_text(table.cell(0, 2), "mIoU (%)↑", 8.2, bold=True)
    replace_cell_text(table.cell(0, 3), "mBIoU (%)↑", 8.2, bold=True)
    replace_cell_text(table.cell(2, 1), "清理标签与基础置信图", 8.2)
    bold_cell(table.cell(len(table.rows) - 1, 2))
    bold_cell(table.cell(len(table.rows) - 1, 3))

    # Table 5: bold the best paired result for each scene and the overall mean.
    table = doc.tables[4]
    best_rows = {1: 6, 2: 3, 3: 3, 4: 3}
    for column, row in best_rows.items():
        bold_cell(table.cell(row, column))
    add_separator_above(table, 1)

    # Table 6: bold each metric maximum separately. The selected tau=0.75 row is also labelled in the first column.
    table = doc.tables[5]
    for column in range(1, 4):
        values = []
        for row in range(1, len(table.rows)):
            parts = re.split(r"\s*/\s*", table.cell(row, column).text.strip(), maxsplit=1)
            values.append((float(parts[0]), float(parts[1])))
        max_left = max(v[0] for v in values)
        max_right = max(v[1] for v in values)
        for row, (left, right) in enumerate(values, start=1):
            emphasize_pair(table.cell(row, column), left == max_left, right == max_right, 7.7)
    for row in range(1, len(table.rows)):
        if table.cell(row, 0).text.strip() == "0.75":
            replace_cell_text(table.cell(row, 0), "0.75*", 7.7, bold=True)

    note = paragraph_after_table(table)
    note.alignment = WD_ALIGN_PARAGRAPH.LEFT
    note.paragraph_format.space_before = Pt(2)
    note.paragraph_format.space_after = Pt(4)
    note.paragraph_format.line_spacing = 1.0
    run = note.add_run("注：粗体表示各列单项指标的最高值；*为本文统一采用的阈值。")
    set_run_font(run, 8)


def main() -> None:
    missing = [str(path) for path in (SOURCE, *FIGURES.values()) if not path.is_file()]
    if missing:
        raise FileNotFoundError("Missing inputs:\n" + "\n".join(missing))

    doc = Document(str(SOURCE))

    # Apply IVC/Elsevier manuscript table conventions to every editable table.
    for index, table in enumerate(doc.tables):
        size = 6.6 if index == 4 else (7.2 if index in (2, 5) or len(table.columns) >= 6 else 8.2)
        apply_three_line_table(table, size)

    apply_main_table_emphasis(doc)

    # Standardize all table captions above their tables.
    for paragraph in doc.paragraphs:
        if re.match(r"^表\s*\d+", paragraph.text.strip()):
            style_table_caption(paragraph)

    lerf_anchor = find_paragraph(doc, "Figurines的运行间波动")
    scene_anchor = find_paragraph(doc, "附录B的旧词表协议")
    ablation_anchor = find_paragraph(doc, "补充的标准LERF-OVS", last=True)
    threshold_anchor = find_paragraph(doc, "阈值升高会明显改变")

    append_reference(lerf_anchor, " 三种子配对分布见图4。")
    append_reference(scene_anchor, " 逐场景变化及真实测试视角的代表性结果分别见图5和图6。")
    append_reference(ablation_anchor, " 组件贡献及运行波动见图7。")
    append_reference(threshold_anchor, " 两个数据集上的阈值响应见图8。")

    insert_figure_after(lerf_anchor, FIGURES["lerf"], CAPTIONS["lerf"], 6.35)
    last = insert_figure_after(scene_anchor, FIGURES["scene"], CAPTIONS["scene"], 6.35)
    insert_figure_after(last, FIGURES["qual"], CAPTIONS["qual"], 6.45)
    insert_figure_after(ablation_anchor, FIGURES["ablation"], CAPTIONS["ablation"], 6.35)
    insert_figure_after(threshold_anchor, FIGURES["threshold"], CAPTIONS["threshold"], 6.35)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUTPUT))
    print(OUTPUT)


if __name__ == "__main__":
    main()
