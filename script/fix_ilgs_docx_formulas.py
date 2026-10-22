from pathlib import Path
import shutil

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


DOC_DIR = Path(r"D:\wlxy")
BACKUP_PATH = DOC_DIR / "ILGS_IPR_formula_backup.docx"
FALLBACK_PATH = DOC_DIR / "ILGS_IPR_formula_fixed.docx"

FORMULAS = {
    "Y_t(x)=i,    w_t(x) ∈ [0,1]": ("Y_t(x) = i,    w_t(x) ∈ [0, 1]", "1"),
    "S_i(c)=Σ_k v_i^k p_i^k(c)": ("S_i(c) = Σ_k v_i^k p_i^k(c)", "2"),
    "P_i=max_c S_i(c) / Σ_c S_i(c)": ("P_i = max_c S_i(c) / Σ_c S_i(c)", "3"),
    "M_i=(S_i(c_1)-S_i(c_2)) / Σ_c S_i(c)": ("M_i = [S_i(c_1) - S_i(c_2)] / Σ_c S_i(c)", "4"),
    "v_i^t = A_i^t · exp[-(√A_i^t - √A_{i,med})² / (2σ²)]": (
        "v_i^t = A_i^t · exp{ - (√A_i^t - √A_{i,med})² / (2σ²) }",
        "5",
    ),
    "L'_id = Σ_x w_t(x) CE(p_t(x), Y_t(x)) / (Σ_x w_t(x) + ε)": (
        "L'_{id} = [Σ_x w_t(x) CE(p_t(x), Y_t(x))] / [Σ_x w_t(x) + ε]",
        "6",
    ),
}


def find_target_docx():
    candidates = [
        p
        for p in DOC_DIR.glob("ILGS_IPR_*.docx")
        if p.name not in {
            "ILGS_IPR_current_progress.docx",
            "ILGS_IPR_draft.docx",
            "ILGS_IPR_formula_backup.docx",
        }
    ]
    if not candidates:
        raise FileNotFoundError("Could not find target ILGS_IPR_*.docx in D:\\wlxy")
    candidates.sort(key=lambda p: p.stat().st_mtime, reverse=True)
    return candidates[0]


def set_run_font(run, size=11, name="Cambria Math"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)


def set_cell_text(cell, text, align, font_name="Cambria Math", font_size=11):
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    set_run_font(run, size=font_size, name=font_name)


def remove_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is not None:
        tbl_pr.remove(borders)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = OxmlElement(f"w:{edge}")
        tag.set(qn("w:val"), "nil")
        borders.append(tag)
    tbl_pr.append(borders)


def insert_equation_table_after(doc, paragraph, equation, number):
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    table.columns[0].width = Inches(5.9)
    table.columns[1].width = Inches(0.55)
    table.cell(0, 0).width = Inches(5.9)
    table.cell(0, 1).width = Inches(0.55)
    remove_table_borders(table)

    set_cell_text(table.cell(0, 0), equation, WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_text(table.cell(0, 1), f"({number})", WD_ALIGN_PARAGRAPH.RIGHT, font_name="Times New Roman", font_size=10.5)

    left_para = table.cell(0, 0).paragraphs[0]
    right_para = table.cell(0, 1).paragraphs[0]
    for p in (left_para, right_para):
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)

    paragraph._p.addnext(table._tbl)
    paragraph._element.getparent().remove(paragraph._element)


def normalize_formula_context(doc):
    replacements = {
        "设训练图像集合为 {I_t}，对应的原始 object mask 为 {Y_t}。其中 Y_t(x)=i 表示第 t 个视角中像素 x 被分配给 object id i。ILGS 将 Y_t 作为 identity pseudo-label 监督渲染出的 identity feature map，并进一步利用 identity 信息约束 language embeddings 的多视角一致性。":
            "设训练图像集合为 {I_t}，对应的原始 object mask 为 {Y_t}。其中 Y_t(x)=i 表示第 t 个视角中像素 x 被分配给 object id i；本文进一步为每个像素定义置信度权重 w_t(x)，如式（1）所示。ILGS 将 Y_t 作为 identity pseudo-label 监督渲染出的 identity feature map，并进一步利用 identity 信息约束 language embeddings 的多视角一致性。",
        "设 object id i 在第 k 个 crop 上对类别 c 的 softmax 相似度为 p_i^k(c)，其可见性权重为 v_i^k，则该 object id 对类别 c 的累积语义响应为 S_i(c)。语义置信度定义为最高类别响应占全部类别响应的比例，semantic margin 定义为 top-1 与 top-2 类别响应之间的差距。当语义置信度或 margin 较低时，说明该 object id 的语义归属不稳定，可能混入多个物体或背景区域。":
            "设 object id i 在第 k 个 crop 上对类别 c 的 softmax 相似度为 p_i^k(c)，其可见性权重为 v_i^k，则该 object id 对类别 c 的累积语义响应 S_i(c) 定义为式（2）。语义置信度 P_i 定义为最高类别响应占全部类别响应的比例，如式（3）所示；semantic margin M_i 定义为 top-1 与 top-2 类别响应之间的差距，如式（4）所示。当语义置信度或 margin 较低时，说明该 object id 的语义归属不稳定，可能混入多个物体或背景区域。",
        "该权重随后被归一化到 [min_visibility_weight, 1]。通过这种方式，高质量视角对语义投票和训练监督贡献更大，异常视角仍被保留但影响降低。":
            "式（5）得到的可见性权重随后被归一化到 [min_visibility_weight, 1]。通过这种方式，高质量视角对语义投票和训练监督贡献更大，异常视角仍被保留但影响降低。",
        "核心训练改动发生在 train.py。原始 ILGS 对 object identity CE loss 做逐像素平均。本文加入 object_confidence_path 后，从 object_confidence_semantic_area_refined 中读取每张图的 confidence map，并把 CE loss 改为加权平均：":
            "核心训练改动发生在 train.py。原始 ILGS 对 object identity CE loss 做逐像素平均。本文加入 object_confidence_path 后，从 object_confidence_semantic_area_refined 中读取每张图的 confidence map，并将 CE loss 改为式（6）所示的加权平均形式：",
    }
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text in replacements:
            paragraph.text = replacements[text]
            for run in paragraph.runs:
                run.font.name = "Times New Roman"
                run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
                run.font.size = Pt(10.5)


def main():
    target = find_target_docx()
    if not BACKUP_PATH.exists():
        shutil.copy2(target, BACKUP_PATH)

    doc = Document(str(target))
    normalize_formula_context(doc)

    changed = 0
    # Collect first, then edit in reverse document order by paragraph index.
    todo = []
    for idx, paragraph in enumerate(doc.paragraphs):
        text = paragraph.text.strip()
        if text in FORMULAS:
            todo.append((idx, paragraph, *FORMULAS[text]))

    for _, paragraph, equation, number in reversed(todo):
        insert_equation_table_after(doc, paragraph, equation, number)
        changed += 1

    saved_path = target
    try:
        doc.save(str(target))
    except PermissionError:
        saved_path = FALLBACK_PATH
        doc.save(str(saved_path))
    print(f"target={target}")
    print(f"backup={BACKUP_PATH}")
    print(f"saved={saved_path}")
    print(f"formulas_fixed={changed}")


if __name__ == "__main__":
    main()
