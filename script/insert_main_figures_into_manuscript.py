from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.shared import Inches, Pt


SOURCE = Path(r"D:\wlxy\ILGS-IPR_中文_统一方法与协议修订_20261018.docx")
OUTPUT = Path(r"D:\wlxy\ILGS-IPR_中文_统一方法与协议修订_主图整合_20261018.docx")
FIGURE_DIR = Path(r"D:\wlxy\ilgs\文章图")

TEASER = FIGURE_DIR / "mask_noise_teaser_current_ipr.png"
PIPELINE = FIGURE_DIR / "pipeline_final_current_ipr.png"
METHOD = FIGURE_DIR / "method_final_current_ipr.png"

CONTENT_WIDTH = Inches(5.77)


def find_paragraph(doc, prefix):
    for paragraph in doc.paragraphs:
        if paragraph.text.strip().startswith(prefix):
            return paragraph
    raise RuntimeError(f"Paragraph not found: {prefix}")


def insert_paragraph_after(paragraph):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    return Paragraph(new_p, paragraph._parent)


def clear_paragraph(paragraph):
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def set_east_asia_font(run, name="宋体"):
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)


def format_figure_paragraph(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.keep_together = True


def format_caption(paragraph, text):
    clear_paragraph(paragraph)
    paragraph.style = "Normal"
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(1)
    paragraph.paragraph_format.space_after = Pt(7)
    paragraph.paragraph_format.line_spacing = 1.15
    paragraph.paragraph_format.keep_with_next = False
    paragraph.paragraph_format.keep_together = True
    run = paragraph.add_run(text)
    set_east_asia_font(run)
    run.font.size = Pt(11)
    run.font.bold = False


def add_figure(paragraph, image_path, alt_text):
    format_figure_paragraph(paragraph)
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=CONTENT_WIDTH)
    doc_pr = run._r.xpath(".//wp:docPr")
    if doc_pr:
        doc_pr[0].set("title", alt_text)
        doc_pr[0].set("descr", alt_text)


for required in (SOURCE, TEASER, PIPELINE, METHOD):
    if not required.is_file():
        raise FileNotFoundError(required)

doc = Document(SOURCE)

# Figure 1: place the visual motivation immediately after the problem statement.
motivation = find_paragraph(doc, "当一个身份区域混入背景或覆盖多个对象时")
if "典型噪声形态及IPR的干预动机如图1所示" not in motivation.text:
    motivation.add_run("典型噪声形态及IPR的干预动机如图1所示。")

teaser_paragraph = insert_paragraph_after(motivation)
add_figure(teaser_paragraph, TEASER, "身份伪标签中的典型噪声及IPR干预动机")
teaser_caption = insert_paragraph_after(teaser_paragraph)
format_caption(
    teaser_caption,
    "图1  真实预测掩码中的噪声及IPR的选择性作用。示例来自3D-OVS的blue_sofa场景，使用seed 0模型，与方法图采用不同场景。a，红框依据评测标注定位文本查询目标，仅用于可视化；b，统一协议下Controlled ILGS的原始预测掩码；c，ILGS+IPR的原始预测掩码。玩偶和遥控器分别选取view 2和view 3中具有代表性的改善案例，音箱使用view 2展示原本较干净的预测基本保持不变。所有预测掩码均未经人工像素编辑，本图用于定性说明，定量结论以正文完整场景与三种子汇总为准。",
)

# Figure 2: replace the existing main pipeline and correct its caption/protocol wording.
method_intro = find_paragraph(doc, "IPR以训练视角RGB图像及其既有身份掩码为输入")
method_intro.text = method_intro.text.replace(
    "整体位置如图1。",
    "整体流程如图2，可靠性估计与加权训练示意见图3。",
)

old_pipeline_caption = find_paragraph(doc, "图1  IPR在ILGS中的干预位置")
pipeline_paragraph = Paragraph(old_pipeline_caption._p.getprevious(), old_pipeline_caption._parent)
clear_paragraph(pipeline_paragraph)
add_figure(pipeline_paragraph, PIPELINE, "ILGS-IPR训练流程")
format_caption(
    old_pipeline_caption,
    "图2  IPR与ILGS的整体流程。a，IPR仅以训练视角RGB图像和既有噪声身份标签为输入；b，稀疏结构清理、查询无关的同身份裁剪一致性与大范围区域风险共同生成清理标签Y′和可靠性图W(p)；c，身份感知语言高斯表示保持不变；d，W(p)仅直接重加权二维身份分类损失，语言监督、RGB重建和动态离群过滤沿用共同ILGS实现。栅格面板均来自3D-OVS covered_desk场景seed 0的真实导出，未进行人工像素编辑；优化分支中的张量为代表性示例，箭头与损失框用于说明流程关系。",
)

# Figure 3: add the method detail at the end of the method section.
method_end = find_paragraph(doc, "式（10）仅替换二维身份分类项")
method_figure_paragraph = insert_paragraph_after(method_end)
add_figure(method_figure_paragraph, METHOD, "IPR可靠性估计与置信加权身份损失")
method_caption = insert_paragraph_after(method_figure_paragraph)
format_caption(
    method_caption,
    "图3  IPR可靠性估计与置信加权训练。a，原始身份伪标签包含占据大量像素的混合身份区域；b，稀疏结构清理、查询无关的同身份裁剪一致性和大范围区域风险形成像素可靠性图W(p)，并与在线动态离群门控o(p)共同重加权二维身份分类损失；c，语言特征、身份特征与查询分割由置信加权后的身份感知语言高斯场产生；d，在文本查询“Portuguese egg tart”下比较Official ILGS与ILGS+IPR的原始预测。栅格面板来自3D-OVS bench场景seed 0，预测掩码未经人工像素编辑；该查询由预先记录的全视角平均IoU改善规则自动选取，完整定量结论以正文多场景结果为准。",
)

doc.save(OUTPUT)
print(OUTPUT)
