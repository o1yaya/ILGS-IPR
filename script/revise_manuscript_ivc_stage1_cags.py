from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.text.paragraph import Paragraph


SOURCE = Path(r"D:\wlxy\ILGS-IPR_中文_3DOVS三种子补充_20261010_实验图增强版.docx")
OUTPUT = Path(r"D:\PycharmProjects\PythonProject\ILGS_release-main\output\manuscript\ILGS-IPR_中文_3DOVS三种子补充_20261010_IVC补强阶段1_CAGS定位.docx")


CAGS_POSITIONING = (
    "与伪标签预处理相邻的一类工作从三维表示内部建模上下文。CAGS [20]针对跨视角粒度不一致，"
    "在高斯局部图上传播上下文特征，并采用掩码级对比学习和邻接关系预计算以减轻三维实例碎片化。"
    "CAGS与本文的干预位置不同：前者改变高斯语义特征学习与实例匹配流程，后者在进入三维训练前估计"
    "二维身份伪标签的可靠性，仅重加权ILGS的身份分类监督，不引入图传播、新的高斯特征表示或额外推理分支。"
    "因此，两者属于可组合的互补路径；当前实验仅验证IPR作为ILGS前端可靠性模块的作用，尚未验证其与CAGS联合使用的效果。"
)

PROTOCOL_BOUNDARY = (
    "本文当前的LERF受控实验与CAGS等近期工作采用的LERF-OVS评测不是同一协议。本文在Figurines、Ramen和"
    "Teatime上按类别与场景汇总mIoU/mBIoU，并使用ILGS的分段扩展解码；LERF-OVS公开对比通常包含Waldo Kitchen、"
    "Ramen、Figurines和Teatime，报告mIoU与mAcc@0.25。由于注释、场景集合、文本选择与指标聚合方式不同，"
    "表1的数值不应与LERF-OVS文献结果直接排名；标准LERF-OVS对比需在独立的四场景协议下重新评测。"
)

CAGS_REFERENCE = (
    "[20] Sun W., Li Y., Jiao J. CAGS: Open-vocabulary 3D scene understanding with "
    "context-aware Gaussian splatting. Image and Vision Computing, 2026, 165: 105830. "
    "doi:10.1016/j.imavis.2025.105830."
)


def find_paragraph(document: Document, prefix: str) -> Paragraph:
    for paragraph in document.paragraphs:
        if paragraph.text.strip().startswith(prefix):
            return paragraph
    raise ValueError(f"Paragraph not found: {prefix}")


def insert_after(anchor: Paragraph, text: str) -> Paragraph:
    new_xml = deepcopy(anchor._p)
    for child in list(new_xml):
        if child.tag.endswith("}r") or child.tag.endswith("}hyperlink"):
            new_xml.remove(child)
    anchor._p.addnext(new_xml)
    paragraph = Paragraph(new_xml, anchor._parent)
    paragraph.add_run(text)
    return paragraph


def main() -> None:
    if not SOURCE.is_file():
        raise FileNotFoundError(SOURCE)
    document = Document(SOURCE)
    if not any("CAGS [20]" in p.text for p in document.paragraphs):
        insert_after(find_paragraph(document, "SAM2Object [11]"), CAGS_POSITIONING)
    if not any("LERF-OVS评测不是同一协议" in p.text for p in document.paragraphs):
        insert_after(find_paragraph(document, "候选词表复用各场景"), PROTOCOL_BOUNDARY)
    if not any(p.text.strip().startswith("[20] Sun W.") for p in document.paragraphs):
        insert_after(find_paragraph(document, "[19] Liu K."), CAGS_REFERENCE)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)
    print("CAGS mentions:", sum("CAGS" in p.text for p in document.paragraphs))


if __name__ == "__main__":
    main()
