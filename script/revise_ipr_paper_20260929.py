"""Rebuild the supplied manuscript with audited methods and controlled results."""
from copy import deepcopy
from pathlib import Path
import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'outputs/paper_revision_20260929'
OUT.mkdir(parents=True, exist_ok=True)
SOURCE = Path('D:/wlxy/ILGS-IPR_中文_删除FullRefine_方法公式修订版.docx')
doc = Document(SOURCE)
refs = [p.text for p in doc.paragraphs if re.match(r'^\[\d+\]', p.text)]
body = doc._element.body
for element in list(body):
    if element.tag != qn('w:sectPr'):
        body.remove(element)

def elem(name, *children, **attrs):
    e = OxmlElement(name)
    for k, v in attrs.items():
        e.set(qn(k), str(v))
    for c in children:
        if isinstance(c, list):
            e.extend(c)
        else:
            e.append(c)
    return e

def mr(text, plain=False):
    r = elem('m:r')
    if plain:
        r.append(elem('m:rPr', elem('m:sty', **{'m:val':'p'})))
    r.append(elem('w:rPr', elem('w:rFonts', **{'w:ascii':'Cambria Math', 'w:hAnsi':'Cambria Math'})))
    t = elem('m:t'); t.text = text
    t.set(qn('xml:space'), 'preserve'); r.append(t)
    return r

def seq(*xs):
    out=[]
    for x in xs:
        out.extend([mr(x)] if isinstance(x,str) else x if isinstance(x,list) else [x])
    return out

def sub(base, index):
    return elem('m:sSub', elem('m:e', seq(base)), elem('m:sub', seq(index)))

def sup(base, power):
    if not isinstance(base,str) and not isinstance(base,list) and base.tag == qn('m:sSub'):
        return elem('m:sSubSup',deepcopy(base.find(qn('m:e'))),deepcopy(base.find(qn('m:sub'))),elem('m:sup',seq(power)))
    return elem('m:sSup', elem('m:e', seq(base)), elem('m:sup', seq(power)))

def frac(a,b):
    return elem('m:f', elem('m:num', seq(a)), elem('m:den', seq(b)))

def sqrt(a):
    return elem('m:rad', elem('m:radPr', elem('m:degHide', **{'m:val':'1'})), elem('m:deg'), elem('m:e',seq(a)))

def summ(index, value):
    return elem('m:nary', elem('m:naryPr', elem('m:chr', **{'m:val':'∑'}), elem('m:limLoc', **{'m:val':'subSup'}),elem('m:supHide',**{'m:val':'1'})),elem('m:sub',seq(index)),elem('m:sup'),elem('m:e',seq(value)))

def hat(base):
    return elem('m:acc',elem('m:accPr',elem('m:chr',**{'m:val':'̃'})),elem('m:e',seq(base)))

def lim(name, index):
    return elem('m:limLow',elem('m:e',mr(name,True)),elem('m:lim',seq(index)))

def math(p, tokens):
    m=elem('m:oMath',seq(tokens)); p._p.append(m); return m

inline = {
    'I_t':sub('I','t'), 'T':mr('T'), 't':mr('t'), 'x':mr('x'), 'i':mr('i'), 'k':mr('k'),
    'f_ik':sub('f','ik'), 'e_c':sub('e','c'), 'c':mr('c'), 'd':mr('d'),
    'c1':sub('c','1'), 'c2':sub('c','2'), 'sigma':mr('σ'),
    'A_med':sub('A',seq('i,',mr('med',True))),
    'Omega_t':sub('Ω','t'), 'Y_t':sub('Y','t'), 'Ytilde_t':hat(sub('Y','t')),
    'w_t':sub('w','t'), 'c0_t':sup(sub('c','t'),'0'), 's_i':sub('s','i'),
    'q_i':sub('q','i'), 'm_i':sub('m','i'), 'S_i':sub('S','i'), 'C_s':sub('C','s'),
    'alpha_ik':sub('α','ik'), 'A_ik':sub('A','ik'), 'r_ik':sub('r','ik'),
    'K_i':sub('K','i'), 'R_ik':sub('R','ik'), 't_ik':sub('t','ik'), 'g_i':sub('g','i'),
    'a_ik':sub('a','ik'), 'C':mr('C'), 'o_t':sub('o','t'), 'ell_t':sub('ℓ','t'),
    'n':mr('n'), 'v_it':sub('v','it'), 'H_t':sub('H','t'),
    'h_t':sub('h','t'), 'Lcls':sub('L',mr('cls',True)),
    'LclsIPR':sup(sub('L',mr('cls',True)),mr('IPR',True)),
}

def font(run, size=11, bold=False):
    run.font.name='Times New Roman'; run.font.size=Pt(size); run.bold=bold
    rp=run._element.get_or_add_rPr(); rf=rp.find(qn('w:rFonts'))
    if rf is None: rf=elem('w:rFonts'); rp.append(rf)
    rf.set(qn('w:eastAsia'),'宋体'); run.font.color.rgb=RGBColor(0,0,0)

def para(text='', style=None, size=11, indent=True):
    p=doc.add_paragraph(style=style)
    pf=p.paragraph_format
    pf.space_after=Pt(5); pf.line_spacing=1.18
    pf.first_line_indent=Pt(22) if indent and not style else Pt(0)
    pf.widow_control=True
    for j, part in enumerate(re.split(r'\$([^$]+)\$', text)):
        if j%2: math(p,deepcopy(inline[part]))
        elif part: font(p.add_run(part),size)
    return p

def heading(text,level=1):
    p=para(text,f'Heading {level}',14 if level==1 else 12)
    p.paragraph_format.keep_with_next=True
    p.paragraph_format.space_before=Pt(12 if level==1 else 8)
    for r in p.runs: r.bold=True
    return p

def equation(tokens,number):
    p=para(indent=False)
    pf=p.paragraph_format; pf.space_before=Pt(5); pf.space_after=Pt(8)
    pf.line_spacing=1.15; pf.keep_together=True
    pf.tab_stops.add_tab_stop(Cm(7.325),WD_TAB_ALIGNMENT.CENTER)
    pf.tab_stops.add_tab_stop(Cm(14.65),WD_TAB_ALIGNMENT.RIGHT)
    p.add_run('\t'); m=math(p,tokens); font(p.add_run('\t('+str(number)+')'))
    for run in m.iter(qn('m:r')):
        run.find(qn('w:rPr')).append(elem('w:sz',**{'w:val':'24'}))
    return p

def caption(text):
    p=para(text,size=10,indent=False); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next=True
    return p

def table(headers, rows, widths):
    t=doc.add_table(rows=1,cols=len(headers)); t.autofit=False
    for c,w in zip(t.columns,widths): c.width=Cm(w)
    for j,s in enumerate(headers): t.rows[0].cells[j].text=s
    for row in rows:
        cells=t.add_row().cells
        for c,s in zip(cells,row): c.text=str(s)
    pr=t._tbl.tblPr
    borders=elem('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        borders.append(elem('w:'+edge,**{'w:val':'single','w:sz':'4','w:color':'B7B7B7'}))
    pr.append(borders)
    for i,row in enumerate(t.rows):
        row._tr.get_or_add_trPr().append(elem('w:cantSplit'))
        for c,w in zip(row.cells,widths):
            c.width=Cm(w)
            tcpr=c._tc.get_or_add_tcPr(); tcpr.append(elem('w:vAlign',**{'w:val':'center'}))
            if i==0: tcpr.append(elem('w:shd',**{'w:fill':'EDEDED'}))
            for p in c.paragraphs:
                p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.first_line_indent=Pt(0)
                p.paragraph_format.left_indent=Pt(0);p.paragraph_format.right_indent=Pt(0)
                p.paragraph_format.space_before=Pt(4);p.paragraph_format.space_after=Pt(4)
                p.paragraph_format.line_spacing=1.1
                for r in p.runs: font(r,9.5,i==0)
    t.rows[0]._tr.get_or_add_trPr().append(elem('w:tblHeader'))
    para('',indent=False).paragraph_format.space_after=Pt(0)
    return t

for name in ['Normal','Title','Heading 1','Heading 2','Heading 3']:
    if name not in doc.styles:
        doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
    st=doc.styles[name]; st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0)
    st._element.get_or_add_rPr().append(elem('w:rFonts',**{'w:eastAsia':'宋体','w:ascii':'Times New Roman','w:hAnsi':'Times New Roman'}))
title=para('面向ILGS的身份空间保持式伪标签可靠性建模','Title',17)
title.alignment=WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after=Pt(12)
para('摘要：身份感知语言高斯依赖跨视角对象伪标签构建语义一致性，但自动生成的身份区域可能包含边界噪声、背景泄漏或多对象混合。仅依赖训练交叉熵难以区分可信身份与可被模型拟合的错误标签。本文提出身份空间保持式伪标签可靠性建模方法（Identity-aware Pseudo-label Refinement，IPR），在已有对象ID集合内进行局部清理，利用跨视角裁剪语义响应和连通区域面积分布估计监督可靠性，并通过像素级置信图加权二维身份分类损失。该方法不创建新ID，也不重新建立全局对象轨迹。对三个LERF场景的三种子受控实验表明，在统一损失归一化与查询解码协议下，IPR将场景平均mIoU由75.89%提高至79.63%，将mBIoU由71.25%提高至74.46%。消融结果表明，大范围区域降权具有较强的平均贡献，语义可靠性进一步提高总体指标，而额外的逐视角可见性乘子未带来收益。结论限定于当前场景词表已知的评测设定；跨数据集泛化和未见词汇适用性仍需进一步验证。',size=10.5,indent=False)
para('关键词：开放词汇三维语义分割；身份伪标签；三维高斯泼溅；监督可靠性；置信度加权',size=10.5,indent=False)

heading('1 引言')
para('开放词汇三维语义分割旨在根据自然语言查询定位三维场景中的目标。NeRF [1]及其语言场扩展LERF [2]通过蒸馏CLIP [3]等视觉语言特征，使场景表示具备文本检索能力。3D Gaussian Splatting（3DGS）[4]进一步采用显式高斯表示和可微光栅化，为高效场景重建与语义查询提供了基础。LangSplat [5]、Feature 3DGS [6]和LEGaussian [7]将二维语言特征提升至高斯表示，但二维监督的噪声及其跨视角不一致仍会影响对象级分割。')
para('身份信息为缓解这一问题提供了结构约束。Gaussian Grouping [9]将多视角对象分组引入高斯表示，ILGS [10]进一步利用身份感知语义一致性，使同身份语言特征聚合、不同身份语言特征分离。本文关注该流程中的一个具体环节：用于训练的对象身份掩码是否足够可靠。由SAM [8]和DEVA [15]等自动方法得到的对象轨迹，并不保证每个ID都对应纯净且完整的单一对象。')
para('当一个身份区域混入背景或覆盖多个对象时，模型可能逐渐拟合该伪标签，并将本不相关的语言特征绑定到同一身份。ILGS的交叉熵离群点过滤可以抑制当前难以拟合的像素，却不能直接判断一个低损失区域在语义上是否合理。因此，本文拟解决的关键问题是：在保留已有跨视角ID来源和分类空间的条件下，如何识别不可信的身份监督并降低其对三维训练的贡献。这里的目标是控制错误监督的作用强度，而非恢复完全正确的对象轨迹。')
para('本文提出IPR，将局部标签清理与监督可靠性估计连接到ILGS身份分类分支。局部清理仅在已有ID集合内重新分配小连通分量；随后，以对象裁剪的语义响应和大范围区域风险构造对象权重，并与基础置信图融合。最终配置不在训练置信图上叠加连续的逐视角可见性乘子，但仍保留语义投票内部的面积权重。该设计将“如何汇总语义证据”与“如何分配训练权重”明确区分。')
para('本文的贡献包括三个方面。第一，明确界定既有身份空间中的监督可靠性问题，并给出不引入新ID的局部处理约束。第二，构建基于语义响应和大范围连通区域风险的置信度加权身份监督，保持ILGS的高斯表示及其他训练分支。第三，在三个LERF场景、三个训练种子和统一查询协议下开展受控比较，区分局部清理、基础置信度及附加可靠性项的作用，同时报告退化情况与适用边界。')

heading('2 相关工作')
heading('2.1 从语言场到对象身份监督',2)
para('LERF [2]将多尺度视觉语言特征融入NeRF，使三维位置能够响应开放词汇查询。LangSplat [5]以高斯表示和语言自编码器降低特征存储与渲染负担，Feature 3DGS [6]及LEGaussian [7]从不同角度研究二维特征到三维表示的提升。此类工作主要改善语言场的表示与学习方式，而对象边界及跨视角语义一致性仍依赖上游二维监督。')
para('Gaussian Grouping [9]引入身份编码以组织三维对象；ILGS [10]通过身份感知约束对齐语言嵌入。本文以ILGS为基础，研究其对象伪标签进入训练前的可靠性估计，以及置信图如何作用于二维身份分类损失。身份表示、语言编码与三维语义一致性机制归属于原有框架。')
heading('2.2 多视角标签一致性与可靠性建模',2)
para('SAM2Object [11]通过SAM2与几何信息整合多视角实例；Unified-Lift [12]在二维到三维实例提升中引入对象级表示、关联学习和噪声标签过滤；SeConGS [13]利用跨视角一致的分割监督服务于语言高斯查询。TrackRef3D [17]以先跟踪后赋义的流程汇总轨迹语义，Multi-Cue Mask Refinement [18]则融合多种线索改进掩码关系。这些研究说明，二维到三维提升的质量不仅取决于特征能力，还取决于监督的组织方式与可靠性。')
para('与上述工作相比，本文研究范围更窄：在已有对象ID空间内，使用离线估计的置信图调整监督贡献。语义投票采用面积加权的跨视角证据汇总，思想上与轨迹级语义聚合相关，但本文不执行语义驱动的全局ID合并。训练阶段的误差过滤与预处理阶段的语义、面积风险估计相互补充；本文不将降权操作等同于噪声标签已被正确修复。')

heading('3 方法')
para('IPR的输入为多视图RGB图像、已有对象身份掩码及相机信息，输出为清理后的身份标签和像素级置信图。整体流程如图1。候选语义词表用于评估可靠性，不直接替换身份标签；置信图仅作用于二维身份分类损失。')
for line in [
    '原始对象ID掩码 → 已有ID内局部清理 → 清理标签与基础置信图',
    '对象连通区域 → 大范围风险判断 ／ 裁剪语义投票 → 对象可靠性',
    '基础置信图 × 对象可靠性 → IPR像素置信图',
    'ILGS身份渲染 → 逐像素交叉熵 → 动态过滤 × IPR置信图 → 加权损失',
    'RGB重建、语言特征监督与三维身份感知语义约束保留',
]:
    p=para(line,size=10,indent=False);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next=True
caption('图1  IPR在ILGS训练流程中的位置（可编辑流程示意）')
para('面积加权语义投票属于离线证据汇总；最终训练权重中不再使用额外的逐视角连续可见性系数。共同的查询解码协议在实验设置中单独说明，不作为IPR的训练创新。')

heading('3.1 任务定义与身份空间约束',2)
para('设第$t$幅输入图像为$I_t$，像素位置$x$属于像素域$Omega_t$，原始身份标签为$Y_t$，清理后的标签为$Ytilde_t$。IPR置信图记为$w_t$。局部清理仅复用原始图像中已有的ID，语义与大范围风险估计阶段不再改变清理后的标签，因此满足：')
equation(seq(mr('ID',True),'(',hat(sub('Y','t')),') ⊆ ',mr('ID',True),'(',sub('Y','t'),'),   ',sup(sub('Y','t'),mr('IPR',True)),'=',hat(sub('Y','t')),',   0.05 ≤ ',sub('w','t'),'(x) ≤ 1'),1)
para('式（1）的保持性是ID集合意义下的约束，不保证所有像素标签或对象形状不变，也不保证原始轨迹本身正确。若一个ID始终混合多个对象，本文只能降低其监督作用，不能将其自动拆分为多个新身份。')
heading('3.2 局部清理与基础置信图',2)
para('对每个ID的连通分量进行检查。面积小于64像素时，在三次二值膨胀形成的外环中统计其他已有ID，并将该分量赋给占比最高的邻居；无可用邻居时保留原标签。该步骤为局部启发式修正，不执行全局语义匹配。')
para('基础置信图$c0_t$由三类局部因素构成：小连通分量乘以0.25，发生重分配时再乘以0.5；身份边界及其扩展邻域乘以0.65；若清理后某ID的总像素面积低于原始跨视图出现面积中位数的0.15倍，或高于其6倍，则该ID区域再乘以0.5。边界由3×3局部最大值与最小值不一致的位置检测，并向外膨胀一次。各因素相乘后截断至[0.05, 1]。')
para('基础置信度中的极端面积检测与后文的大范围风险规则不同：前者比较同一ID在不同图像中的总面积，后者检查单个连通分量占整幅图像的比例。基础置信图和最终置信图均保存为8位PNG，训练读取时除以255；由量化产生的下限偏差小于1/255。')
heading('3.3 跨视角裁剪语义可靠性',2)
para('对清理后的每个非忽略ID $i$，收集面积不小于500像素的连通分量，将其中第$k$个分量的像素集合记为$R_ik$；$k$枚举分量而不是视角，同一视角可贡献多个分量。按像素面积降序选取至多40个裁剪。裁剪外扩12像素，掩码外背景以该分量像素均值填充。对于未进入上述集合的像素，后续保留基础置信度。')
para('使用OpenCLIP ViT-B/16对裁剪和候选词表$C_s$编码。令图像与文本嵌入均经过L2归一化，类别概率由余弦相似度经温度缩放后的softmax给出：')
equation(seq(sub('p','ik'),'(c)=',frac(seq(mr('exp',True),'(35⟨',sub('f','ik'),',',sub('e','c'),'⟩)'),summ(seq('d∈',sub('C','s')),seq(mr('exp',True),'(35⟨',sub('f','ik'),',',sub('e','d'),'⟩)')))),2)
para('其中$f_ik$为裁剪图像的归一化嵌入，$e_c$为类别$c$的归一化文本嵌入，$d$遍历候选词表，尖括号表示内积。对某个ID，使用面积投票权重$alpha_ik$累积响应，并以最大类别得分和前两类差值定义语义置信度$q_i$及间隔$m_i$：')
equation(seq(sub('S','i'),'(c)=',summ('k',seq(sub('α','ik'),sub('p','ik'),'(c)'))),3)
equation(seq(sub('q','i'),'=',frac(seq(sub('S','i'),'(',sub('c','1'),')'),summ(seq('c∈',sub('C','s')),seq(sub('S','i'),'(c)'))),',   ',sub('m','i'),'=',frac(seq(sub('S','i'),'(',sub('c','1'),')−',sub('S','i'),'(',sub('c','2'),')'),summ(seq('c∈',sub('C','s')),seq(sub('S','i'),'(c)')))),4)
para('这里$c1$和$c2$分别为累计响应最高和次高的候选类别。至少有一个有效裁剪且候选词表非空时，式（4）的分母为正；词表仅含一个类别时，实现令$m_i$等于$q_i$。这些量表示给定词表下的相对语义集中程度，不能解释为校准后的标签正确率。')
heading('3.4 面积加权投票及其与训练权重的区别',2)
para('为汇总不同质量的裁剪证据，语义投票沿用面积加权规则。设$A_ik$为选中裁剪对应分量的像素面积，$A_med$为这些选中分量的中位面积，$sigma$为尺度参数。原始面积分数$r_ik$及投票权重$alpha_ik$为：')
equation(seq(sub('r','ik'),'=',sub('A','ik'),mr('exp',True),'[−',frac(sup(seq('(',sqrt(sub('A','ik')),'−',sqrt(sub('A','i,med')),')'),'2'),seq('2',sup('σ','2'))),'],   ',sub('α','ik'),'=',mr('clip',True),'(',frac(sub('r','ik'),seq(lim('max','j'),sub('r','ij'))),',0.25,1)'),5)
para('式（5）中的面积以像素为单位，$sigma$取100；若最大原始分数数值下溢为0，实现将该ID所有投票权重设为1。这里的$alpha_ik$仅用于式（3）。最终IPR禁用了置信图生成阶段额外的逐视角可见性乘子$v_it$，但未禁用上述面积投票，也未移除基础置信图中的极端面积降权。因此，“去掉可见性”只指训练权重中的这一附加乘子。')
heading('3.5 大范围区域风险与分段对象权重',2)
para('大面积背景或多个对象可能被归入同一个ID。本文使用连通分量的相对面积作为风险代理，而不将其视为真实混合区域的判定器。对该ID全部满足500像素门槛的分量计算面积比例，不受最多40个裁剪的限制。令$t_ik$为分量所在视角，定义：')
equation(seq(sub('a','ik'),'=',frac(seq('|',sub('R','ik'),'|'),seq('|',sub('Ω',sub('t','ik')),'|')),',   ',sub('g','i'),'=𝟙[',sub(mr('P90',True),'k'),'(',sub('a','ik'),') ≥ 0.35 ∨ ',lim('max','k'),sub('a','ik'),' ≥ 0.45]'),6)
para('其中P90为这些分量面积比例的90%分位数，𝟙为条件成立时取1的指示函数。原实现还含中位数比例不低于0.55的分支，但在当前最大值阈值0.45下属于冗余条件，式（6）为等价有效规则。大范围区域先被标记并跳过CLIP裁剪编码，其他ID才计算语义置信度。对象权重采用互斥的优先级规则：')
cases=elem('m:m',elem('m:mPr'))
for left,right in [('0.35',seq(sub('g','i'),'=1')),('0.65',seq(sub('g','i'),'=0 ∧ (',sub('q','i'),'<0.45 ∨ ',sub('m','i'),'<0.08)')),('1',mr('otherwise',True))]:
    cases.append(elem('m:mr',elem('m:e',seq(left)),elem('m:e',seq(right))))
bracket=elem('m:d',elem('m:dPr',elem('m:begChr',**{'m:val':'{'}),elem('m:endChr',**{'m:val':''})),elem('m:e',cases))
equation(seq(sub('s','i'),'=',bracket),7).paragraph_format.keep_with_next=True
para('式（7）不是语义权重与大区域权重的重复相乘。高风险大区域直接取0.35；其余ID根据语义集中程度取0.65或1。真实大对象也可能触发面积规则，这构成方法的适用边界。')
heading('3.6 像素置信图与身份损失',2)
para('令$h_t$在当前像素属于已收集的有效分量时取其ID的$s_i$，其余位置取1。最终像素权重为：')
equation(seq(sub('w','t'),'(x)=',mr('clip',True),'(',sup(sub('c','t'),'0'),'(x)',sub('h','t'),'(x),0.05,1)'),8)
para('训练时令$ell_t$为渲染身份分布与清理标签之间的逐像素交叉熵。沿用动态过滤阈值：第$n$次迭代将阈值设为当前图像平均交叉熵的(2−n/30000)倍；有效掩码$o_t$在交叉熵不高于阈值时取1，否则取0。$C$表示身份分类器的输出通道数，区别于语义候选词表大小。加权身份损失写为：')
equation(seq(sup(sub('L',mr('cls',True)),mr('IPR',True)),'=',frac(summ(seq('x∈',sub('Ω','t')),seq(sub('o','t'),'(x)',sub('w','t'),'(x)',sub('ℓ','t'),'(x)')),seq(mr('log',True),' C · ',mr('max',True),'{',summ(seq('x∈',sub('Ω','t')),seq(sub('o','t'),'(x)',sub('w','t'),'(x)')),',1}'))),9)
para('分母按有效置信质量归一化，下限为1。所有受控无置信图基线使用同一归一化规则，并令$w_t$恒为1。该约定与早期代码在过滤后仍对全图像素直接取均值的做法不同，因而本文将对照称为“ILGS受控基线”，不将其等同于原论文的逐项复现。')
equation(seq(sup(sub('L',mr('total',True)),mr('IPR',True)),'=',sub('L',mr('ILGS',True)),'[',sub('L',mr('cls',True)),' ← ',sup(sub('L',mr('cls',True)),mr('IPR',True)),']'),10)
para('式（10）表示以式（9）替换身份分类项。RGB重建、语言特征监督和三维身份感知语义一致性沿用共同代码及调度。置信图通过改变二维身份监督间接影响下游身份与语言表示，并未直接重写三维语义一致性损失。')

heading('4 实验与分析')
heading('4.1 数据集与受控评测协议',2)
para('主要实验使用LERF中的Figurines、Ramen和Teatime三个场景。每个设置分别使用训练种子0、1、2，并在相同图像划分、输入分辨率、语言特征及自编码器配置下训练30000次迭代。实验使用两张NVIDIA RTX 3090，每个训练任务使用一张GPU。IPR在LERF消融结果上确定最终配置，再固定该配置开展3D-OVS外部验证；因此LERF结果兼具方法选择与评估用途，不能视为独立于配置选择的测试。')
para('所有比较统一采用seed_ref查询掩码生成：优先选择得分最高的非零身份作为种子，在非环绕的身份邻接图上扩展，对正种子相似度使用相对种子得分的3%差值阈值，并排除背景ID 0。该公共解码器与旧official模式在种子、参考相似度和阈值等方面不同。本文将其视为固定评测条件，不将解码器变化带来的收益计入IPR贡献，也不混用两种模式的表格。')
para('mIoU和mBIoU首先按类别汇总各评测视图，再对场景内类别等权平均。边界带宽采用图像对角线的2%，最少1像素；预测尺寸与GT不一致时使用最近邻插值。每个场景报告三种子的均值和样本标准差，跨场景平均为三个场景均值的等权平均，不表示像素加权汇总。配对胜率按同场景同种子计算，不能替代显著性检验。')
para('候选词表复用各场景的评测查询文本，未使用真实像素分割标注生成IPR权重。当前预处理脚本遍历场景图像及对应伪标签，包含留出视图的无标注图像，三维训练则按训练划分读取视图。因此应将实验表述为场景词表已知、无标注多视图可参与预处理的设定，不能据此声称未见词汇或严格只使用训练视图的归纳泛化。')
heading('4.2 实现与参数',2)
para('局部清理与基础置信度使用第3.2节固定参数。语义阶段使用OpenCLIP ViT-B/16的laion2b_s34b_b88k权重，分量最小面积500、最大裁剪数40、外扩12像素、softmax缩放35。面积投票尺度100、最低投票权重0.25；语义阈值为0.45与0.08；低语义权重0.65，大区域权重0.35，最终置信度截断下限0.05。大区域门槛为P90比例0.35或最大比例0.45。最终配置对应实验目录中的no_visibility，正文统一称为IPR。')
para('不同阶段的面积统计量不能混淆：基础置信图采用ID总面积，语义投票采用选中裁剪的分量像素面积，大区域规则采用全部合格分量的归一化面积。像素门槛与面积投票仍具有分辨率依赖性；当前外部实验固定处理分辨率，尚未验证跨分辨率不变性。')
heading('4.3 LERF主结果',2)
para('表1给出统一协议下的ILGS受控基线与IPR结果。IPR在三个场景的种子均值上均提高两项指标，跨场景平均mIoU和mBIoU分别增加3.74和3.21个百分点。Figurines的mIoU样本标准差由10.69降至0.56，Ramen由2.12降至0.22，Teatime由5.76降至2.94，说明当前三种子下的波动有所减小。')
caption('表1  LERF三种子比较（均值 ± 样本标准差，%；每格为mIoU / mBIoU）')
table(['场景','ILGS受控基线','IPR'],[
['Figurines','66.57 ± 10.69 / 64.31 ± 10.58','71.74 ± 0.56 / 69.12 ± 0.44'],
['Ramen','84.88 ± 2.12 / 76.65 ± 0.69','88.33 ± 0.22 / 78.90 ± 0.40'],
['Teatime','76.22 ± 5.76 / 72.78 ± 5.13','78.82 ± 2.94 / 75.35 ± 3.35'],
['场景平均','75.89 / 71.25','79.63 / 74.46']], [2.5,6.075,6.075])
para('在九组同场景同种子配对中，IPR的mIoU胜出8组、mBIoU胜出7组；对应配对差值中位数为2.43和1.92个百分点。Figurines的seed 2仍低于受控基线，因此改进并非每次运行都成立。样本仅覆盖三个场景和三个种子，本文据此报告受控条件下的改善，不作普遍稳定性或统计显著性的推断。')
heading('4.4 局部清理与可靠性消融',2)
para('表2先比较局部清理与基础置信度，再比较附加可靠性项。所有附加配置均建立在相同清理标签和基础置信图上；“仅语义”与“仅大区域”中的“仅”只针对附加权重，不表示移除了基础置信图。S表示语义降权，G表示大区域降权，V表示训练置信图中的逐视角可见性乘子。语义投票内部的面积权重在相关配置中保留。')
para('各附加权重消融复用同一份对象语义与大区域标记。关闭G时，已被标记的大区域不补算语义评分，而是保留基础置信度（若启用V，再叠加V）。因此“基础置信度＋S”检验的是当前门控结构内的语义分支，不能解释为不经大区域筛选、对全部ID独立计算的纯语义方法。默认忽略ID集合为空，背景ID也可进入离线风险估计；查询阶段排除背景ID是另一独立操作。')
caption('表2  三个LERF场景与三个种子的受控消融（场景平均，%）')
table(['设置','mIoU','mBIoU'],[
['ILGS受控基线','75.89','71.25'],
['仅局部清理','74.26','70.18'],
['局部清理＋基础置信度','77.71','73.14'],
['基础置信度＋S','75.74','71.64'],
['基础置信度＋G','78.43','73.78'],
['基础置信度＋S＋G（IPR）','79.63','74.46'],
['基础置信度＋S＋G＋V','77.34','72.14'],
['基础置信度＋G＋V','79.40','74.10'],
['基础置信度＋S＋V','77.87','73.42']], [7.15,3.75,3.75])
para('局部清理单独使用时低于基线，说明重新分配小区域并不保证提高下游质量。引入基础置信度后，相对清理配置增加3.45 mIoU和2.96 mBIoU个百分点；最终IPR相对“清理＋基础置信度”再增加1.92和1.32个百分点，并在两项指标上均胜出7/9组配对。该比较支持对监督贡献进行分级建模，而不是仅依赖标签形态修正。')
para('仅大区域附加降权达到78.43 / 73.78，语义与大区域联合达到79.63 / 74.46，联合配置平均增加1.20 / 0.68个百分点。但该语义增量只在5/9组mIoU和6/9组mBIoU配对中为正：Teatime中联合配置的mIoU为78.82，低于仅大区域的80.20。因此，语义可靠性是当前平均结果中的辅助项，不能声称对所有场景均有增益。大区域比例同样只是风险代理，不能由这些结果推断其识别的每个区域都包含噪声。')
para('额外加入V后，平均指标降至77.34 / 72.14，比最终IPR低2.29 / 2.32个百分点。投影面积随视角和遮挡变化，附加乘子可能压低仍有价值的监督，但这一机制尚需有效权重分布和受控噪声实验验证。本文据现有消融锁定不含该训练乘子的配置，而不是宣称所有面积线索无效。')
heading('4.5 3D-OVS外部验证进展',2)
para('3D-OVS [19]包含Bed、Bench、Blue Sofa、Covered Desk、Lawn、Office Desk、Room、Snacks、Sofa和Table十个场景，每个场景使用五个带标注视图评测。两种方法共享相同DEVA伪标签、图像划分、语言自编码器和seed_ref解码器。DEVA使用短边480像素的图像副本生成轨迹，标签通过最近邻插值恢复到原图尺寸；语言特征与高斯训练采用原图宽高各缩小4倍的分辨率。')
para('当前适配对SIMPLE_RADIAL相机使用焦距对应的针孔近似，未显式校正径向畸变；基线与IPR采用相同近似。此结果属于当前适配协议，不能直接与不同相机处理或分辨率下的文献数值进行排名。外部阶段沿用在LERF确定的配置，不按各3D-OVS场景分别调整阈值。')
caption('表3  Bed先导实验（seed 0；非十场景汇总，%）')
table(['方法','mIoU','mBIoU'],[['ILGS受控基线','97.58','92.39'],['IPR','97.59','93.05'],['差值','+0.01','+0.66']],[7.15,3.75,3.75])
para('截至本稿修订时，已得到Bed的单种子对照结果，其余九个场景已完成伪标签与语言特征准备，完整训练及评测汇总尚待补充。Bed的区域IoU接近饱和；camera类别的边界IoU由86.81%提高至91.24%，但hand类别由91.44%降至90.62%。该先导结果提示局部边界变化，尚不能证明十场景泛化或跨种子稳定收益。')
heading('4.6 局限性与证据边界',2)
para('首先，语义可靠性依赖给定场景词表，未涵盖的对象或多个相似候选可能产生较低置信度。候选词表与评测查询重合，以及无标注留出图像可参与预处理，限定了当前结果的适用协议。独立候选词表和严格训练视图预处理需要单独验证。其次，面积规则可能误伤真实的大对象；高置信度的错误身份也可能被保留，完全遗漏的对象无法通过降权恢复。')
para('再次，最终配置是在LERF消融中选择的，三种子结果描述了当前观察到的波动，不能替代独立验证。已有SAM2 [14]和DINOv3 [16]相关探索使用过不同标签及评测路径，未纳入本次受控主表，不能用于判定基础模型强弱。下一阶段需要完成固定配置的3D-OVS比较，并补充受控ID合并、背景泄漏、缺失区域及权重分布分析，区分模型选择、损失缩放与噪声抑制的影响。预处理时间和存储开销尚未形成完整统计。')

heading('5 结论')
para('本文研究ILGS既有身份空间中的伪标签可靠性问题。IPR通过局部清理、面积加权语义证据汇总及大范围区域风险估计生成像素置信图，并作用于二维身份分类损失。最终配置省去额外的逐视角训练可见性乘子，保留原有ID词表和三维身份感知学习结构。')
para('在统一损失归一化和查询解码的三个LERF场景三种子实验中，IPR取得79.63% mIoU与74.46% mBIoU，相对受控基线提高3.74和3.21个百分点。消融说明局部清理本身不足以保证收益，基础置信度及区域风险降权对总体结果有贡献，而语义增量具有场景依赖性。当前结论限于所述场景词表与多视图预处理协议，外部数据集完整结果及直接抗噪证据仍待补充。')

heading('参考文献')
for ref in refs:
    if ref.startswith('[12]'): ref='[12] Zhu R., Qiu S., Liu Z., et al. Rethinking End-to-End 2D to 3D Scene Segmentation in Gaussian Splatting. CVPR, 2025: 3656–3665.'
    if ref.startswith('[18]'): ref='[18] Park H., Cho D. Consistent Scene Understanding in 3D Gaussian Splatting via Multi-Cue Mask Refinement. ICPR, 2026. arXiv:2607.01708.'
    p=para(ref,size=9.5,indent=False);p.paragraph_format.space_after=Pt(3)
para('[19] Liu K., Zhan F., Zhang J., et al. Weakly Supervised 3D Open-vocabulary Segmentation. Advances in Neural Information Processing Systems, 2023.',size=9.5,indent=False)

for section in doc.sections:
    for p in section.header.paragraphs:
        if p.text: p.clear()
    foot=section.footer.paragraphs[0];foot.clear();foot.alignment=WD_ALIGN_PARAGRAPH.CENTER
    field=elem('w:fldSimple',**{'w:instr':'PAGE'});foot._p.append(field)
doc.core_properties.title='面向ILGS的身份空间保持式伪标签可靠性建模'
doc.core_properties.subject='统一受控协议下的方法与实验修订'
output=OUT/'ILGS-IPR_中文_受控实验与方法修订_20260929.docx'
doc.save(output)
print(output)
print('paragraphs',len(doc.paragraphs),'tables',len(doc.tables),'math',len(doc._element.xpath('.//m:oMath')))
