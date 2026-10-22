from copy import deepcopy
from pathlib import Path
from zipfile import ZipFile

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


SOURCE = Path(r"D:\wlxy\ILGS-IPR_中文_方法一致性修订_v5.docx")
OUTPUT = Path(r"D:\wlxy\ILGS-IPR_中文_方法协议修订_v6.docx")
WORK = Path(r"D:\PycharmProjects\PythonProject\ILGS_release-main\.codex_doc_assets_v6")
PIPELINE = WORK / "pipeline_corrected.png"
ABLATION = WORK / "ablation_preliminary.png"


def el(tag):
    return OxmlElement(tag)


def math_run(text):
    run = el("m:r")
    node = el("m:t")
    node.text = text
    run.append(node)
    return run


def math_container(tag, items):
    node = el(tag)
    for item in items:
        node.append(math_run(item) if isinstance(item, str) else item)
    return node


def sub(base, subscript):
    node = el("m:sSub")
    node.append(math_container("m:e", [base]))
    node.append(math_container("m:sub", [subscript]))
    return node


def sup(base, superscript):
    node = el("m:sSup")
    base_items = base if isinstance(base, list) else [base]
    node.append(math_container("m:e", base_items))
    node.append(math_container("m:sup", [superscript]))
    return node


def subsup(base, subscript, superscript):
    node = el("m:sSubSup")
    node.append(math_container("m:e", [base]))
    node.append(math_container("m:sub", [subscript]))
    node.append(math_container("m:sup", [superscript]))
    return node


def frac(numerator, denominator):
    node = el("m:f")
    node.append(math_container("m:num", numerator))
    node.append(math_container("m:den", denominator))
    return node


def clear_paragraph(paragraph):
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def set_run_font(run, ascii_name="Times New Roman", east_asia="宋体", size=10.5, bold=None):
    run.font.name = ascii_name
    run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), ascii_name)
    rfonts.set(qn("w:hAnsi"), ascii_name)
    rfonts.set(qn("w:eastAsia"), east_asia)


def add_body_run(paragraph, text, bold=False):
    run = paragraph.add_run(text)
    set_run_font(run, bold=bold)
    return run


def add_code_run(paragraph, text):
    run = paragraph.add_run(text)
    set_run_font(run, ascii_name="Consolas", east_asia="宋体", size=9.2)
    return run


def append_inline_math(paragraph, items):
    math = el("m:oMath")
    for item in items:
        math.append(math_run(item) if isinstance(item, str) else item)
    paragraph._p.append(math)


def set_pieces(paragraph, pieces):
    clear_paragraph(paragraph)
    for kind, value in pieces:
        if kind == "text":
            add_body_run(paragraph, value)
        elif kind == "bold":
            add_body_run(paragraph, value, bold=True)
        elif kind == "code":
            add_code_run(paragraph, value)
        elif kind == "math":
            append_inline_math(paragraph, value)
        else:
            raise ValueError(kind)


def word_run(text, size=10.5):
    run = el("w:r")
    rpr = el("w:rPr")
    rfonts = el("w:rFonts")
    rfonts.set(qn("w:ascii"), "Times New Roman")
    rfonts.set(qn("w:hAnsi"), "Times New Roman")
    rfonts.set(qn("w:eastAsia"), "宋体")
    rpr.append(rfonts)
    sz = el("w:sz")
    sz.set(qn("w:val"), str(int(size * 2)))
    rpr.append(sz)
    run.append(rpr)
    node = el("w:t")
    node.text = text
    run.append(node)
    return run


def set_formula(paragraph, items, number):
    clear_paragraph(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = None
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.keep_together = True
    ppr = paragraph._p.get_or_add_pPr()
    tabs = ppr.find(qn("w:tabs"))
    if tabs is not None:
        ppr.remove(tabs)
    tabs = el("w:tabs")
    tab = el("w:tab")
    tab.set(qn("w:val"), "right")
    tab.set(qn("w:pos"), str(int(6.2 * 1440)))
    tabs.append(tab)
    ppr.append(tabs)
    math = el("m:oMath")
    for item in items:
        math.append(math_run(item) if isinstance(item, str) else item)
    paragraph._p.append(math)
    tab_run = el("w:r")
    tab_run.append(el("w:tab"))
    paragraph._p.append(tab_run)
    paragraph._p.append(word_run(f"({number})"))


def insert_after(paragraph, text="", style=None):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    from docx.text.paragraph import Paragraph

    result = Paragraph(new_p, paragraph._parent)
    if style:
        result.style = style
    if text:
        result.add_run(text)
    return result


def replace_picture(paragraph, image_path, width=6.55, alt=""):
    clear_paragraph(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(4)
    picture = paragraph.add_run().add_picture(str(image_path), width=Inches(width))
    picture._inline.docPr.set("descr", alt)


def font(size, bold=False):
    name = "arialbd.ttf" if bold else "arial.ttf"
    candidate = Path(r"C:\Windows\Fonts") / name
    if candidate.exists():
        return ImageFont.truetype(str(candidate), size)
    return ImageFont.load_default()


def centered(draw, box, text, fnt, fill, spacing=4):
    x0, y0, x1, y1 = box
    bounds = draw.multiline_textbbox((0, 0), text, font=fnt, spacing=spacing, align="center")
    width = bounds[2] - bounds[0]
    height = bounds[3] - bounds[1]
    draw.multiline_text(
        ((x0 + x1 - width) / 2, (y0 + y1 - height) / 2),
        text,
        font=fnt,
        fill=fill,
        spacing=spacing,
        align="center",
    )


def extract_media():
    WORK.mkdir(parents=True, exist_ok=True)
    extracted = []
    with ZipFile(SOURCE) as zf:
        for name in zf.namelist():
            if not name.startswith("word/media/"):
                continue
            path = WORK / Path(name).name
            path.write_bytes(zf.read(name))
            try:
                with Image.open(path) as image:
                    extracted.append((path, image.size))
            except Exception:
                pass
    return extracted


def build_corrected_figures(media):
    pipeline_src = next(path for path, size in media if size == (1920, 1080))
    chart_src = next(path for path, size in media if size == (1256, 601))

    image = Image.open(pipeline_src).convert("RGB")
    draw = ImageDraw.Draw(image)
    panel1 = (244, 412, 424, 463)
    draw.rectangle(panel1, fill=(243, 247, 252))
    centered(draw, panel1, "raw object IDs  Yt", font(20, True), (53, 82, 119))

    header = (470, 195, 1032, 275)
    draw.rectangle(header, fill=(255, 247, 238))
    draw.text((498, 211), "2  Anchor-space-preserving IPR", font=font(27, True), fill=(207, 105, 0))
    draw.text((500, 246), "local cleanup + semantic-area reliability", font=font(18), fill=(144, 83, 25))

    middle = (677, 271, 817, 405)
    draw.rectangle(middle, fill=(255, 247, 238))
    draw.rounded_rectangle((683, 277, 811, 399), radius=12, fill=(255, 255, 255), outline=(238, 143, 0), width=3)
    centered(draw, (690, 288, 804, 388), "area\nvisibility", font(22, True), (164, 94, 0))

    output_box = (548, 568, 948, 636)
    draw.rounded_rectangle(output_box, radius=12, fill=(239, 255, 245), outline=(78, 220, 135), width=3)
    centered(draw, output_box, "cleaned labels  Y~t  +  confidence map  wt", font(21, True), (0, 119, 53))

    total = (1592, 768, 1857, 925)
    draw.rounded_rectangle(total, radius=12, fill=(237, 245, 255), outline=(78, 140, 255), width=3)
    centered(
        draw,
        total,
        "Ltotal^IPR = LILGS\nwith Lcls replaced by\nLcls^IPR",
        font(21, True),
        (35, 83, 205),
        spacing=7,
    )

    footer = (280, 953, 1660, 1025)
    draw.rectangle(footer, fill="white")
    centered(
        draw,
        footer,
        "No new IDs or trajectories are introduced; local cleanup reuses existing IDs and reliability weights the identity loss",
        font(22, True),
        (25, 38, 55),
    )
    image.save(PIPELINE, quality=95)

    chart = Image.open(chart_src).convert("RGB")
    draw = ImageDraw.Draw(chart)
    draw.rectangle((75, 538, 1248, 600), fill="white")
    labels = [
        (182, "Official"),
        (409, "Area\nvariant*"),
        (634, "Semantic\nvariant*"),
        (858, "Cleanup\nno confidence*"),
        (1082, "Full IPR"),
    ]
    for x, label in labels:
        centered(draw, (x - 82, 538, x + 82, 600), label, font(17), (20, 20, 20), spacing=1)
    draw.rectangle((930, 15, 1240, 43), fill="white")
    draw.text((940, 18), "* preliminary historical runs", font=font(14), fill=(85, 85, 85))
    chart.save(ABLATION, quality=95)


def set_text(paragraph, text):
    clear_paragraph(paragraph)
    add_body_run(paragraph, text)


def main():
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    media = extract_media()
    build_corrected_figures(media)
    doc = Document(str(SOURCE))
    p = doc.paragraphs
    if len(p) != 108:
        raise RuntimeError(f"Unexpected paragraph count: {len(p)}")

    set_text(p[0], "面向ILGS噪声身份伪标签的锚点空间保持式语义—面积可靠性建模")
    p[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in p[0].runs:
        set_run_font(run, east_asia="黑体", size=16)

    set_text(
        p[1],
        "摘要：开放词汇三维语义分割根据自然语言查询定位三维场景中的目标区域。ILGS通过身份感知语义一致性将对象身份监督引入语言高斯学习，但其监督依赖自动生成的object_mask及跨视角对象ID。当伪标签包含背景泄漏、多对象合并或跨视角身份漂移时，身份损失可能强化错误关系。本文提出锚点空间保持式语义—面积身份伪标签可靠性建模方法（Semantic-Area Identity-aware Pseudo-label Refinement，Semantic-Area IPR）。该方法首先在不创建新ID或重新建立轨迹的前提下，对孤立小连通域进行局部清理；随后依据对象裁剪的语义置信度、跨视角面积可见性以及大范围混合区域风险生成像素级置信图，并以该置信图加权ILGS二维身份分类损失。三个LERF场景上的初步单次运行结果显示，相较修正后的官方复现基线，Semantic-Area IPR将平均mIoU由72.69%提升至80.50%，将平均mBIoU由68.59%提升至74.75%。figurines场景上的探索实验还表明，更强的单帧分割或跨视角匹配模块若不能稳定对齐既有身份空间，并不必然获得更高的下游分割指标。"
    )

    set_text(
        p[8],
        "因此，本文研究的问题可以表述为：给定自动生成的object_mask及其跨视角对象ID，在不增加人工标注、不重新建立全局对象轨迹且不引入新ID的条件下，如何识别不可靠的身份监督，并降低其在ILGS优化中的影响？本文关注的不是重新构建一套身份对应关系，而是在既有ID空间中进行保守的局部清理与可靠性估计。"
    )
    set_text(
        p[9],
        "针对这一问题，本文提出Semantic-Area IPR。该方法由两个连续阶段组成：首先，将孤立的小连通域重新分配给其邻域中已有的主导ID，同时对边界和异常面积区域降低初始置信度；其次，从语义可靠性、面积可见性和大范围混合区域风险三个方面估计对象ID的可靠程度。最终输出不包含新ID，也不重新建立跨视角轨迹，而是由清理后的身份标签与像素级置信图共同构成训练监督。"
    )
    set_text(
        p[10],
        "本文的主要贡献包括：（1）将ILGS中的object_mask噪声表述为身份监督的可靠性与错误放大问题，并区分局部像素噪声和跨视角身份空间；（2）提出一种不创建新ID、不重新跟踪的两阶段伪标签处理方法，通过局部清理、语义置信度、面积可见性和混合区域检测产生像素级可靠性；（3）将可靠性显式注入ILGS二维身份分类损失，并通过三个LERF场景上的主结果、预备消融和探索性二维模型实验分析其效果与适用边界。"
    )

    set_text(
        p[22],
        "近期工作开始直接处理二维基础模型伪标签的跨视角可靠性。SAM2Object [11]利用SAM2的视频分割能力和三维几何先验整合多视角掩码；Unified-Lift [12]将噪声标签过滤和面积感知映射纳入二维到三维高斯分割；SeConGS [13]通过跟踪分割掩码建立语义一致监督。TrackRef3D [17]进一步采用track-then-label范式，将对象发现与语义赋值解耦，并使用轨迹感知语义投票。Multi-Cue Mask Refinement [18]则联合语义、几何和结构线索进行掩码合并及跨视角ID匹配。"
    )
    set_text(
        p[23],
        "上述方法主要通过重新跟踪、合并或匹配掩码建立更一致的身份关系。本文研究与其互补的设定：当ILGS已有一套可用但含噪的跨视角ID空间时，不重新构建全局轨迹，而是保守修正局部噪声，并估计各身份监督对训练的可信贡献。因而，本文的技术重点是既有身份空间中的可靠性建模及其向identity loss的传播。"
    )

    replace_picture(
        p[26],
        PIPELINE,
        alt="Two-stage anchor-space-preserving IPR preprocessing and confidence-weighted ILGS identity loss",
    )
    set_text(
        p[27],
        "图1  Semantic-Area IPR在ILGS中的位置。方法先在已有ID空间内执行局部清理，再估计语义权重、面积可见性和混合区域风险；输出为不含新ID的清理标签与像素级置信图。ILGS骨干、官方CE离群点过滤、RGB重建、语言特征监督和推理流程保持不变。"
    )
    set_text(
        p[28],
        "图1展示完整数据路径。原始RGB视图、对象ID、CLIP特征和相机信息作为输入。局部清理只允许把孤立像素重新分配给当前标签空间中已有的邻近ID，不生成新ID，也不重新建立对象轨迹。语义—面积阶段保持清理后的标签不变，仅计算像素级置信图。训练时，ILGS先根据渲染身份logits与清理标签计算逐像素交叉熵，并使用官方动态阈值生成有效像素掩码；随后将该掩码与IPR置信图相乘，得到置信度加权的二维身份分类损失。"
    )

    set_pieces(
        p[30],
        [
            ("text", "设原始训练图像为"), ("math", [sub("I", "t"), ", t=1,…,T"]),
            ("text", "，像素域为"), ("math", [sub("Ω", "t")]),
            ("text", "，原始像素级身份标签为"), ("math", [sub("Y", "t"), "(x)"]),
            ("text", "。记原始标签中出现的ID集合为"), ("math", ["I(", sub("Y", "t"), ")"]),
            ("text", "，局部清理后的标签为"), ("math", [sup(sub("Y", "t"), "~"), "(x)"]),
            ("text", "，IPR像素级可靠性为"), ("math", [sub("w", "t"), "(x)"]),
            ("text", "。本文不新建ID或重建跨视角轨迹，但允许孤立小区域在已有ID集合内重新归属。"),
        ],
    )
    set_text(
        p[31],
        "式（1）给出锚点空间保持约束：清理后的标签只能使用原始标签中已有的ID，语义—面积阶段不再改变清理标签，只生成取值为0至1的可靠性权重。这里的“保持”指ID词表与跨视角轨迹来源不变，而不是保证每个像素标签完全不变。"
    )
    set_formula(
        p[32],
        ["I(", sup(sub("Y", "t"), "~"), ")⊆I(", sub("Y", "t"), "),    ", subsup("Y", "t", "IPR"), "(x)=", sup(sub("Y", "t"), "~"), "(x),    0≤", sub("w", "t"), "(x)≤1"],
        1,
    )

    p[33].text = "3.3 语义置信度估计"
    p[40].text = "3.4 面积可见性建模"
    p[45].text = "3.5 大范围混合区域抑制"
    p[49].text = "3.6 可靠性融合与置信度加权训练"

    set_text(
        p[34],
        "对于每个对象ID，本文从多个视角提取面积不小于500像素的连通分量，并按面积选取至多40个裁剪区域。裁剪在目标外保留12像素上下文，掩码外区域以目标像素均值填充。给定场景候选语义词表，使用OpenCLIP ViT-B/16分别编码裁剪图像和文本，并计算归一化相似度及缩放后的softmax概率。候选词表只用于估计身份标签的相对语义稳定性，不会写入object ID。"
    )
    set_pieces(
        p[35],
        [
            ("text", "设场景候选词表为"), ("math", [sub("C", "s")]),
            ("text", "，对象ID "), ("math", ["i"]), ("text", "在第"), ("math", ["k"]),
            ("text", "个裁剪上属于类别"), ("math", ["c∈", sub("C", "s")]), ("text", "的概率为"),
            ("math", [subsup("p", "i", "k"), "(c)"]), ("text", "，面积可见性权重为"),
            ("math", [subsup("v", "i", "k")]), ("text", "。累积语义响应"),
            ("math", [sub("S", "i"), "(c)"]), ("text", "、top-1置信度"),
            ("math", [sub("q", "i")]), ("text", "和top-1/top-2间隔"),
            ("math", [sub("m", "i")]), ("text", "由式（2）至式（4）定义。"),
        ],
    )

    set_text(
        p[41],
        "语义置信度难以单独识别背景泄漏、遮挡碎片和异常投影，因此本文进一步使用对象连通分量的跨视角面积分布估计可见性。该权重并不要求各视角面积相同，而是偏向中等且稳定的可见区域：面积过小的碎片和与中位尺度偏差过大的区域获得较低权重。由于当前LERF图像分辨率一致，代码直接使用像素面积；跨分辨率实验应改用归一化面积或相对尺度参数。"
    )
    set_text(
        p[44],
        "式（5）与历史实现一致：未归一化分数同时考虑连通分量面积和平方根面积相对中位面积的偏差，再在同一对象ID内部归一化至[0.25, 1]。因此，本模块更准确地称为面积可见性加权，而不是要求投影面积严格一致。"
    )

    set_pieces(
        p[47],
        [
            ("text", "令对象ID "), ("math", ["i"]), ("text", "在视角"), ("math", ["t"]),
            ("text", "中的归一化面积为"), ("math", [subsup("a", "i", "t")]),
            ("text", "。历史代码还包含中位数阈值，但该条件被最大值阈值完全包含，不会单独改变判定结果。为避免冗余，本文按等价的有效规则使用P90和最大面积比例定义混合区域指示量"),
            ("math", [sub("g", "i")]), ("text", "，如式（6）所示。"),
        ],
    )
    set_formula(
        p[48],
        [
            subsup("a", "i", "t"), "=", frac(["|", subsup("M", "i", "t"), "|"], ["|", sub("Ω", "t"), "|"]),
            ",    ", sub("g", "i"), "=1[P90", sub("t", ""), "(", subsup("a", "i", "t"), ")≥", sub("τ", "p90"), "  ∨  ", sub("max", "t"), "(", subsup("a", "i", "t"), ")≥", sub("τ", "max"), "]",
        ],
        6,
    )

    set_pieces(
        p[50],
        [
            ("text", "对象级语义权重记为"), ("math", [sub("s", "i")]),
            ("text", "。若"), ("math", [sub("g", "i"), "=1"]), ("text", "，则"),
            ("math", [sub("s", "i"), "=0.35"]), ("text", "；否则，当"),
            ("math", [sub("q", "i"), "<0.45"]), ("text", "或"),
            ("math", [sub("m", "i"), "<0.08"]), ("text", "时取"),
            ("math", [sub("s", "i"), "=0.65"]), ("text", "，其余情况取1。设"),
            ("math", [subsup("c", "t", "0"), "(x)"]), ("text", "为局部清理阶段产生的初始置信图，最终像素权重由式（7）给出。"),
        ],
    )
    set_formula(
        p[53],
        [subsup("L", "total", "IPR"), "=", sub("L", "ILGS"), "(", sub("L", "cls"), "←", subsup("L", "cls", "IPR"), ")"],
        9,
    )
    set_text(
        p[54],
        "训练阶段首先计算逐像素交叉熵，并沿用ILGS官方动态阈值产生有效像素掩码。式（8）使用有效掩码与IPR置信权重的乘积计算加权二维身份分类损失，分母下限截断为1，并保留官方log C归一化。式（9）表示实现中的唯一训练改动：以加权身份分类损失替换ILGS原身份分类损失。RGB重建、CLIP特征监督、身份感知三维语义一致性及其启用时刻和间隔均保持官方调度不变，而不是在每次迭代中同时计算全部损失。"
    )

    local_heading = insert_after(p[32], "3.2 保持ID空间的局部清理", "Heading 2")
    local_p1 = insert_after(
        local_heading,
        "完整历史流程以原始object_mask为输入。对于面积小于64像素的孤立连通分量，方法在半径为3像素的外环邻域中统计已有ID；若存在有效邻居，则将该分量重新分配给出现次数最多的邻近ID，否则保留原标签。该操作只复用当前标签空间中的ID，不创建新ID，也不根据语义重新建立跨视角轨迹。",
        "Normal",
    )
    insert_after(
        local_p1,
        "局部清理同时生成初始置信图：小连通分量首先乘以0.25，发生重新分配时再乘以0.5；宽度为2像素的身份边界乘以0.65；当某ID在当前视角的面积低于其跨视角中位数的0.15倍或高于6倍时再乘以0.5。所有权重最终截断至[0.05, 1]。该初始置信图随后作为式（7）中的输入置信度。",
        "Normal",
    )

    set_text(
        p[58],
        "所有结果均基于ILGS官方代码路径复现，并使用相同场景划分、骨干网络、30000次训练迭代和评价脚本。修正后的官方基线只处理特征解码、掩码生成与评价接口的兼容性，不读取IPR置信图。需要明确的是，本文现有历史实验在可靠性估计阶段复用了各LERF场景的评测查询词作为候选词表，因此当前结果属于scene-vocabulary-aware的可靠性研究，不能单独证明对未见词汇的零样本泛化。最终投稿版本需要补充与评测查询无关的自动词表实验。"
    )
    set_text(
        p[60],
        "局部清理阶段采用：小连通分量阈值64、邻域半径3、边界宽度2、边界权重0.65、异常面积下界0.15倍中位数、上界6倍中位数、异常面积权重0.50。语义—面积阶段采用：连通分量阈值500、每个ID最多40个裁剪、上下文宽度12、可见性尺度100、最低可见性权重0.25、语义置信度阈值0.45、语义间隔阈值0.08、低语义权重0.65。有效的大区域规则为P90面积比例不低于0.35或最大面积比例不低于0.45，对应权重0.35。最终置信度下限为0.05，CLIP softmax缩放系数为35。"
    )
    set_text(
        p[62],
        "表1给出当前单次运行的系统级结果。在相同复现设置下，完整Semantic-Area IPR在三个场景上均高于修正后的官方基线，平均mIoU由72.69%提高至80.50%，平均mBIoU由68.59%提高至74.75%。这些结果支持完整处理流程的有效性，但在补充多随机种子和独立候选词表实验前，本文不把单次运行差值解释为统计显著性，也不据此声称超过现有方法。"
    )
    set_text(
        p[65],
        "表2汇总历史实验目录中的预备变体。需要注意，语义和面积模块在当前代码中主要通过置信图影响训练，而语义—面积脚本本身不改写输入标签。因此，不读取置信图的设置不能称为“语义—面积融合但无置信度”，更准确的名称是“标签清理、无置信图”。Only area和Only semantic沿用历史运行名称，其具体数据源和开关仍需在最终受控消融中统一。现阶段，表2可以说明完整IPR系统优于这些历史变体，但不能严格量化每个模块的独立因果贡献。"
    )
    table = doc.tables[1]
    table.cell(0, 0).text = "预备设置"
    table.cell(1, 0).text = "修正后的官方基线"
    table.cell(2, 0).text = "面积变体（历史运行）"
    table.cell(3, 0).text = "语义变体（历史运行）"
    table.cell(4, 0).text = "标签清理（无置信图）"
    table.cell(5, 0).text = "完整Semantic-Area IPR"
    # The source caption forced a new page, which stranded the preceding
    # ablation explanation on an otherwise empty page.
    p[66].paragraph_format.page_break_before = False
    replace_picture(p[67], ABLATION, alt="Preliminary historical variants and full Semantic-Area IPR")
    set_text(p[68], "图2  三个LERF场景上的预备变体结果。带星号的设置沿用历史运行命名，最终模块消融需使用统一代码入口和随机种子重新验证。")

    p[69].text = "4.5 探索性定性结果与可选Full-refine"
    set_text(
        p[70],
        "图3至图5展示IPR训练后再启用Full-refine的组合系统，只用于说明查询级候选修正能够产生的现象，不用于归因训练阶段IPR。teatime中的sheep和spoon handle显示组合系统能够减少部分背景响应并恢复细长目标，但其中既包含训练置信加权的影响，也包含推理候选重排序的影响。"
    )
    set_text(
        p[71],
        "ramen中的glass of water展示了组合系统对透明目标候选区域的修正；pork belly样例中两种方法的单图IoU相同，因此该样例不能作为数值提升证据，只能用于观察响应区域变化。figurines中的red apple有所改善，而old camera进一步下降，说明查询级后处理对候选身份区域选择较为敏感。"
    )
    set_text(
        p[72],
        "Full-refine根据文本相似度、候选身份区域面积、可见性和边界连通性执行重排序、合并与过滤。它不参与训练，也不属于锚点空间保持式可靠性建模的核心组成。由于它可能改变查询与身份区域的最终对应关系，正文主结论以IPR-only数值为准；组合结果作为探索性扩展单独报告。"
    )
    set_text(
        p[74],
        "表3显示，Full-refine平均mIoU和mBIoU分别比IPR-only高0.68和1.91个百分点，但只在teatime取得场景级提升，在figurines和ramen均下降。因此，平均值不能被解释为稳定增益。该模块应保持可选，并在补充类别级结果和失败率统计后再讨论其适用条件。"
    )
    set_text(p[76], "图3  teatime场景的探索性组合结果。最后一列同时包含Semantic-Area IPR训练和Full-refine推理后处理，不能单独归因于IPR。")
    set_text(p[78], "图4  ramen场景的探索性组合结果。glass of water得到改善，而pork belly单图IoU未提高，说明查询级修正具有类别依赖性。")
    set_text(p[80], "图5  figurines场景的探索性组合结果与失败案例。red apple有所改善，但old camera下降，反映候选身份区域选择错误。")
    set_text(
        p[82],
        "本文在figurines场景上探索了SAM2语义掩码、DEVA与SAM2对齐以及DINOv3/SAM2 V2-anchor严格匹配。表4结果仅反映当前实现和单一场景：这些策略均未超过Semantic-Area IPR，但不能据此得出更强二维模型普遍无效的结论。该实验支持的有限结论是，二维边界质量和特征能力必须与ILGS现有身份空间稳定对齐，才能转化为三维训练收益。"
    )
    set_text(
        p[84],
        "SAM2改善单帧边界，DEVA提供跨视角轨迹，DINOv3/V2-anchor执行特征级匹配，但三者都会重新产生或重新对齐身份区域。当前失败案例说明，对齐误差可能削弱identity-aware supervision。相比之下，IPR优先复用已有ID空间并降低可疑监督权重。该观察构成本文方法动机，而不是对上述模型能力的普遍比较。"
    )

    limitations = p[86].insert_paragraph_before("4.7 当前实验协议限制与补充方向")
    limitations.style = "Heading 2"
    lim_p = p[86].insert_paragraph_before(
        "当前证据仍有三项边界。第一，可靠性词表与LERF评测查询重合，因此需要补充自动生成或与评测类别独立的词表实验。第二，现有表格来自单次训练，需要增加多随机种子均值和标准差。第三，验证数据集中在三个LERF场景，最终版本还应加入3D-OVS或其他场景，并通过受控ID合并、背景泄漏和跨视角漂移实验直接测量抗噪能力。上述补充完成前，本文结论限定为当前scene-vocabulary-aware协议下的系统级改进。"
    )
    lim_p.style = "Normal"

    set_text(
        p[87],
        "本文针对ILGS依赖含噪身份伪标签的问题，提出锚点空间保持式Semantic-Area IPR。方法不创建新ID或重新建立全局轨迹，而是在已有ID空间内执行保守的局部清理，并结合语义置信度、面积可见性和大范围混合区域风险生成像素级置信图。训练阶段仅以置信度加权的二维身份分类损失替换ILGS原身份分类损失，其他网络结构、损失调度和推理流程保持不变。"
    )
    set_text(
        p[88],
        "三个LERF场景上的初步单次运行结果显示，完整IPR系统相较修正后的官方复现基线平均mIoU提高7.81个百分点，平均mBIoU提高6.16个百分点。历史变体和二维模型探索说明，监督可靠性及其与既有身份空间的兼容性值得单独建模；但模块级因果贡献、未见词汇泛化和跨数据集稳定性仍需通过受控消融、多随机种子和独立词表实验验证。Full-refine只作为查询级可选扩展，不参与核心方法结论。"
    )

    ref17 = doc.add_paragraph(
        "[17] Tan Y., Zhang R., Zhang H., Li A., Tan X. TrackRef3D: Multi-View Consistent Track-then-Label for Open-World Referring Segmentation in 3D Gaussian Splatting. arXiv:2605.26576, 2026."
    )
    ref18 = doc.add_paragraph(
        "[18] Park H., Cho D. Consistent Scene Understanding in 3D Gaussian Splatting via Multi-Cue Mask Refinement. In: International Conference on Pattern Recognition (ICPR), 2026."
    )

    # Keep body typography and compact the references so the final page is not sparse.
    for paragraph in doc.paragraphs:
        if paragraph.style.name == "Normal":
            for run in paragraph.runs:
                if not run.font.name or run.font.name != "Consolas":
                    set_run_font(run)
    ref_start = next(i for i, paragraph in enumerate(doc.paragraphs) if paragraph.text.strip() == "参考文献")
    for paragraph in doc.paragraphs[ref_start + 1 :]:
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.paragraph_format.line_spacing = 1.0
        for run in paragraph.runs:
            set_run_font(run, size=9.0)

    for paragraph in (local_p1, lim_p, ref17, ref18):
        for run in paragraph.runs:
            set_run_font(run)

    doc.save(str(OUTPUT))
    print(f"saved={OUTPUT}")
    print(f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)}")


if __name__ == "__main__":
    main()
