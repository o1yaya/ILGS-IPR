from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


DOCX_PATH = Path(r"D:\wlxy\ILGS-IPR.docx")


def el(tag):
    return OxmlElement(tag)


def mr(text):
    r = el("m:r")
    t = el("m:t")
    t.text = text
    r.append(t)
    return r


def add_seq(parent, items):
    for item in items:
        if isinstance(item, str):
            parent.append(mr(item))
        else:
            parent.append(item)
    return parent


def container(tag, items):
    node = el(tag)
    add_seq(node, items)
    return node


def sub(base, subscript):
    node = el("m:sSub")
    node.append(container("m:e", [base]))
    node.append(container("m:sub", [subscript]))
    return node


def sup(base_items, superscript):
    node = el("m:sSup")
    node.append(container("m:e", base_items if isinstance(base_items, list) else [base_items]))
    node.append(container("m:sup", [superscript]))
    return node


def subsup(base, subscript, superscript):
    node = el("m:sSubSup")
    node.append(container("m:e", [base]))
    node.append(container("m:sub", [subscript]))
    node.append(container("m:sup", [superscript]))
    return node


def frac(num_items, den_items):
    node = el("m:f")
    node.append(container("m:num", num_items))
    node.append(container("m:den", den_items))
    return node


def clear_paragraph(paragraph):
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def set_math_paragraph(paragraph, formula_items):
    clear_paragraph(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = None
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(3)

    math_para = el("m:oMathPara")
    math = el("m:oMath")
    add_seq(math, formula_items)
    math_para.append(math)
    paragraph._p.append(math_para)


def set_number_paragraph(paragraph, number):
    clear_paragraph(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.paragraph_format.first_line_indent = None
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(f"({number})")
    run.font.name = "Times New Roman"
    run.font.size = Pt(10.5)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")


def remove_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is not None:
        tbl_pr.remove(borders)
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = OxmlElement(f"w:{side}")
        tag.set(qn("w:val"), "nil")
        borders.append(tag)
    tbl_pr.append(borders)


def insert_formula_table_after(doc, paragraph, formula_items, number):
    table = doc.add_table(rows=1, cols=3)
    table.autofit = False
    remove_table_borders(table)

    widths = [Inches(0.65), Inches(4.65), Inches(0.65)]
    for cell, width in zip(table.rows[0].cells, widths):
        cell.width = width

    left_p = table.rows[0].cells[0].paragraphs[0]
    left_p.alignment = WD_ALIGN_PARAGRAPH.LEFT

    formula_p = table.rows[0].cells[1].paragraphs[0]
    set_math_paragraph(formula_p, formula_items)

    number_p = table.rows[0].cells[2].paragraphs[0]
    set_number_paragraph(number_p, number)

    tbl = table._tbl
    paragraph._p.addnext(tbl)
    paragraph._element.getparent().remove(paragraph._element)


def main():
    doc = Document(str(DOCX_PATH))

    formulas = {
        33: (
            [sub("Y", "t"), "(x)=i,    ", sub("w", "t"), "(x)∈[0,1]"],
            1,
        ),
        37: (
            [
                sub("S", "i,c"),
                "=",
                sub("∑", "k"),
                sub("α", "i,k"),
                sub("p", "i,k,c"),
            ],
            2,
        ),
        38: (
            [
                sub("q", "i"),
                "=",
                frac(
                    [sub("max", "c"), sub("S", "i,c")],
                    [sub("∑", "c"), sub("S", "i,c"), "+ε"],
                ),
            ],
            3,
        ),
        39: (
            [
                sub("m", "i"),
                "=",
                frac(
                    [sub("S", "i,c₁"), "−", sub("S", "i,c₂")],
                    [sub("∑", "c"), sub("S", "i,c"), "+ε"],
                ),
            ],
            4,
        ),
        44: (
            [
                subsup("v", "i", "t"),
                "=",
                subsup("A", "i", "t"),
                "·exp(",
                frac(
                    [
                        "−",
                        sup(
                            ["(", subsup("A", "i", "t"), "−", sub("A", "i,med"), ")"],
                            "2",
                        ),
                    ],
                    ["2", sup("σ", "2")],
                ),
                ")",
            ],
            5,
        ),
        51: (
            [
                sub("L′", "id"),
                "=",
                frac(
                    [
                        sub("∑", "x"),
                        sub("w", "t"),
                        "(x)CE(",
                        sub("p", "t"),
                        "(x),",
                        sub("Y", "t"),
                        "(x))",
                    ],
                    [sub("∑", "x"), sub("w", "t"), "(x)+ε"],
                ),
            ],
            6,
        ),
    }

    # Process from the bottom so replacing paragraphs does not shift indices
    # that still need to be edited.
    for idx, (items, number) in sorted(formulas.items(), reverse=True):
        insert_formula_table_after(doc, doc.paragraphs[idx], items, number)

    doc.save(str(DOCX_PATH))
    print(f"saved {DOCX_PATH}")


if __name__ == "__main__":
    main()
