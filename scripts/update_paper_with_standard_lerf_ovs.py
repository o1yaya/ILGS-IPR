from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.table import Table


SOURCE = Path(r"D:\wlxy\ILGS-IPR_中文1010.docx")
OUTPUT = Path(r"D:\wlxy\ILGS-IPR_中文1010_标准四场景LERFOVS补充版.docx")


def replace_paragraph(paragraph, text):
    run_properties = None
    if paragraph.runs and paragraph.runs[0]._r.rPr is not None:
        run_properties = deepcopy(paragraph.runs[0]._r.rPr)

    for child in list(paragraph._p):
        if child is not paragraph._p.pPr:
            paragraph._p.remove(child)

    run = paragraph.add_run(text)
    if run_properties is not None:
        run._r.insert(0, run_properties)
    return paragraph


def insert_paragraph_after(document, anchor, reference, text):
    paragraph = document.add_paragraph()
    if paragraph._p.pPr is not None:
        paragraph._p.remove(paragraph._p.pPr)
    if reference._p.pPr is not None:
        paragraph._p.insert(0, deepcopy(reference._p.pPr))

    run = paragraph.add_run(text)
    if reference.runs and reference.runs[0]._r.rPr is not None:
        run._r.insert(0, deepcopy(reference.runs[0]._r.rPr))

    anchor_element = anchor._p if hasattr(anchor, "_p") else anchor._tbl
    anchor_element.addnext(paragraph._p)
    return paragraph


def set_cell_text(cell, text):
    paragraph = cell.paragraphs[0]
    run_properties = None
    if paragraph.runs and paragraph.runs[0]._r.rPr is not None:
        run_properties = deepcopy(paragraph.runs[0]._r.rPr)

    for extra in cell.paragraphs[1:]:
        cell._tc.remove(extra._p)
    for child in list(paragraph._p):
        if child is not paragraph._p.pPr:
            paragraph._p.remove(child)

    run = paragraph.add_run(text)
    if run_properties is not None:
        run._r.insert(0, run_properties)


def main():
    document = Document(SOURCE)
    paragraphs = document.paragraphs

    replacements = {
        1: (
            "摘要：身份感知语言高斯依赖跨视角对象伪标签构建语义一致性，但自动生成的身份区域可能包含边界噪声、"
            "背景泄漏或多对象混合。仅依赖训练交叉熵难以区分可信身份与可被模型拟合的错误标签。本文提出身份空间"
            "保持式伪标签可靠性建模方法（Identity-aware Pseudo-label Refinement，IPR），在已有对象ID集合内进行局部"
            "清理，利用跨视角裁剪语义响应和连通区域面积分布估计监督可靠性，并通过像素级置信图加权二维身份分类损失。"
            "在统一损失归一化与查询解码协议下，三个LERF场景的三种子受控实验中，IPR将平均mIoU由75.89%提高至"
            "79.63%，mBIoU由71.25%提高至74.46%。在仅使用训练视图且不使用评测查询构造权重的标准四场景LERF-OVS"
            "评测中，三种子平均mIoU由50.51%提高至52.02%，mAcc@0.25由58.41%提高至60.23%，增益主要来自Waldo "
            "Kitchen。在固定参数的3D-OVS十场景三种子实验中，平均mIoU由82.01%提高至87.07%，mBIoU由75.69%提高至"
            "81.22%，且两个指标的场景均值均在9/10个场景提高；Covered Desk仍出现退化。结果表明，IPR能够在不改变身份"
            "集合的条件下缓解部分错误身份监督，但收益具有场景依赖性。"
        ),
        8: (
            "本文的贡献包括三个方面。第一，明确界定既有身份空间中的监督可靠性问题，并给出不引入新ID的局部处理约束；"
            "第二，构建基于语义响应和大范围连通区域风险的置信度加权身份监督，保持ILGS的高斯表示及其他训练分支；第三，"
            "在三个LERF场景的受控实验、标准四场景LERF-OVS评测和3D-OVS十场景外部验证中开展匹配种子的重复比较，"
            "同时报告平均增益、运行间波动、退化案例与适用边界。"
        ),
        56: (
            "图2  IPR的可靠性建模与作用路径。a，原始身份标签中的背景泄漏；b，结构清理、跨视角语义证据和面积先验共同"
            "形成像素置信图，并用于加权身份损失；c，置信加权训练后的身份感知语言高斯在新视角产生语言、身份与分段相似度；"
            "d，同一文本查询下，错误身份区域受到抑制，预测集中到目标对象。定量结论见第4节。"
        ),
        59: (
            "3DGS训练包含随机视角采样、参数初始化和自适应密度控制，单次运行可能将方法差异与优化波动混在一起。受控"
            "LERF实验使用Figurines、Ramen和Teatime三个场景；标准LERF-OVS评测进一步加入Waldo Kitchen。每个设置分别"
            "以训练种子0、1、2独立训练，并保持方法间的图像划分、输入分辨率、语言特征及自编码器配置一致，训练均进行"
            "30000次迭代。实验使用两张NVIDIA RTX 3090，每个训练任务占用一张GPU。IPR在三场景LERF消融结果上确定最终"
            "配置，再以固定参数开展3D-OVS外部验证；因此，该三场景结果同时承担方法选择与评估功能，不能视为独立于配置"
            "选择的测试。标准四场景LERF-OVS结果作为公开协议下的补充验证单独报告。"
        ),
        60: (
            "受控LERF与3D-OVS比较统一采用seed_ref查询掩码生成：优先选择得分最高的非零身份作为种子，在非环绕的身份"
            "邻接图上扩展，对正种子相似度使用相对种子得分的3%差值阈值，并排除背景ID 0。该公共解码器与旧official模式"
            "在种子、参考相似度和阈值等方面不同，本文将其视为固定评测条件，不将解码器变化带来的收益计入IPR贡献。标准"
            "LERF-OVS评测则使用官方查询解码和帧-查询对聚合；两套协议及其数值不混用。"
        ),
        61: (
            "受控实验的mIoU和mBIoU首先按类别汇总各评测视图，再对场景内类别等权平均。边界带宽采用图像对角线的2%，"
            "最少1像素；预测尺寸与GT不一致时使用最近邻插值。标准LERF-OVS按帧-查询对聚合，报告mIoU和阈值0.25下的"
            "mAcc@0.25。每个设置均由匹配种子0、1、2独立运行；均值用于方法比较，样本标准差仅描述运行间波动。由于n=3"
            "不足以可靠刻画单场景结果分布，本文不据此作单场景显著性推断。跨场景平均均为场景均值的等权平均，不表示像素"
            "加权汇总。"
        ),
        62: (
            "三场景受控LERF与3D-OVS实验的候选词表复用各场景评测查询文本，未使用真实像素分割标注生成IPR权重；其"
            "预处理还允许无标注留出视图参与。因此，这部分结果属于场景词表已知、无标注多视图可参与预处理的设定。作为"
            "更严格的补充，标准LERF-OVS实验仅处理训练视图，并以查询无关的CLIP裁剪一致性估计对象可靠性，不读取评测"
            "查询、像素标注或留出视图。两种设置分别回答受控机制比较和公开协议泛化问题。"
        ),
        71: "4.5 局部清理与可靠性消融",
        72: (
            "表3先比较局部清理与基础置信度，再比较附加可靠性项。所有附加设置均建立在相同的清理标签和基础置信图上；"
            "“无语义”与“无大区域”中的“无”只针对附加权重，不表示移除了基础置信图。S表示语义降权，G表示大区域降权，"
            "V表示训练视图中的跨视角可见性乘子。面积投票在补充表中展开，正文仅保留关键比较。"
        ),
        74: "表3  三个LERF场景与三个种子的受控消融（场景平均，%）",
        81: "4.6 3D-OVS外部验证",
        84: "表4  3D-OVS十场景三种子受控对照（均值 ± 样本标准差；每格为mIoU / mBIoU，%）",
        86: (
            "表4显示，IPR的十场景等权平均mIoU为87.07%，高于受控基线的82.01%；mBIoU为81.22%，高于基线的75.69%，"
            "分别增加5.06和5.53个百分点。按场景均值比较，两个指标均在9/10个场景提高；Covered Desk的mIoU和mBIoU"
            "分别下降2.15和1.99个百分点。Bench、Blue Sofa、Office Desk和Snacks贡献了主要增量，而Bed、Lawn、Sofa和"
            "Table的改善较小。因此，整体提升主要来自部分错误对象对应的恢复，不能解释为对所有场景的均匀改善。"
        ),
        96: "4.7 局限性与证据边界",
        97: (
            "首先，三场景受控LERF与3D-OVS实验中的语义可靠性依赖给定场景词表，未涵盖对象或多个相似候选可能产生较低"
            "置信度；评测查询与候选词表重合、无标注留出图像可参与预处理也限定了该部分结果的适用范围。标准LERF-OVS"
            "四场景实验通过仅使用训练视图的查询无关一致性估计消除了这两项依赖，但覆盖场景仍有限。其次，面积规则可能误伤"
            "真实大对象，高置信度的错误身份也可能被保留，完全遗漏的对象无法通过降权恢复。3D-OVS中Covered Desk的单视角"
            "大幅退化和Sofa的持续失败表明，外部场景中的错误并未被本方法普遍消除。"
        ),
        98: (
            "再次，最终配置是在LERF消融中选择，并以固定参数迁移到3D-OVS。三次重复能够暴露训练波动，但不足以稳定估计"
            "单场景方差或建立分布假设，因此本文不据此作单场景显著性声明。标准LERF-OVS上的总体收益主要由Waldo Kitchen"
            "贡献，Ramen与Figurines的mIoU略有下降，不能据此声称逐场景一致改进。3D-OVS总体平均改善与场景级95%置信区间"
            "相符，但区间仍较宽，且Covered Desk持续退化；Bed、Blue Sofa和Room的跨种子波动也表明，当前证据不能被表述为"
            "对任意场景均稳定改善。较大的外部收益主要集中于若干查询的对象对应恢复，现有实验尚不能把因果贡献独立分配给"
            "局部清理、语义响应或大区域降权中的某一项。后续需在更多外部数据集上进行多种子验证，并补充受控ID合并、背景"
            "泄漏、缺失区域、权重分布以及预处理时间和存储开销分析。"
        ),
        100: (
            "本文研究ILGS既有身份空间中的伪标签可靠性问题。IPR通过局部清理、面积加权语义证据汇总及大范围区域风险估计"
            "生成像素置信图，并作用于二维身份分类损失。最终配置省去额外的逐视角训练可见性乘子，保留原有ID词表和三维身份"
            "感知学习结构。面向标准LERF-OVS的严格协议，本文进一步采用仅训练视图、查询无关的跨视角CLIP一致性估计，避免"
            "评测词表和留出视图进入可靠性预处理。"
        ),
        101: (
            "在统一损失归一化和查询解码的三个LERF场景三种子受控实验中，IPR取得79.63% mIoU与74.46% mBIoU，相对"
            "受控基线提高3.74和3.21个百分点。在标准四场景LERF-OVS评测中，三种子平均mIoU由50.51%提高至52.02%，"
            "mAcc@0.25由58.41%提高至60.23%；Waldo Kitchen上的提升最明显，其余三个场景总体相当。在固定参数的3D-OVS"
            "十场景三种子实验中，IPR取得87.07% mIoU与81.22% mBIoU，提高5.06和5.53个百分点，并在9/10个场景的均值上"
            "同时改善两项指标。消融结果表明，基础置信度与区域风险降权对总体结果有贡献，语义增量则具有场景依赖性。"
            "Covered Desk的退化、Waldo Kitchen之外较集中的标准协议收益以及若干场景的跨种子波动共同限定了结论范围。"
        ),
    }

    for index, text in replacements.items():
        replace_paragraph(paragraphs[index], text)

    anchor = paragraphs[70]
    heading = insert_paragraph_after(
        document,
        anchor,
        paragraphs[71],
        "4.4 标准LERF-OVS四场景评测",
    )
    protocol = insert_paragraph_after(
        document,
        heading,
        paragraphs[70],
        "为检验受控实验结论能否在公开评测协议下复现，本文补充Ramen、Figurines、Teatime和Waldo Kitchen四个标准"
        "LERF-OVS场景。官方ILGS与IPR采用相同训练划分、输入、语言特征、官方查询解码和匹配随机种子。严格IPR配置仅"
        "利用训练视图，以查询无关的CLIP裁剪一致性估计对象可靠性，阈值固定为0.75；该过程不使用评测查询文本、像素标注"
        "或留出视图。评测遵循帧-查询对聚合并报告mIoU与mAcc@0.25。",
    )
    caption = insert_paragraph_after(
        document,
        protocol,
        paragraphs[68],
        "表2  标准四场景LERF-OVS三种子结果（每格为mIoU / mAcc@0.25，%；正文表报告三种子均值）",
    )
    caption.paragraph_format.keep_with_next = True

    table_element = deepcopy(document.tables[0]._tbl)
    table_element.tr_lst[-1].addprevious(deepcopy(table_element.tr_lst[-2]))
    caption._p.addnext(table_element)
    table = Table(table_element, document._body)

    rows = [
        ("场景", "官方ILGS", "ILGS + IPR"),
        ("Ramen", "51.32 / 57.28", "51.19 / 56.81"),
        ("Figurines", "53.97 / 59.52", "53.77 / 60.12"),
        ("Teatime", "66.70 / 72.88", "67.24 / 74.01"),
        ("Waldo Kitchen", "30.04 / 43.94", "35.88 / 50.00"),
        ("四场景平均", "50.51 / 58.41", "52.02 / 60.23"),
    ]
    for row, values in zip(table.rows, rows):
        for cell, value in zip(row.cells, values):
            set_cell_text(cell, value)

    result = insert_paragraph_after(
        document,
        table,
        paragraphs[70],
        "表2显示，IPR将四场景平均mIoU从50.51%提高至52.02%，将mAcc@0.25从58.41%提高至60.23%，分别增加"
        "1.51和1.82个百分点。提升主要来自Waldo Kitchen：mIoU和mAcc@0.25分别增加5.84和6.06个百分点；Ramen与"
        "Figurines的mIoU分别下降0.13和0.20个百分点，Teatime提高0.54个百分点。因此，该结果支持IPR在噪声较重场景中"
        "缓解监督偏差，但不支持逐场景一致提升的表述。",
    )
    insert_paragraph_after(
        document,
        result,
        paragraphs[70],
        "为保持主表可读性，表2仅给出三次独立运行的均值。逐场景样本标准差范围为：官方ILGS的mIoU 0.52–2.01、"
        "mAcc@0.25 0.81–2.94，IPR的mIoU 0.36–1.84、mAcc@0.25 0.00–1.96。三种子结果仅用于检查随机优化波动，"
        "不作为显著性证据，也不将复验本身列为方法贡献。",
    )

    # Keep each table caption with its header and at least one data row.
    for caption_paragraph in (paragraphs[68], paragraphs[74], paragraphs[84]):
        caption_paragraph.paragraph_format.keep_with_next = True
    for current_table in (document.tables[0], document.tables[1], document.tables[2], document.tables[3]):
        for cell in current_table.rows[0].cells:
            cell.paragraphs[0].paragraph_format.keep_with_next = True
        for cell in current_table.rows[1].cells:
            cell.paragraphs[0].paragraph_format.keep_with_next = True

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)

    check = Document(OUTPUT)
    print(f"Saved: {OUTPUT}")
    print(f"Paragraphs: {len(check.paragraphs)}")
    print(f"Tables: {len(check.tables)}")
    print(f"New heading present: {any(p.text == '4.4 标准LERF-OVS四场景评测' for p in check.paragraphs)}")
    print(f"New table rows: {len(check.tables[1].rows)}")


if __name__ == "__main__":
    main()
