from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt


DOCX_PATH = Path(r"D:\wlxy\ILGS-IPR.docx")

BODY_FONT = "宋体"
EN_FONT = "Times New Roman"
MATH_FONT = "Cambria Math"


REPLACEMENTS = {
    31: (
        "body",
        "设训练图像集合为 {Iₜ}ₜ₌₁ᵀ，对应的原始 object mask 为 Mᵢᵗ。其中，Mᵢᵗ(x)=1 "
        "表示第 t 个视角中像素 x 被分配给 object id i，Mᵢᵗ(x)=0 表示不属于该 id。"
        "本文用 Yₜ(x) 表示像素级 identity 伪标签，用 wₜ(x) 表示像素级置信度权重；"
        "当 Yₜ(x)=i 时，wₜ(x) 由 object id i 在第 t 个视角中的可靠性决定。ILGS 将 Mᵢᵗ "
        "作为 identity 伪标签监督渲染出的 identity feature map，并进一步利用 identity 信息约束 "
        "language embeddings 的多视角一致性。",
    ),
    32: (
        "body",
        "本文关注的问题是：当 Mᵢᵗ 中存在 noisy identity labels 时，如何在不增加人工标注、"
        "不重构 ILGS 主干的前提下降低错误监督。与重新生成完整 mask 不同，本文假设原 object id "
        "中仍包含有价值的跨视角 identity anchor，因此采用保守 refinement：保留大部分 id 结构，"
        "只为每个像素估计置信度 wₜ(x)，并在训练中降低不可靠区域的影响。",
    ),
    33: ("formula", "Yₜ(x)=i,    wₜ(x)∈[0,1]                                      (1)"),
    36: (
        "body",
        "设 object id i 在第 k 个 crop 上对类别 c 的 softmax 相似度为 pᵢ,ₖ,꜀，"
        "其可见性权重为 αᵢ,ₖ，则该 object id 对类别 c 的累积语义响应 Sᵢ,꜀ 定义为式（2）。"
        "语义置信度 qᵢ 定义为最高类别响应占全部类别响应的比例，如式（3）所示；"
        "semantic margin mᵢ 定义为 top-1 与 top-2 类别响应之间的差距，如式（4）所示。"
        "当 qᵢ 或 mᵢ 较低时，说明该 object id 的语义归属不稳定，可能混入多个物体或背景区域。",
    ),
    37: ("formula", "Sᵢ,꜀ = ∑ₖ αᵢ,ₖ pᵢ,ₖ,꜀                                      (2)"),
    38: ("formula", "qᵢ = max꜀ Sᵢ,꜀ / (∑꜀ Sᵢ,꜀ + ε)                              (3)"),
    39: ("formula", "mᵢ = (Sᵢ,꜀₁ − Sᵢ,꜀₂) / (∑꜀ Sᵢ,꜀ + ε)                         (4)"),
    43: (
        "body",
        "本文参考 visibility-aware keyframe selection 的思想，将面积变化转化为可见性权重，"
        "而不是直接删除样本。设 Aᵢᵗ=|Mᵢᵗ|/|Ωₜ| 为 object id i 在第 t 个视角中的 mask 面积比例，"
        "其中 Ωₜ 表示第 t 个视角的图像像素域；Aᵢ,med 为其跨视角面积比例中位数，"
        "σ 控制对面积偏差的容忍度，则面积可见性权重 vᵢᵗ 可写为：",
    ),
    44: ("formula", "vᵢᵗ = Aᵢᵗ · exp(−(Aᵢᵗ − Aᵢ,med)² / (2σ²))                  (5)"),
    51: (
        "formula",
        "L′ᵢd = [∑ₓ wₜ(x) CE(pₜ(x), Yₜ(x))] / [∑ₓ wₜ(x)+ε]             (6)",
    ),
    52: (
        "body",
        "其中，zₜ(x) 为渲染得到的 identity logits，pₜ(x)=softmax(zₜ(x)) 为对应的 identity "
        "概率分布，Yₜ(x) 为 object id 伪标签，wₜ(x) 为本文估计的像素级置信度，ε 为防止分母为 0 "
        "的极小常数。该设计不需要激进删除伪标签，也不会完全破坏跨视角 id 对应关系；"
        "不可靠区域仍参与训练，但对参数更新的影响显著降低。",
    ),
}


def clear_paragraph(paragraph):
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def set_paragraph(paragraph, text, kind):
    clear_paragraph(paragraph)
    if kind == "formula":
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.first_line_indent = None
        paragraph.paragraph_format.space_before = Pt(3)
        paragraph.paragraph_format.space_after = Pt(3)
        font_name = MATH_FONT
    else:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        paragraph.paragraph_format.first_line_indent = Pt(21)
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        font_name = EN_FONT

    run = paragraph.add_run(text)
    run.font.name = font_name
    run.font.size = Pt(10.5)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT if kind == "body" else MATH_FONT)


def main():
    doc = Document(str(DOCX_PATH))
    for index, (kind, text) in REPLACEMENTS.items():
        set_paragraph(doc.paragraphs[index], text, kind)
    doc.save(str(DOCX_PATH))
    print(f"saved {DOCX_PATH}")


if __name__ == "__main__":
    main()
