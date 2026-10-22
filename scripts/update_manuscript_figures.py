from copy import deepcopy
from pathlib import Path
import shutil
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


DOC_KEY = "\u516c\u5f0f\u4fee\u6b63\u7248"
DOC_DIR = Path(r"D:\wlxy")
FIG_DIR = Path(r"C:\Users\hp\Downloads\paper_figures_compact")


def find_doc() -> Path:
    matches = [
        p for p in DOC_DIR.glob("*.docx")
        if DOC_KEY in p.name
        and "\u5907\u4efd" not in p.name
        and not p.name.startswith("~$")
    ]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one formula-fixed docx, found: {matches}")
    return matches[0]


def set_run_font(run, size=10.5, bold=False):
    run.font.name = "\u5b8b\u4f53"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "\u5b8b\u4f53")
    run.font.size = Pt(size)
    run.bold = bold


def format_body(paragraph):
    paragraph.paragraph_format.first_line_indent = Cm(0.74)
    paragraph.paragraph_format.line_spacing = 1.25
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    for run in paragraph.runs:
        set_run_font(run, 10.5)


def format_caption(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = None
    paragraph.paragraph_format.line_spacing = 1.15
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(3)
    for run in paragraph.runs:
        set_run_font(run, 9)


def insert_after(anchor, paragraphs):
    current = anchor
    for paragraph in paragraphs:
        current._p.addnext(paragraph._p)
        current = paragraph
    return current


def new_paragraph(doc, text="", caption=False):
    paragraph = doc.add_paragraph()
    if text:
        run = paragraph.add_run(text)
        set_run_font(run, 9 if caption else 10.5)
    if caption:
        format_caption(paragraph)
    else:
        format_body(paragraph)
    return paragraph


def new_picture_paragraph(doc, image_path: Path, width_cm=15.2):
    if not image_path.exists():
        raise FileNotFoundError(image_path)
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = None
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=Cm(width_cm))
    return paragraph


def remove_existing_inserted_block(doc):
    markers = [
        "\u56fe 1  teatime",
        "\u56fe 2  ramen",
        "\u56fe 3  figurines",
        "\u56fe 2  teatime",
        "\u56fe 3  ramen",
        "\u56fe 4  figurines",
        "\u4e3a\u8fdb\u4e00\u6b65\u5c55\u793a IPR",
        "\u4ece\u53ef\u89c6\u5316\u7ed3\u679c\u770b",
    ]
    for paragraph in list(doc.paragraphs):
        text = paragraph.text.strip()
        if any(marker in text for marker in markers):
            paragraph._element.getparent().remove(paragraph._element)


def find_anchor(doc):
    for paragraph in doc.paragraphs:
        if paragraph.text.strip().startswith(
            "\u8fdb\u4e00\u6b65\u89c2\u5bdf\u7c7b\u522b\u7ea7\u7ed3\u679c"
        ):
            return paragraph
    raise RuntimeError("Could not find result-analysis anchor paragraph.")


def add_bottom_border(paragraph, size="8", color="auto"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def update_doc():
    doc_path = find_doc()
    backup = doc_path.with_name(doc_path.stem + "_插图前备份.docx")
    shutil.copy2(doc_path, backup)

    doc = Document(doc_path)
    remove_existing_inserted_block(doc)

    anchor = find_anchor(doc)
    inserted = []

    inserted.append(new_paragraph(
        doc,
        "\u4e3a\u8fdb\u4e00\u6b65\u5c55\u793a IPR \u4e0e Full-refine "
        "\u7684\u4f5c\u7528\u65b9\u5f0f\uff0c\u672c\u6587\u4ece\u4e09\u4e2a\u573a\u666f\u4e2d\u9009\u53d6"
        "\u5177\u6709\u4ee3\u8868\u6027\u7684\u53ef\u89c6\u5316\u6837\u4f8b\u3002\u56fe 1 \u5c55\u793a teatime "
        "\u4e2d\u591a\u7c7b\u5178\u578b\u5931\u6548\u60c5\u51b5\uff0c\u5176\u4e2d sheep \u548c spoon handle "
        "\u5206\u522b\u5bf9\u5e94\u533a\u57df\u6269\u5f20\u6f02\u79fb\u548c\u90e8\u4ef6\u7ea7\u5c0f\u76ee\u6807\u5206\u5272\u4e24\u7c7b\u95ee\u9898\u3002"
        "\u4e0e Official-fixed \u76f8\u6bd4\uff0cIPR + Full-refine \u80fd\u591f\u5c06\u8fc7\u5ea6\u6269\u5f20\u7684\u7ed3\u679c"
        "\u6536\u7f29\u5230\u4e0e\u76ee\u6807\u66f4\u4e00\u81f4\u7684 identity \u533a\u57df\uff0c\u5e76\u5728 spoon handle "
        "\u8fd9\u7c7b\u6781\u5c0f\u533a\u57df\u4e0a\u6062\u590d\u6709\u6548\u9884\u6d4b\u3002"
    ))
    inserted.append(new_picture_paragraph(
        doc, FIG_DIR / "fig_teatime_mask_comparison_compact.png", width_cm=15.2
    ))
    inserted.append(new_paragraph(
        doc,
        "\u56fe 1  teatime \u573a\u666f\u4e2d sheep \u548c spoon handle \u7684\u63a9\u7801\u5bf9\u6bd4\u7ed3\u679c\u3002RGB \u4e3a\u6e32\u67d3\u56fe\u50cf\uff0c"
        "GT \u4e3a\u4eba\u5de5\u6807\u6ce8\uff0cOfficial-fixed \u4e3a\u4fee\u590d\u6587\u4ef6\u547d\u540d\u4e0e\u6d4b\u8bd5\u7d22\u5f15\u540e\u7684"
        "\u5b98\u65b9\u63a9\u7801\u751f\u6210\u6d41\u7a0b\uff0cIPR + Full-refine \u4e3a\u672c\u6587\u63a8\u8350\u540e\u5904\u7406\u7ed3\u679c\u3002",
        caption=True,
    ))

    inserted.append(new_paragraph(
        doc,
        "\u56fe 2 \u7ed9\u51fa ramen \u573a\u666f\u4e2d pork belly \u548c glass of water \u7684 full-refine \u5019\u9009\u63a9\u7801\u4e2d\u95f4\u8fc7\u7a0b\u3002"
        "pork belly \u793a\u4f8b\u8868\u660e\uff0c\u5355\u7eaf\u91c7\u7528\u5b98\u65b9\u6269\u5f20\u7ed3\u679c\u65f6\u4f1a\u5c06\u76ee\u6807\u533a\u57df"
        "\u8fde\u540c\u5468\u56f4\u98df\u6750\u4e00\u8d77\u6fc0\u6d3b\uff0c\u800c seed-ref \u80fd\u4fdd\u7559\u4e0e\u76ee\u6807\u66f4\u5339\u914d\u7684\u4e3b\u8981"
        " identity \u7ec4\u4ef6\u3002\u5bf9 glass of water \u548c egg \u7b49\u7c7b\u522b\uff0c\u4e0d\u540c\u5019\u9009\u7ed3\u679c\u7684"
        "\u5dee\u5f02\u8f83\u5c0f\uff0c\u8bf4\u660e\u8be5\u7b56\u7565\u4e0d\u4f1a\u5bf9\u5df2\u7ecf\u53ef\u9760\u7684\u9884\u6d4b\u9020\u6210\u660e\u663e\u7834\u574f\u3002"
    ))
    inserted.append(new_picture_paragraph(
        doc, FIG_DIR / "fig_ramen_full_refine_candidates_compact.png", width_cm=15.2
    ))
    inserted.append(new_paragraph(
        doc,
        "\u56fe 2  ramen \u573a\u666f\u4e2d Full-refine \u5019\u9009\u63a9\u7801\u7684\u9009\u62e9\u8fc7\u7a0b\u3002"
        "Full official \u8868\u793a\u5b98\u65b9\u6269\u5f20\u7ed3\u679c\uff0cSeed-ref \u8868\u793a\u4fdd\u7559\u521d\u59cb identity "
        "\u79cd\u5b50\u533a\u57df\uff0cSelected \u4e3a\u6839\u636e\u5931\u6548\u5224\u65ad\u9009\u51fa\u7684\u6700\u7ec8\u63a9\u7801\u3002",
        caption=True,
    ))

    inserted.append(new_paragraph(
        doc,
        "\u56fe 3 \u5219\u5c55\u793a figurines \u573a\u666f\u7684\u6539\u8fdb\u4e0e\u5c40\u9650\u3002"
        "red toy chair \u5728\u5b98\u65b9\u6d41\u7a0b\u4e2d\u51fa\u73b0"
        "\u63a9\u7801\u7f3a\u5931\uff0cIPR + Full-refine \u80fd\u6062\u590d\u90e8\u5206\u76ee\u6807\u533a\u57df\uff1b\u4f46 old camera "
        "\u4ecd\u7136\u96be\u4ee5\u5b8c\u5168\u6062\u590d\uff0c\u8fd9\u8bf4\u660e\u5f53\u6587\u672c\u76f8\u4f3c\u5ea6\u6392\u5e8f\u672c\u8eab\u5c06\u4e0d\u5b8c\u6574\u7684"
        "\u5c0f\u8fde\u901a\u57df\u9009\u4e3a\u9996\u9009 identity \u65f6\uff0c\u4ec5\u4f9d\u8d56\u540e\u5904\u7406\u96be\u4ee5\u4ece\u9519\u8bef\u79cd\u5b50\u4e2d"
        "\u5b8c\u5168\u6062\u590d\u76ee\u6807\u3002\u56e0\u6b64\uff0c\u672c\u6587\u5728\u63a8\u8350\u7ed3\u679c\u4e2d\u5bf9 figurines \u66f4\u503e\u5411"
        "\u91c7\u7528 IPR \u800c\u975e\u7edf\u4e00\u542f\u7528 Full-refine\u3002"
    ))
    inserted.append(new_picture_paragraph(
        doc, FIG_DIR / "fig_figurines_mask_comparison_compact.png", width_cm=15.2
    ))
    inserted.append(new_paragraph(
        doc,
        "\u56fe 3  figurines \u573a\u666f\u4e2d red toy chair \u548c old camera \u7684\u6539\u8fdb\u6548\u679c\u4e0e\u5931\u8d25\u6848\u4f8b\u3002"
        "\u8be5\u56fe\u7528\u4e8e\u8bf4\u660e Full-refine \u5bf9\u90e8\u5206\u7c7b\u522b\u6709\u6548\uff0c\u4f46\u5728\u6587\u672c\u76f8\u4f3c\u5ea6"
        "\u521d\u59cb\u9009\u62e9\u9519\u8bef\u7684\u60c5\u51b5\u4e0b\u4ecd\u5b58\u5728\u5c40\u9650\u3002",
        caption=True,
    ))

    insert_after(anchor, inserted)

    # Strengthen discussion with an explicit figure-backed interpretation.
    discussion_anchor = None
    for paragraph in doc.paragraphs:
        if paragraph.text.strip().startswith("\u8fd9\u79cd\u62c6\u5206\u4e5f\u89e3\u91ca"):
            discussion_anchor = paragraph
            break
    if discussion_anchor is not None:
        extra = new_paragraph(
            doc,
            "\u7ed3\u5408\u56fe 1\u2014\u56fe 3 \u53ef\u4ee5\u770b\u5230\uff0c"
            "Full-refine \u7684\u4e3b\u8981\u4ef7\u503c\u5e76\u4e0d\u5728\u4e8e\u5bf9\u6240\u6709\u7c7b\u522b\u7edf\u4e00\u63d0\u9ad8"
            "\u76f8\u4f3c\u5ea6\uff0c\u800c\u662f\u8bc6\u522b\u5b98\u65b9\u63a9\u7801\u751f\u6210\u4e2d\u7684\u5178\u578b\u5931\u6548\u6a21\u5f0f\u3002"
            "\u5f53\u5b98\u65b9\u6269\u5f20\u7ed3\u679c\u51fa\u73b0\u5927\u9762\u79ef\u6f02\u79fb\u65f6\uff0c\u79cd\u5b50\u4fdd\u62a4\u53ef\u4ee5\u907f\u514d"
            "\u76ee\u6807\u4e0e\u80cc\u666f\u6216\u76f8\u90bb\u7269\u4f53\u7c98\u8fde\uff1b\u5f53\u76ee\u6807\u5c5e\u4e8e\u90e8\u4ef6\u7ea7\u5c0f\u533a\u57df\u65f6\uff0c"
            "\u5c40\u90e8\u8fde\u901a\u57df\u7b5b\u9009\u53ef\u4ee5\u51cf\u5c11\u5927 identity \u533a\u57df\u5bf9\u5c0f\u76ee\u6807\u7684\u6df9\u6ca1\u3002"
            "\u4f46\u662f\uff0c\u5982\u679c\u6587\u672c\u67e5\u8be2\u5728\u521d\u59cb\u9636\u6bb5\u5df2\u7ecf\u9009\u5230\u9519\u8bef\u6216\u4e25\u91cd\u4e0d\u5b8c\u6574"
            "\u7684 identity\uff0c\u540e\u5904\u7406\u7684\u53ef\u6062\u590d\u7a7a\u95f4\u4f1a\u53d7\u5230\u660e\u663e\u9650\u5236\u3002"
        )
        discussion_anchor._p.addnext(extra._p)

    # Put a thin separator after the last newly inserted visual block to keep layout tidy.
    for paragraph in doc.paragraphs:
        if paragraph.text.strip().startswith("\u56fe 3  figurines"):
            add_bottom_border(paragraph)
            break

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text.startswith("\u56fe 1  3DOVS"):
            paragraph.text = text.replace("\u56fe 1", "\u56fe 4", 1)
            format_caption(paragraph)
            break

    doc.save(doc_path)
    print(f"updated: {doc_path}")
    print(f"backup:  {backup}")


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    update_doc()
