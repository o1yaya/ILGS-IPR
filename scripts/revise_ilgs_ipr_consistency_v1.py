from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph


SOURCE = Path(r"D:\wlxy\ILGS-IPR_中文.docx")
OUTPUT = Path(r"D:\wlxy\ILGS-IPR_中文_方法一致性修订_v5.docx")
FIGURE = Path(
    r"D:\PycharmProjects\PythonProject\ILGS_release-main\tmp_ilgs_outputs"
    r"\pipeline_figure_v23_for_docx.png"
)


def el(tag):
    return OxmlElement(tag)


def math_run(text):
    run = el("m:r")
    text_node = el("m:t")
    text_node.text = text
    run.append(text_node)
    return run


def math_container(tag, items):
    node = el(tag)
    for item in items:
        node.append(math_run(item) if isinstance(item, str) else item)
    return node


def sub(base, subscript):
    node = el("m:sSub")
    node.append(math_container("m:e", [base]))
    node.append(math_container("m:sub", [subscript]))
    return node


def sup(base, superscript):
    node = el("m:sSup")
    base_items = base if isinstance(base, list) else [base]
    node.append(math_container("m:e", base_items))
    node.append(math_container("m:sup", [superscript]))
    return node


def subsup(base, subscript, superscript):
    node = el("m:sSubSup")
    node.append(math_container("m:e", [base]))
    node.append(math_container("m:sub", [subscript]))
    node.append(math_container("m:sup", [superscript]))
    return node


def frac(numerator, denominator):
    node = el("m:f")
    node.append(math_container("m:num", numerator))
    node.append(math_container("m:den", denominator))
    return node


def clear_paragraph(paragraph):
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def set_run_font(run, ascii_name="Times New Roman", east_asia="宋体", size=10.5, bold=False):
    run.font.name = ascii_name
    run.font.size = Pt(size)
    run.font.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), ascii_name)
    rfonts.set(qn("w:hAnsi"), ascii_name)
    rfonts.set(qn("w:eastAsia"), east_asia)


def add_body_run(paragraph, text, bold=False):
    run = paragraph.add_run(text)
    set_run_font(run, bold=bold)
    return run


def add_code_run(paragraph, text):
    run = paragraph.add_run(text)
    set_run_font(run, ascii_name="Consolas", east_asia="宋体", size=9.5)
    return run


def append_inline_math(paragraph, items):
    math = el("m:oMath")
    for item in items:
        math.append(math_run(item) if isinstance(item, str) else item)
    paragraph._p.append(math)


def set_paragraph_pieces(paragraph, pieces):
    clear_paragraph(paragraph)
    for kind, value in pieces:
        if kind == "text":
            add_body_run(paragraph, value)
        elif kind == "bold":
            add_body_run(paragraph, value, bold=True)
        elif kind == "code":
            add_code_run(paragraph, value)
        elif kind == "math":
            append_inline_math(paragraph, value)
        else:
            raise ValueError(f"Unknown piece kind: {kind}")


def set_paragraph_with_code_terms(paragraph, text, code_terms):
    pieces = []
    cursor = 0
    while cursor < len(text):
        matches = [(text.find(term, cursor), term) for term in code_terms if text.find(term, cursor) >= 0]
        if not matches:
            pieces.append(("text", text[cursor:]))
            break
        index, term = min(matches, key=lambda item: item[0])
        if index > cursor:
            pieces.append(("text", text[cursor:index]))
        pieces.append(("code", term))
        cursor = index + len(term)
    set_paragraph_pieces(paragraph, pieces)


def word_run(text, size=10.5):
    run = el("w:r")
    rpr = el("w:rPr")
    rfonts = el("w:rFonts")
    rfonts.set(qn("w:ascii"), "Times New Roman")
    rfonts.set(qn("w:hAnsi"), "Times New Roman")
    rfonts.set(qn("w:eastAsia"), "宋体")
    rpr.append(rfonts)
    sz = el("w:sz")
    sz.set(qn("w:val"), str(int(size * 2)))
    rpr.append(sz)
    run.append(rpr)
    text_node = el("w:t")
    text_node.text = text
    run.append(text_node)
    return run


def word_tab():
    run = el("w:r")
    run.append(el("w:tab"))
    return run


def set_formula(paragraph, items, number):
    clear_paragraph(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = None
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.keep_together = True

    ppr = paragraph._p.get_or_add_pPr()
    tabs = ppr.find(qn("w:tabs"))
    if tabs is not None:
        ppr.remove(tabs)
    tabs = el("w:tabs")
    tab = el("w:tab")
    tab.set(qn("w:val"), "right")
    tab.set(qn("w:pos"), str(int(6.2 * 1440)))
    tabs.append(tab)
    ppr.append(tabs)

    math = el("m:oMath")
    for item in items:
        math.append(math_run(item) if isinstance(item, str) else item)
    paragraph._p.append(math)
    paragraph._p.append(word_tab())
    paragraph._p.append(word_run(f"({number})"))


def paragraph_after(paragraph):
    new_p = OxmlElement("w:p")
    if paragraph._p.pPr is not None:
        new_p.append(deepcopy(paragraph._p.pPr))
    paragraph._p.addnext(new_p)
    return Paragraph(new_p, paragraph._parent)


def replace_figure(paragraph):
    clear_paragraph(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run()
    picture = run.add_picture(str(FIGURE), width=Inches(6.55))
    picture._inline.docPr.set("descr", "Semantic-Area IPR pipeline and outlier-aware weighted identity loss")


def set_cell(cell, text, bold=False):
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    set_run_font(run, size=9.5, bold=bold)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def update_table_four(doc):
    table = doc.tables[3]
    while len(table.rows) < 6:
        table.add_row()
    data = [
        ["探索性设置", "mIoU (%)", "mBIoU (%)"],
        ["修正后的官方基线", "65.99", "64.22"],
        ["Semantic-Area IPR", "76.30", "73.02"],
        ["SAM2语义掩码", "74.06", "70.99"],
        ["DEVA + SAM2对齐", "70.58", "68.22"],
        ["DINOv3/SAM2 V2-anchor严格匹配", "71.19", "68.13"],
    ]
    for row_index, row_data in enumerate(data):
        row = table.rows[row_index]
        prevent_row_split(row)
        for col_index, value in enumerate(row_data):
            set_cell(row.cells[col_index], value, bold=(row_index == 0 or row_index == 2))


def align_all_tables_left(doc):
    for table in doc.tables:
        for row in table.rows:
            prevent_row_split(row)
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for paragraph in cell.paragraphs:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT


def keep_table_together(table):
    """Keep compact result tables on one page when Word can fit them."""
    for row in table.rows[:-1]:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_with_next = True


def main():
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    if not FIGURE.exists():
        raise FileNotFoundError(FIGURE)

    doc = Document(str(SOURCE))
    p = doc.paragraphs
    if len(p) < 107:
        raise RuntimeError(f"Unexpected paragraph count: {len(p)}")

    replace_figure(p[26])

    p[10].text = (
        "本文的主要贡献包括以下三个方面：（1）将ILGS中的object_mask噪声进一步表述为身份锚点可靠性问题，"
        "并指出背景泄漏、对象合并和跨视角身份漂移会使身份感知损失放大错误监督；（2）提出一种保持锚点结构的"
        "语义—面积可靠性建模方法，在不重新生成完整轨迹、不改变原始身份标签的情况下生成像素级置信图；"
        "（3）通过三个LERF场景上的主实验、消融实验、定性分析以及更强二维模型的探索性实验，验证所提方法的"
        "有效性及其适用边界。"
    )

    p[27].text = (
        "图1  Semantic-Area IPR在ILGS中的改进流程。IPR在训练前对原始身份标签进行可靠性建模，输出保持ID不变的"
        "标签和像素级置信图。训练时，置信图与ILGS官方CE离群点掩码共同作用于二维身份分类损失；RGB重建、CLIP"
        "特征监督、身份感知语义一致性、ILGS骨干网络及推理流程均保持不变。"
    )
    p[28].text = (
        "图1展示了Semantic-Area IPR在ILGS流程中的位置。原始RGB视图和对象ID作为跨视角身份锚点输入IPR。"
        "IPR不重新生成对象轨迹，也不修改对象ID，而是根据语义置信度、跨视角面积可见性和大范围混合区域风险生成"
        "像素级置信图。ILGS训练阶段仍先根据渲染身份logits与原始身份标签计算逐像素交叉熵，并使用官方阈值生成"
        "离群点掩码；随后将该掩码与IPR置信图相乘，计算置信度加权的二维身份分类损失。"
    )

    set_paragraph_pieces(p[30], [
        ("text", "设训练图像集合为"),
        ("math", [sub("I", "t"), ", t=1,…,T"]),
        ("text", "，"),
        ("math", [sub("Ω", "t")]),
        ("text", "表示第"),
        ("math", ["t"]),
        ("text", "个视角的像素域，对应对象ID "),
        ("math", ["i"]),
        ("text", "的原始二值掩码记为"),
        ("math", [subsup("M", "i", "t")]),
        ("text", "。像素级身份标签记为"),
        ("math", [sub("Y", "t"), "(x)"]),
        ("text", "，像素级可靠性权重记为"),
        ("math", [sub("w", "t"), "(x)"]),
        ("text", "。ILGS利用身份标签监督渲染身份logits，并以该身份信息约束语言高斯场。"),
    ])
    p[31].text = (
        "本文只优化身份监督的可靠性，不重新生成二维掩码或全局对象轨迹。式（1）表示IPR保持原始身份标签及其"
        "跨视角锚点关系不变，仅为每个像素计算可靠性权重。"
    )
    set_formula(p[32], [sup(sub("Y", "t"), "′"), "(x)=", sub("Y", "t"), "(x),    0≤", sub("w", "t"), "(x)≤1"], 1)

    set_paragraph_pieces(p[34], [
        ("text", "对于每个对象ID，本文从其多个视角的可见区域提取连通分量。仅保留面积大于"),
        ("code", "min_component_area"),
        ("text", "的连通分量，并选取面积最大的至多"),
        ("code", "max_crops_per_id"),
        ("text", "个裁剪区域。实验中二者分别设为500和40。每个裁剪区域在目标外保留"),
        ("code", "crop_padding"),
        ("text", "个像素的上下文，避免纯掩码裁剪丢失必要的语义信息。"),
    ])
    set_paragraph_pieces(p[35], [
        ("text", "设对象ID "), ("math", ["i"]), ("text", "在第"), ("math", ["k"]),
        ("text", "个裁剪区域上与类别"), ("math", ["c"]), ("text", "的softmax概率为"),
        ("math", [subsup("p", "i", "k"), "(c)"]), ("text", "，对应的面积可见性权重为"),
        ("math", [subsup("v", "i", "k")]), ("text", "。累积语义响应"),
        ("math", [sub("S", "i"), "(c)"]), ("text", "、语义置信度"),
        ("math", [sub("q", "i")]), ("text", "和top-1/top-2归一化间隔"),
        ("math", [sub("m", "i")]), ("text", "分别由式（2）至式（4）定义，其中"),
        ("math", [sub("c", "1")]), ("text", "和"), ("math", [sub("c", "2")]),
        ("text", "表示响应最高和次高的类别。"),
    ])
    set_formula(p[36], [sub("S", "i"), "(c)=", sub("Σ", "k"), subsup("v", "i", "k"), subsup("p", "i", "k"), "(c)"], 2)
    set_formula(p[37], [sub("q", "i"), "=", frac([sub("max", "c"), sub("S", "i"), "(c)"], [sub("Σ", "c"), sub("S", "i"), "(c)+ε"])], 3)
    set_formula(p[38], [sub("m", "i"), "=", frac([sub("S", "i"), "(", sub("c", "1"), ")−", sub("S", "i"), "(", sub("c", "2"), ")"], [sub("Σ", "c"), sub("S", "i"), "(c)+ε"])], 4)
    set_paragraph_pieces(p[39], [
        ("text", "较低的语义置信度并不直接表示该ID错误。代码采用离散降权规则：当"),
        ("math", [sub("q", "i")]), ("text", "低于"), ("code", "semantic_conf_threshold = 0.45"),
        ("text", "，或"), ("math", [sub("m", "i")]), ("text", "低于"),
        ("code", "semantic_margin_threshold = 0.08"), ("text", "时，将该ID的语义权重设为"),
        ("code", "low_semantic_weight = 0.65"), ("text", "；否则保持为1。该规则保留不确定区域，但降低其训练贡献。"),
    ])

    set_paragraph_pieces(p[42], [
        ("text", "设"), ("math", [subsup("A", "i", "t"), "=|", subsup("M", "i", "t"), "|"]),
        ("text", "为对象ID "), ("math", ["i"]), ("text", "在视角"), ("math", ["t"]),
        ("text", "中的连通分量像素面积，"), ("math", [sub("A", "i,med")]),
        ("text", "为其跨视角中位数。代码先计算未归一化分数"), ("math", [subsup("r", "i", "t")]),
        ("text", "，再除以同一ID的最大分数并截断至"), ("math", ["[", sub("w", "min"), ",1]"]),
        ("text", "，得到式（5）所示的面积可见性权重。"),
    ])
    set_formula(p[43], [
        subsup("r", "i", "t"), "=", subsup("A", "i", "t"), "exp",
        frac(["−", sup(["(√", subsup("A", "i", "t"), "−√", sub("A", "i,med"), ")"], "2")], ["2", sup("σ", "2")]),
        ",    ", subsup("v", "i", "t"), "=max{", sub("w", "min"), ",",
        frac([subsup("r", "i", "t")], [sub("max", "j"), subsup("r", "i", "j")]), "}",
    ], 5)
    set_paragraph_with_code_terms(p[44], (
        "式（5）与实现保持一致：可见性分数基于连通分量的像素面积和平方根面积偏差，并在同一对象ID内部归一化。"
        "实验中，visibility_sigma设为100，min_visibility_weight设为0.25。异常视角仍被保留，但其语义投票和训练权重较低。"
    ), ["visibility_sigma", "min_visibility_weight"])

    set_paragraph_pieces(p[47], [
        ("text", "大范围混合区域通过面积比例统计检测。令"),
        ("math", [subsup("a", "i", "t"), "="]),
        ("math", [frac(["|", subsup("M", "i", "t"), "|"], ["|", sub("Ω", "t"), "|"])]),
        ("text", "。本文统计其中位数、P90和最大值，如式（6）所示。当任一统计量超过对应阈值时，该ID被优先标记为大范围混合区域，"
                 "其语义权重设为large_region_weight。"),
    ])

    set_formula(p[50], [
        subsup("a", "i", "t"), "=", frac(["|", subsup("M", "i", "t"), "|"], ["|", sub("Ω", "t"), "|"]),
        ",  ", sub("a", "i,med"), "=", sub("median", "t"), "(", subsup("a", "i", "t"), ")",
        ",  ", sub("a", "i,p90"), "=", sub("P90", "t"), "(", subsup("a", "i", "t"), ")",
        ",  ", sub("a", "i,max"), "=", sub("max", "t"), "(", subsup("a", "i", "t"), ")",
    ], 6)
    p[47]._p.addnext(p[50]._p)

    set_paragraph_pieces(p[49], [
        ("text", "代码使用语义权重"), ("math", [sub("s", "i")]),
        ("text", "表示对象级可靠性。若ID被标记为大范围混合区域，则"), ("math", [sub("s", "i")]),
        ("text", "设为"), ("code", "large_region_weight = 0.35"), ("text", "；否则，当"),
        ("math", [sub("q", "i"), "<0.45"]), ("text", "或"), ("math", [sub("m", "i"), "<0.08"]),
        ("text", "时，"), ("math", [sub("s", "i")]), ("text", "设为"),
        ("code", "low_semantic_weight = 0.65"), ("text", "；其余ID的"), ("math", [sub("s", "i")]),
        ("text", "为1。该优先级规则与实现中的if/elif判断一致。设"),
        ("math", [subsup("c", "t", "0"), "(x)"]), ("text", "为输入置信图；若未提供，则其初值为1。最终像素级权重由式（7）给出。"),
    ])
    set_formula(p[51], [
        sub("w", "t"), "(x)=clip(", subsup("c", "t", "0"), "(x)·", sub("s", "i"), "·", subsup("v", "i", "t"), ",",
        sub("c", "min"), ",1),    x∈", subsup("M", "i", "t"),
    ], 7)
    set_formula(p[52], [
        subsup("L", "cls", "IPR"), "=", frac(["1"], ["log C"]), "·",
        frac(
            [sub("Σ", "x"), sub("o", "t"), "(x)", sub("w", "t"), "(x)CE(", sub("p", "t"), "(x),", sub("Y", "t"), "(x))"],
            ["max{", sub("Σ", "x"), sub("o", "t"), "(x)", sub("w", "t"), "(x),1}"],
        ),
    ], 8)
    formula9 = paragraph_after(p[52])
    set_formula(formula9, [
        subsup("L", "total", "IPR"), "=", sub("λ", "rgb"), sub("L", "rgb"), "+", sub("λ", "cls"), subsup("L", "cls", "IPR"),
        "+", sub("λ", "clip"), sub("L", "clip"), "+", sub("λ", "cons"), sub("L", "cons"),
    ], 9)
    set_paragraph_pieces(p[53], [
        ("text", "训练阶段首先计算逐像素交叉熵"),
        ("math", [sub("e", "t"), "(x)=CE(", sub("p", "t"), "(x),", sub("Y", "t"), "(x))"]),
        ("text", "，并沿用ILGS官方阈值生成有效像素掩码"),
        ("math", [sub("o", "t"), "(x)=1[", sub("e", "t"), "(x)≤τ(t)]"]),
        ("text", "。式（8）使用"), ("math", [sub("o", "t"), "(x)"]),
        ("text", "与IPR置信权重"), ("math", [sub("w", "t"), "(x)"]),
        ("text", "的乘积对二维身份分类损失进行加权，分子和分母均只统计有效像素，分母下限截断为1，并保留官方代码中的"),
        ("math", ["log C"]), ("text", "归一化。式（9）给出总目标，其中RGB重建损失、CLIP特征损失和身份感知语义一致性损失均保持不变；IPR只改变二维身份分类损失的计算。"),
    ])

    p[57].text = (
        "所有实验均基于ILGS官方流程复现。不同可靠性建模设置采用相同的场景划分、ILGS骨干网络、训练迭代次数和评价协议。"
        "实验保留ILGS官方掩码处理及CE阈值离群点过滤，Semantic-Area IPR只新增训练前可靠性估计和置信度加权的二维身份"
        "分类损失。修正后的官方基线仅修复官方掩码生成或评价兼容性问题，不使用IPR置信图。实验场景为teatime、figurines和ramen。"
    )
    parameter_text = (
        "参数设置与代码默认值一致。min_component_area = 500，max_crops_per_id = 40，crop_padding = 12，"
        "visibility_sigma = 100，min_visibility_weight = 0.25。semantic_conf_threshold = 0.45，"
        "semantic_margin_threshold = 0.08，low_semantic_weight = 0.65。大范围混合区域的中位数、P90和最大面积比例阈值"
        "分别为0.55、0.35和0.45，large_region_weight = 0.35。最终置信图截断下限min_confidence = 0.05，"
        "CLIP softmax缩放系数clip_softmax_scale = 35。"
    )
    set_paragraph_with_code_terms(p[59], parameter_text, [
        "min_component_area", "max_crops_per_id", "crop_padding", "visibility_sigma",
        "min_visibility_weight", "semantic_conf_threshold", "semantic_margin_threshold",
        "low_semantic_weight", "large_region_weight", "min_confidence", "clip_softmax_scale",
    ])
    p[64].text = (
        "为验证各模块的贡献，本文比较Only area、Only semantic、Semantic-Area no-conf和Full Semantic-Area IPR。"
        "Only area的平均mIoU为73.03%，但平均mBIoU低于修正后的官方基线，说明单独使用面积权重无法区分真实小目标与异常区域。"
        "Only semantic将平均mBIoU提高至69.87%，但无法捕捉跨视角面积异常。Semantic-Area no-conf融合语义与面积信息，"
        "但不使用最终像素级置信图，因此残留噪声仍以相同强度参与身份监督。完整方法取得最高的平均mIoU和mBIoU。"
    )
    p[65].paragraph_format.page_break_before = True
    p[65].paragraph_format.keep_with_next = True
    p[72].paragraph_format.keep_with_next = True
    p[69].text = (
        "图3至图5展示组合系统的定性结果，其中标记为IPR + Full-refine的结果同时包含Semantic-Area IPR训练和可选推理后处理，"
        "因此这些图不单独用于归因训练阶段IPR的增益。图3中的teatime勺柄属于细长小目标；组合系统能够在抑制背景假阳性的同时"
        "保留目标位置。对于sheep查询，结果也显示背景泄漏有所减少。"
    )

    p[82].text = "表4  figurines场景上的探索性结果及主要方法参照（mIoU / mBIoU，%）"
    p[82].paragraph_format.page_break_before = True
    p[82].paragraph_format.keep_with_next = True
    update_table_four(doc)
    p[81].text = (
        "本文在figurines场景上探索了SAM2、DEVA以及DINOv3/SAM2 V2-anchor等更强二维策略。表4同时列出修正后的官方基线"
        "和Semantic-Area IPR作为参照。SAM2语义掩码能够改善部分类别的边界和局部目标，但其mIoU和mBIoU仍低于"
        "Semantic-Area IPR；另外两种跨视角对齐策略也未超过IPR。该实验用于分析适用边界，不作为跨数据集的普遍结论。"
    )

    align_all_tables_left(doc)
    for table_index in (1, 2, 3):
        keep_table_together(doc.tables[table_index])

    for index in (10, 27, 28, 31, 44, 49, 53, 57, 59, 64, 69, 81, 82):
        for run in p[index].runs:
            set_run_font(run)

    doc.save(str(OUTPUT))
    print(f"saved={OUTPUT}")
    print(f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)}")


if __name__ == "__main__":
    main()
